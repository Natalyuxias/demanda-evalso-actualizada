"""Matriz del ferry ARISTO: conteos medidos y asignación OD condicionada.

Reutiliza las lecturas verificadas de cuatro tareas. Nunca ejecuta Etapas 07–13.
La asignación principal exige misma fecha, sentido y clase. El perfil agrupado
de campaña solo se conserva como contraste, no completa días sin muestra.
"""
from pathlib import Path
import argparse
import hashlib
import json
import math
import time

import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Font, PatternFill

KEY = ['campana', 'fecha', 'sentido_cruce', 'segmento']
CAMPAIGN = ['campana', 'sentido_cruce', 'segmento']
MODES = {'Auto': 'VL', 'Camioneta': 'VL', 'Vehículo privado': 'VP',
         'Camión o furgón de carga': 'CAM', 'Pasajero (peatón)': 'PEATON',
         'Moto': 'MOTO', 'Bicicleta': 'BICI'}
SENSE = {'De Chile Chico a Pto. Ibáñez': 'Chile Chico → Puerto Ibáñez',
         'De Pto. Ibáñez a Chile Chico': 'Puerto Ibáñez → Chile Chico'}


def sha(p):
    return hashlib.sha256(Path(p).read_bytes()).hexdigest()


def rows(d):
    return d.astype(object).where(pd.notna(d), None).to_dict('records')


def workbook(path, sheets):
    w = Workbook(write_only=True)
    for title, frame in sheets.items():
        s = w.create_sheet(title); s.freeze_panes = 'B2'
        header = []
        for i, name in enumerate(frame.columns, 1):
            s.column_dimensions[__import__('openpyxl').utils.get_column_letter(i)].width = min(55, max(18, len(name)+3))
            c = WriteOnlyCell(s, name); c.font = Font(bold=True, color='FFFFFF'); c.fill = PatternFill('solid', fgColor='225D64'); header.append(c)
        s.append(header)
        for record in rows(frame):
            cells = []
            for value in record.values():
                cell = WriteOnlyCell(s, value)
                if isinstance(value, str):
                    assert len(value) <= 32767
                    cell.data_type = 's'
                cells.append(cell)
            s.append(cells)
        s.auto_filter.ref = f'A1:{__import__("openpyxl").utils.get_column_letter(len(frame.columns))}{len(frame)+1}'
    w.save(path)


def assign(observations, strata):
    """N incluye todas las respuestas válidas del estrato, aunque no tengan zona."""
    n = observations.groupby(KEY, dropna=False).size().rename('N_recalculado').reset_index()
    st = strata.merge(n, how='left', on=KEY, validate='one_to_one')
    st['N_recalculado'] = st.N_recalculado.fillna(0).astype(int)
    assert st.N_recalculado.eq(st.N_entrevistas_fecha).all()
    a = observations.merge(st, how='left', on=KEY, validate='many_to_one', suffixes=('', '_conteo'))
    a['coef_asignacion_fecha'] = a.F_total_filas_fecha.div(a.N_recalculado.where(a.N_recalculado.gt(0)))
    a['aporte_flujo_medido'] = a.coef_asignacion_fecha
    return a, st


def aggregate(a, keys):
    # min_count=1 mantiene el faltante para grupos sin conteo compatible.
    g = a.groupby(keys, dropna=False, sort=True)
    d = g.agg(entrevistas_validas=('id_encuesta', 'size'),
              entrevistas_con_coeficiente=('aporte_flujo_medido', 'count'),
              fechas_encuestas=('fecha', lambda s: '|'.join(sorted(set(s)))),
              coeficiente_minimo=('coef_asignacion_fecha', 'min'),
              coeficiente_maximo=('coef_asignacion_fecha', 'max')).reset_index()
    q = g.aporte_flujo_medido.sum(min_count=1).rename('flujo_asignado_exposicion_medida').reset_index()
    return d.merge(q, on=keys, validate='one_to_one')


def describe(d):
    d = d.copy()
    d['procedimiento'] = 'Suma por fecha de F_fecha × n_OD_fecha / N_todas_validas_fecha; composición condicionada'
    d['cobertura_temporal'] = 'Solo filas de embarque medidas de fechas indicadas; no 24h, semana, estación completa o año'
    d['confiabilidad'] = 'Asignación OD condicional; no probabilidades de inclusión ni representatividad certificadas'
    d['unidad'] = d.segmento.map({'VL': 'vehículos livianos cuatro ruedas', 'CAM': 'vehículos de carga',
                                'VP': 'unidades privadas agregadas (VL+moto+bici), no VL estricto',
                                'PEATON': 'personas registradas como peatones, no pasajeros totales',
                                'MOTO': 'motocicletas', 'BICI': 'bicicletas',
                                'OTRO_PASAJERO': 'entrevistas de personas, flujo faltante', 'BUS_TP': 'buses, no pasajeros'})
    d['fuente'] = 'Cap6 Base de datos agregada.xlsx/BBDD-1+BBDD-2 por ID2; Consolidado Flujos.xlsx/Transbordadores'
    d['supuestos'] = 'Perfil de respuestas válido representa composición OD del flujo contado en misma fecha/sentido/clase; unidad encuestada y no respuesta no certificadas'
    d['uso_proyecto'] = 'Usuarios del ferry actual; elegibilidad/captación Bahía Jara–Levicán no estimadas'
    return d


def balances(st, a):
    assigned = a[a.OD_zonal_completo].groupby(KEY).aporte_flujo_medido.sum(min_count=1).rename('flujo_asignado_zonas').reset_index()
    b = st.merge(assigned, how='left', on=KEY, validate='one_to_one')
    # Cero de asignación contable si no hay muestra/zona: no es cero de demanda.
    b['flujo_asignado_zonas'] = b.flujo_asignado_zonas.fillna(0).where(b.F_total_filas_fecha.notna())
    b['flujo_sin_OD_zonal'] = b.F_total_filas_fecha - b.flujo_asignado_zonas
    b['motivo_no_asignacion'] = ['Conteo no identificado' if pd.isna(f) else 'Sin entrevistas de clase en misma fecha; no transferir perfil de otro día' if not n else 'OD genérico/ambiguo sin zona' if residual > 1e-10 else 'Completo contablemente; no valida representatividad' for f, n, residual in zip(b.F_total_filas_fecha, b.N_recalculado, b.flujo_sin_OD_zonal)]
    return b


def main():
    start = time.monotonic()
    p = argparse.ArgumentParser(); p.add_argument('--revision', type=Path, required=True)
    p.add_argument('--anexos', type=Path, required=True); p.add_argument('--salida', type=Path, required=True)
    args = p.parse_args(); r, root, out = args.revision, args.anexos, args.salida
    if out.exists():
        raise FileExistsError('No sobrescribir: elegir una carpeta nueva')
    protected = json.loads((r/'integridad_antes.json').read_text())
    assert all(sha(x['archivo']) == x['sha256'] for x in protected)
    verified_sources = json.loads((r/'microdatos/final/validacion_microdatos.json').read_text())['sha256_fuentes']
    assert all(sha(root/path) == digest for path, digest in verified_sources.items())
    micro = pd.read_csv(r/'microdatos/final/microdatos_ferry_264.csv')
    usr = pd.read_csv(r/'usuarios/auditoria_usuarios_264.csv')
    assert micro.id_encuesta.is_unique and usr.id_encuesta.is_unique and set(micro.id_encuesta) == set(usr.id_encuesta)
    indicators = ['nacionalidad_estado', 'patente_estado', 'nacionalidad_extranjera_observada', 'patente_extranjera_observada', 'OD_internacional_observado', 'ambas_nacionalidad_patente_desconocidas', 'alerta_moto_ocupacion_atipica', 'alerta_peaton_columna_ocupantes']
    a = micro.merge(usr[['id_encuesta']+indicators], on='id_encuesta', validate='one_to_one')
    a['sentido_cruce_original'] = a.sentido_cruce
    a['sentido_cruce'] = a.sentido_cruce.map(SENSE); assert a.sentido_cruce.notna().all()
    a['segmento'] = a.tipo_usuario_BBDD2.map(MODES).fillna('OTRO_PASAJERO')
    assert a.loc[a.segmento.eq('VP'), 'campana'].eq('Otoño').all()
    a['fecha'] = pd.to_datetime(a['b1__Día']).dt.strftime('%Y-%m-%d')
    a['origen'] = a.origen_nombre_BBDD2; a['destino'] = a.destino_nombre_BBDD2
    a['zona_origen'] = a.origen_zona_modelacion_ARISTO.astype('Int64')
    a['zona_destino'] = a.destino_zona_modelacion_ARISTO.astype('Int64')
    a['OD_zonal_completo'] = a.zona_origen.notna() & a.zona_destino.notna()
    a['OD_referencia_completo'] = a.origen.notna() & a.destino.notna()
    a['nivel_geografico'] = 'Nombre/escala BBDD2; no comuna convertida a localidad'
    a['alerta_validacion_cruzada'] = a['b1__Duplicada'].eq(1)
    allst = pd.read_csv(r/'aforos/conteos_y_coeficientes_por_fecha.csv')
    pooled = pd.read_csv(r/'aforos/estratos_y_coeficientes.csv')
    use = pooled[pooled.uso_principal_sugerido][CAMPAIGN].drop_duplicates()
    st = allst.merge(use, on=CAMPAIGN, how='inner', validate='many_to_one')
    # Las clases privadas agrupadas de Otoño sustituyen VL/MOTO/BICI en principal.
    assert not st[st.campana.eq('Otoño')].segmento.isin(['VL', 'MOTO', 'BICI']).any()
    # Otros pasajeros no tienen conteo de personas compatible; se conservan sin flujo.
    other = a[a.segmento.eq('OTRO_PASAJERO')].groupby(KEY).size().rename('N_entrevistas_fecha').reset_index()
    for col in st.columns:
        if col not in other: other[col] = None
    st = pd.concat([st, other[st.columns]], ignore_index=True)
    a, st = assign(a, st)
    a['estado_OD'] = a.OD_zonal_completo.map({True: 'OD zonal identificado', False: 'OD parcial: extremo internacional genérico sin zona'})
    a['evidencia_ferry'] = 'Interceptación en ruta Puerto Ibáñez–Chile Chico; campo original Rampa o Transbordador, sin bandera individual de embarque'
    a['elegibilidad_nuevo_proyecto'] = 'No evaluada: demanda física del ferry conservada, incluidos viajes largos'
    a['fuente_conteo'] = 'Cap 6/Mediciones de flujo/Consolidado Flujos.xlsx; hoja Transbordadores; filas_fuente del estrato'
    zone = describe(aggregate(a[a.OD_zonal_completo], CAMPAIGN+['zona_origen', 'zona_destino']))
    detail = describe(aggregate(a[a.OD_zonal_completo], CAMPAIGN+['origen', 'destino', 'zona_origen', 'zona_destino']))
    partial = describe(aggregate(a[~a.OD_zonal_completo], CAMPAIGN+['origen', 'destino', 'zona_origen', 'zona_destino']))
    dated = describe(aggregate(a[a.OD_zonal_completo], KEY+['origen', 'destino', 'zona_origen', 'zona_destino']))
    bal = balances(st, a)
    # Detalle referencial original, no topónimo ni domicilio exacto certificado.
    for side, search, direction in [('origen', 'b1__Búsqueda Origen', 'b1__Dirección Origen'), ('destino', 'b1__Búsqueda destino', 'b1__Dirección destino')]:
        a['referencia_'+side] = a[search].fillna(a[direction]).fillna(a[side])
    refs = describe(aggregate(a, CAMPAIGN+['origen', 'destino', 'referencia_origen', 'referencia_destino', 'zona_origen', 'zona_destino']))
    refs['advertencia_geocodificacion'] = 'Texto original de búsqueda/dirección; puede ser referencia geocodificada. No acredita punto exacto ni localidad desagregada.'
    # Tres dimensiones independientes: nunca sumar sus resultados como extranjeros únicos.
    foreign = []
    for dimension, column, eligible in [('nacionalidad_declarada_peaton', 'nacionalidad_estado', a.segmento.eq('PEATON')),
                                        ('patente_vehicular', 'patente_estado', a.segmento.isin(['VL', 'VP', 'CAM', 'MOTO', 'BICI'])),
                                        ('OD_internacional', 'OD_internacional_observado', pd.Series(True, index=a.index))]:
        d = a[eligible].copy(); d['categoria_indicador'] = d[column].astype(str)
        f = describe(aggregate(d, KEY+['categoria_indicador']))
        f['indicador'] = dimension; f['interpretacion'] = 'Flujo de la unidad de clase con ese indicador bajo igual asignación condicional; no nacionalidad de ocupantes ni residencia'
        foreign.append(f)
    foreign = pd.concat(foreign, ignore_index=True)
    # Comparación material con 07: misma aritmética histórica, separada del observado.
    old = pd.read_csv(root/'demanda_lago_github/outputs/matrices_od_referencia/matriz_od_barcaza.csv')
    old['campana'] = old.estacion; old['sentido_cruce'] = old.sentido_control.map({'Chile Chico a Ibanez': 'Chile Chico → Puerto Ibáñez', 'Ibanez a Chile Chico': 'Puerto Ibáñez → Chile Chico'})
    oldsum = old.groupby(CAMPAIGN, dropna=False).agg(N_07=('encuestas_validas', 'sum'), Q_07_diario=('demanda_expandida_viajes_diarios_promedio', lambda s: s.sum(min_count=1))).reset_index()
    compare = oldsum.merge(pooled, on=CAMPAIGN, how='left', validate='one_to_one')
    compare['frecuencia_07_semana'] = compare.campana.map({'Primavera': 8, 'Verano': 11, 'Otoño': 8})
    compare['Q_07_reconstruida_misma_regla'] = compare.media_por_fila*compare.frecuencia_07_semana/7
    compare.loc[compare.N_07.eq(0), 'Q_07_reconstruida_misma_regla'] = None
    compare['diferencia_aritmetica_07'] = compare.Q_07_diario - compare.Q_07_reconstruida_misma_regla
    compare['comparabilidad'] = 'Q07 diario condicionado no comparable directamente con total observado. Reconstrucción solo verifica regla histórica, no estima demanda diaria actual.'
    cmpbal = bal.groupby(CAMPAIGN).agg(F_nuevo_medido=('F_total_filas_fecha', lambda s: s.sum(min_count=1)), Q_nuevo_zonas=('flujo_asignado_zonas', lambda s: s.sum(min_count=1)), sin_OD_nuevo=('flujo_sin_OD_zonal', lambda s: s.sum(min_count=1))).reset_index()
    compare = compare.merge(cmpbal, on=CAMPAIGN, how='outer', validate='one_to_one')
    compare['alerta_clase'] = compare.segmento.map({'VP': 'VP Otoño no equivale a VL07; incluye unidades privadas sin desagregación en EOD', 'VL': 'VL Otoño07 imponía privado=VL; nuevo principal usa VP'}).fillna('Unidad de clase conservada')
    # Sensibilidad de estratificación: igual F y perfil, fecha vs campaña (sin frecuencias).
    profile = a.merge(pooled[CAMPAIGN+['F_total_filas_medidas', 'N_entrevistas_validas_clase']], on=CAMPAIGN, how='left', validate='many_to_one')
    profile['coef_campana_SOLO_contraste'] = profile.F_total_filas_medidas/profile.N_entrevistas_validas_clase.where(profile.N_entrevistas_validas_clase.gt(0))
    sensitivity = profile.groupby(CAMPAIGN+['origen', 'destino', 'zona_origen', 'zona_destino'], dropna=False).agg(Q_por_fecha=('aporte_flujo_medido', lambda s: s.sum(min_count=1)), Q_perfil_campana=('coef_campana_SOLO_contraste', lambda s: s.sum(min_count=1))).reset_index()
    sensitivity['diferencia_por_estratificacion'] = sensitivity.Q_por_fecha-sensitivity.Q_perfil_campana
    sensitivity['uso'] = 'Fecha es principal; campaña es contraste que transfiere composición a días sin muestra, no rellenar principal'
    # Balances por clase, sin total que mezcle personas y vehículos o campañas.
    summary = {'entrevistas_originales': 290, 'entrevistas_validas_BBDD2': len(a),
               'OD_zonal_completo': int(a.OD_zonal_completo.sum()), 'OD_zonal_incompleto': int((~a.OD_zonal_completo).sum()),
               'entrevistas_con_coeficiente': int(a.aporte_flujo_medido.notna().sum()), 'entrevistas_sin_conteo_personas': int(a.segmento.eq('OTRO_PASAJERO').sum()),
               'entrevistas_OD_zonal_y_coeficiente': int((a.OD_zonal_completo & a.aporte_flujo_medido.notna()).sum()),
               'celdas_zonales': len(zone), 'celdas_zonales_con_flujo': int(zone.flujo_asignado_exposicion_medida.notna().sum()),
               'celdas_detalle': len(detail), 'celdas_detalle_con_flujo': int(detail.flujo_asignado_exposicion_medida.notna().sum()),
               'filas_07_historicas': len(old), 'filas_07_con_muestra': int(old.encuestas_validas.gt(0).sum()),
               'filas_07_estructurales_sin_muestra': int(old.encuestas_validas.eq(0).sum()),
               'estratos_fecha_principales': len(st), 'conteos_zarpe_identificados': False, 'pasajeros_totales_estimados': None,
               'nacionalidad_extranjera_muestra': int(a.nacionalidad_extranjera_observada.sum()),
               'patente_extranjera_muestra': int(a.patente_extranjera_observada.sum()), 'OD_internacional_muestra': int(a.OD_internacional_observado.sum()),
               'matriz_CITRA_modificada': False, 'fuentes_sumadas': False, 'frecuencias_en_principal': False}
    checks = {
        '264_ID2_unicos': a.id_encuesta.nunique() == len(a) == 264,
        '258_con_coef_6_personas_sin_aforo_compatible': a.aporte_flujo_medido.notna().sum() == 258,
        'N_todas_validas_misma_fecha_sin_filtro_OD': st.N_recalculado.eq(st.N_entrevistas_fecha).all(),
        'sin_trasladar_perfil_a_fechas_sin_muestra': bal.loc[bal.N_recalculado.eq(0) & bal.F_total_filas_fecha.notna(), 'flujo_asignado_zonas'].eq(0).all(),
        'balance_F_asignado_mas_residual': ((bal.flujo_asignado_zonas+bal.flujo_sin_OD_zonal-bal.F_total_filas_fecha).abs().dropna() < 1e-9).all(),
        '6_OD_genericos_no_imputados': (~a.OD_zonal_completo).sum() == 6,
        'sin_VP_mas_componentes_en_principal': not st[st.campana.eq('Otoño')].segmento.isin(['VL', 'MOTO', 'BICI']).any(),
        'indicadores_extranjeros_cotejados_auditoria_independiente': all(int(a[c].sum()) == int(usr[c].sum()) for c in ['nacionalidad_extranjera_observada', 'patente_extranjera_observada', 'OD_internacional_observado']),
        'geografia_y_coordenadas_conservadas': all(micro[c].equals(a[c]) for c in micro if c != 'sentido_cruce'),
        'totales_07_reconstruidos_sin_reestimacion': compare.diferencia_aritmetica_07.abs().dropna().lt(1e-9).all(),
        'no_se_estiman_pasajeros_totales': summary['pasajeros_totales_estimados'] is None,
        'originales_previos_y_CITRA_intactos': all(sha(x['archivo']) == x['sha256'] for x in protected),
    }
    # Reconciliar detalle y zona sin convertir ausencias de flujo en cero poblacional.
    dettotal = detail.groupby(CAMPAIGN+['zona_origen', 'zona_destino']).flujo_asignado_exposicion_medida.sum(min_count=1)
    ztotal = zone.set_index(CAMPAIGN+['zona_origen', 'zona_destino']).flujo_asignado_exposicion_medida
    checks['balance_detalle_zonas'] = bool((dettotal-ztotal).abs().dropna().lt(1e-9).all() and dettotal.isna().equals(ztotal.isna()))
    checks = {k: bool(v) for k, v in checks.items()}; assert all(checks.values()), checks
    out.mkdir(parents=True); (out/'csv').mkdir()
    tables = {'matriz_zonas_ARISTO': zone, 'matriz_detalle_BBDD2': detail, 'matriz_por_fecha': dated,
              'OD_parciales_sin_zona': partial, 'referencias_originales': refs, 'auditoria_264': a,
              'factores_y_N_fecha': st, 'balance_por_fecha': bal, 'balance_por_campana': cmpbal,
              'extranjeros_indicadores_separados': foreign, 'comparacion_07': compare,
              'contraste_perfil_campana_vs_fecha': sensitivity}
    for name, frame in tables.items(): frame.to_csv(out/'csv'/f'{name}.csv', index=False, encoding='utf-8-sig')
    notes = pd.DataFrame([
        {'tema': 'Estimación', 'detalle': 'Demanda física observada en embarques medidos. OD asignado condicionalmente por misma fecha, sentido y clase. No demanda capturable, 24h, estación completa o año.'},
        {'tema': 'Coeficientes', 'detalle': 'F/N transfiere composición de respuestas. N>F en primavera; no es factor de inclusión individual ni diseño representativo certificado.'},
        {'tema': 'Clases', 'detalle': 'Otoño: VP agregado sustituye VL/MOTO/BICI en principal. No sumar VP con componentes. Buses y pasajeros de buses son distintas unidades.'},
        {'tema': 'Pasajeros', 'detalle': 'Peatones son aforo peatonal, no total de pasajeros. Ocupantes y grupos descriptivos, sin multiplicación automática; total de personas no identificado.'},
        {'tema': 'Geografía', 'detalle': 'Originales y zonas55 separados de IDs BBDD2. Seis OD internacionales genéricos sin zona; conservados sin imputación. Búsquedas/direcciones son referencias, no ubicaciones exactas certificadas.'},
        {'tema': 'Frecuencias', 'detalle': '9 semanales en TDoriginal y11/8 ordinarias confirmadas se conservan en contraste. Ninguna frecuencia se aplica a matriz observada principal.'},
        {'tema': 'Fuentes', 'detalle': 'Excel: Base de datos agregada/BBDD1+BBDD2+TDTransb; Consolidado Flujos/Transbordadores; DiccionarioAPP; Cap8/Zonas. PDFsecciones/tablas en informeHTML y archivos de tareas.'},
    ])
    workbook(out/'01_matrices_lacustres.xlsx', {'Leer primero': notes, 'Zonas ARISTO': zone, 'Detalle BBDD2': detail, 'Por fecha': dated, 'OD parciales no asignados': partial, 'Referencias originales': refs, 'Balance por fecha': bal})
    workbook(out/'02_auditoria_encuestas.xlsx', {'Leer primero': notes, '264 validas enlazadas': a, '290 originales ferry': pd.read_csv(r/'microdatos/final/290_originales_ferry.csv'), '26 fuera BBDD2': pd.read_csv(r/'microdatos/final/26_originales_no_en_BBDD2.csv'), 'Catalogo zonas55': pd.read_csv(r/'microdatos/final/catalogo_original_zonas_55.csv'), 'Cruce IDs y zonas': pd.read_csv(r/'microdatos/final/correspondencias_BBDD2_zonas.csv')})
    workbook(out/'03_conteos_factores_comparacion.xlsx', {'Leer primero': notes, 'Conteos originales': pd.read_csv(r/'aforos/filas_originales_conteos.csv'), 'Conteos normalizados': pd.read_csv(r/'aforos/conteos_normalizados.csv'), 'N y coef fecha': st, 'Balance fecha': bal, 'Contraste campaña': pooled, 'Comparacion matriz07': compare, 'Perfil fecha vs campaña': sensitivity, 'TDTransb original': pd.read_csv(r/'aforos/auditoria_TD_Transb.csv'), 'Excel vs Informe': pd.read_csv(r/'aforos/contraste_Excel_Informe.csv'), 'Frecuencias NO observado': pd.read_csv(r/'aforos/escenarios_frecuencia_NO_observados.csv')})
    workbook(out/'04_pasajeros_extranjeros.xlsx', {'Leer primero': notes, 'Indicadores y flujos': foreign, 'Auditoria usuarios': usr, 'Intersecciones no sumar': pd.read_csv(r/'usuarios/intersecciones_con_desconocidos.csv'), 'Ocupacion descriptiva': pd.read_csv(r/'usuarios/ocupacion_grupos_descriptivos.csv'), 'Disponibilidad campos': pd.read_csv(r/'usuarios/disponibilidad_campos_por_usuario.csv')})
    summary['verificaciones'] = checks; summary['segundos_integracion'] = round(time.monotonic()-start, 2)
    (out/'validacion.json').write_text(json.dumps(summary, ensure_ascii=False, indent=2)+'\n')
    from informe_lacustre import generate
    generate(out, r, summary, zone, detail, bal, cmpbal, compare)
    sourcefiles = [r/'integridad_antes.json'] + [x for folder in ['microdatos/final', 'aforos', 'usuarios', 'comparacion'] for x in (r/folder).rglob('*') if x.is_file() and x.suffix in ['.csv', '.json', '.py', '.txt']]
    sourcefiles += list(Path(__file__).parent.glob('*.py'))
    manifest = {'fuentes': [{'archivo': str(x.resolve()), 'sha256': sha(x)} for x in sourcefiles],
                'resultados': [{'archivo': str(x.relative_to(out)), 'sha256': sha(x)} for x in out.rglob('*') if x.is_file()], 'protegidos_sin_cambios': checks['originales_previos_y_CITRA_intactos']}
    (out/'manifest_integridad.json').write_text(json.dumps(manifest, ensure_ascii=False, indent=2)+'\n')
    print(json.dumps(summary, ensure_ascii=False, indent=2))


if __name__ == '__main__': main()
