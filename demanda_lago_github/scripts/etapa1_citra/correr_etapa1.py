"""Etapa1 CITRA reproducible: originales + capa geográfica y decisiones aprobadas.
No importa modelos de demanda, logit ni resultados de ARISTO/Etapa13.
"""
from pathlib import Path
import argparse,csv,json,collections,copy,math,sys,html,hashlib,datetime
import geopandas as gp
from shapely.geometry import Point
from openpyxl import Workbook,load_workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Font,PatternFill
from fuentes import readcsv,sha,read_originals,kmz,expansion,canon,norm,num
from red import Red
from pantallas import build

def writecsv(path,rows,fields=None):
 rows=list(rows);fields=fields or list(dict.fromkeys(k for r in rows for k in r)) or ['Sin registros']
 with Path(path).open('w',encoding='utf-8-sig',newline='') as f:
  w=csv.DictWriter(f,fieldnames=fields);w.writeheader();w.writerows(rows)

def workbook(path,sheets):
 w=Workbook(write_only=True)
 for title,rows,fields in sheets:
  rows=list(rows);fields=fields or list(dict.fromkeys(k for r in rows for k in r)) or ['Sin registros'];s=w.create_sheet(title);s.freeze_panes='B2';s.auto_filter.ref=f'A1:{__import__("openpyxl").utils.get_column_letter(len(fields))}{len(rows)+1}'
  for i,k in enumerate(fields,1):s.column_dimensions[__import__('openpyxl').utils.get_column_letter(i)].width=min(55,max(18,len(k)+2))
  hs=[]
  for k in fields:
   c=WriteOnlyCell(s,k);c.font=Font(bold=True,color='FFFFFF');c.fill=PatternFill('solid',fgColor='225D65');hs.append(c)
  s.append(hs)
  for r in rows:
   vals=[]
   for k in fields:
    v=r.get(k)
    if isinstance(v,float) and not math.isfinite(v):v=None
    c=WriteOnlyCell(s,v)
    if isinstance(v,str):
     if len(v)>32767:raise ValueError('Texto demasiado largo; no se truncará: '+k)
     c.data_type='s'
    vals.append(c)
   s.append(vals)
 w.save(path)

def apply_geography(basegeos,approved,pending,raw,zonemap,comunas):
 byraw={r['id_encuesta']:r for r in raw};geos={(r['id_encuesta'],r['extremo']):dict(r) for r in basegeos};assert len(geos)==2*len(raw)
 change=[];allowed=set();finalcols=['comuna_homologada','cut_comuna','zona_modelacion_aristo','nombre_zona_aristo','localidad_homologada','estado','metodo','motivo','pais_homologado']
 for p in approved:
  key=p['id_encuesta'],p['extremo'];assert key not in allowed;allowed.add(key);old=geos[key];source=byraw[key[0]]
  assert source[key[1]+'_original']==p['literal_original']==old['literal_original']
  assert old['lat_original']==p['lat_original'] and old['lon_original']==p['lon_original']
  assert source['hash_registro']==p['hash_registro_original']
  cut=p['CUT_propuesto'];zone=int(p['zona_propuesta']);assert zonemap[cut]['zona']==zone
  if p['coordenada_individual_disponible']=='True':
   point=Point(float(p['lon_original']),float(p['lat_original']));hits=comunas.iloc[comunas.sindex.query(point,predicate='intersects')];assert cut in hits.CUT_COM.astype(str).tolist()
  elif p['nivel_evidencia'].startswith('B'):assert not p['lat_original'] and not p['lon_original']
  prior={k:old.get(k,'') for k in finalcols};old.update(comuna_homologada=zonemap[cut]['comuna'],cut_comuna=cut,zona_modelacion_aristo=str(zone),nombre_zona_aristo=zonemap[cut]['nombre_zona'],localidad_homologada=p.get('localidad_propuesta',''),estado='Homologado comuna y zona',metodo='Propuesta aprobada por usuario: '+p['nivel_evidencia'],motivo=p['justificacion'],pais_homologado='Chile',aprobacion_29=True,alerta_limite_aprobada=p.get('alerta_precision_limite',''),sin_coordenada_individual=p['coordenada_individual_disponible']!='True')
  change.append(dict(id_encuesta=key[0],extremo=key[1],nivel_evidencia=p['nivel_evidencia'],CUT_antes=prior['cut_comuna'],CUT_despues=cut,zona_antes=prior['zona_modelacion_aristo'],zona_despues=zone,lat_original=old['lat_original'],lon_original=old['lon_original'],alerta_limite=old.get('alerta_limite_aprobada',''),justificacion=p['justificacion'],fuentes=p['fuentes_evidencia']))
 assert len(change)==len(approved)
 for r in basegeos:
  key=r['id_encuesta'],r['extremo'];new=geos[key]
  assert all(new[k]==r[k] for k in r if k not in finalcols)
  if key not in allowed:assert new==r
 for r in pending:
  new=geos[r['id_encuesta'],r['extremo']];assert not new['cut_comuna'] and not new['zona_modelacion_aristo'] and new['estado']=='Pendiente'
 return geos,change

def route_decision(r):
 reasons=[]
 if not r['OD_homologado']:reasons.append('OD no determinado por homologación; no imputado')
 if r.get('separacion_control_red_m',float('inf'))>250:reasons.append('Control sin enlace cercano acreditado a red')
 if any(v is None or v>2500 for v in [r.get('separacion_origen_red_m'),r.get('separacion_destino_red_m')]):reasons.append('Localización de extremo sin conexión acreditada a red')
 if not r.get('control_en_camino_Chile'):reasons.append('Control fuera del camino terrestre de referencia')
 if not r.get('sentido_compatible'):reasons.append('Sentido incompatible con trayectoria de referencia')
 if not r.get('corredor_R7_y_R265'):reasons.append('No se acredita corredor completo pertinente R7-R265; no inclusión por sola Ruta7')
 if r.get('alternativa_barcaza_compatible'):reasons.append('Acceso lacustre puede explicar punto y sentido; ruta no identificada')
 if r.get('referencia_internacional'):reasons.append('Referencia internacional o mixta no delimitada')
 auxiliary=not reasons
 principal_reasons=list(reasons)
 if r.get('alternativa_Argentina_compatible'):principal_reasons.append('Itinerario con Argentina también compatible en red: no identifica recorrido íntegro por Chile')
 return not principal_reasons,auxiliary,'; '.join(principal_reasons) or 'Ruta Chile identificada con las pruebas disponibles','; '.join(reasons) or 'Interceptación y trayectoria chilena compatibles; sin certificar itinerario total'

def principal_sin_identificacion(rows):
 out=[]
 labels=['id_OD_zonal','origen','destino','zona_origen','zona_destino','comuna_origen','comuna_destino','CUT_origen','CUT_destino','campana','segmento','sentido_OD']
 for r in rows:
  v={k:r[k] for k in labels if k in r}
  v.update(entrevistas_aceptadas_recorrido_Chile_completo=0,factor_expansion_principal=None,pantalla_seleccionada_principal=None,flujo_terrestre_Chile_confirmado_12h=None,periodo_representado='Sin estimación del itinerario completo; ver fechas del auxiliar por separado',cobertura_espacial='No identificada para el recorrido íntegro por Chile',estado='No estimable con las fuentes disponibles; NO es demanda cero',unidad='Vehículos de clase; flujo no identificado',ruta='Chile completo sin imputación')
  out.append(v)
 return out

def main():
 p=argparse.ArgumentParser();p.add_argument('--anexos',type=Path,required=True);p.add_argument('--homologacion-previa',type=Path,required=True);p.add_argument('--decisiones',type=Path,required=True);p.add_argument('--salida',type=Path,required=True);args=p.parse_args();B=args.anexos.resolve();PRE=args.homologacion_previa.resolve();D=args.decisiones.resolve();O=args.salida.resolve()
 if (O/'validacion.json').exists():raise RuntimeError('Entrega anterior existente: elegir una carpeta nueva para conservarla.')
 O.mkdir(parents=True,exist_ok=True);(O/'csv').mkdir(exist_ok=True);(O/'ejecucion_config.json').write_text(json.dumps({k:str(v) for k,v in vars(args).items()},ensure_ascii=False,indent=2))
 protected=[]
 protected+=list(PRE.rglob('*'))+list(D.rglob('*'))
 protected+=[B/'eod aysen'/n for n in ['EOD Aysen Febrero 2026 vf.xlsx','EOD Aysen Abril 2026.xlsx','Mediciones Aysen (23-02-26).kmz']]
 for folder,stem in [(B/'Cap 7/SIG Red_Vial_Final','Red_Vial_Final'),(B/'Cap 7/SIG Red_Vial_Final','Nodos_RedVial'),(B/'Cap 7/SIG Red_Vial_Final','Nodos_Puntos_Pasada'),(B/'Cap 7/SIG Red_Vial_Final','Zonificación'),(B/'SIG/Coberturas/IDE','COMUNAS')]:
  protected+=list(folder.glob(stem+'.*'))
 protected.append(B/'Cap 8/BBDD_para_Pijas_actualizada_EOD_Final.xlsx')
 protected=list(dict.fromkeys(x for x in protected if x.is_file() and '__pycache__' not in x.parts))
 hashes=[dict(archivo=str(x),sha256=sha(x)) for x in protected];(O/'integridad_antes.json').write_text(json.dumps(hashes,ensure_ascii=False,indent=2))
 obs,flows,headers=read_originals(B);byid={r['id_encuesta']:r for r in obs};old=readcsv(D/'alcance_y_clasificacion_anteriores.csv');oldmap={r['id_encuesta']:r for r in old};assert set(byid)==set(oldmap)
 scope=[r['id_encuesta'] for r in old if r['estado_filtro']=='Incluido' and r['ruta_utilizada']=='Tierra confirmada'];assert len(scope)==112
 oldrecords={r['id_encuesta']:r for r in readcsv(PRE/'correspondencias_por_registro.csv')};assert all(byid[rid]['hash_registro']==r['hash_registro_original'] for rid,r in oldrecords.items())
 approved=readcsv(D/'homologaciones_aprobadas_29.csv');pending=readcsv(D/'pendientes_no_imputar_8.csv');assert len(approved)==29 and len(pending)==8
 admin=gp.read_file(B/'SIG/Coberturas/IDE/COMUNAS.shp').to_crs(4326);zones=gp.read_file(B/'Cap 7/SIG Red_Vial_Final/Zonificación.shp');zonemap={str(r.CUT_COM):dict(zona=int(r.Zona),comuna=r.COMUNA,nombre_zona=r.COMUNA) for _,r in zones.iterrows() if int(r.Zona)<=50}
 basegeos=readcsv(PRE/'extremos_homologados.csv');geos,changed=apply_geography(basegeos,approved,pending,obs,zonemap,admin)
 assert len(geos)==17656
 geolist=list(geos.values());writecsv(O/'csv/homologacion_cerrada_extremos.csv',geolist);writecsv(O/'csv/homologaciones_29_aplicadas.csv',changed)
 geo_records=[]
 for r in obs:
  g=dict(id_encuesta=r['id_encuesta'],archivo=r['archivo'],fila_excel=r['fila_excel'],hash_registro_original=r['hash_registro'],estado_area_previo=oldmap[r['id_encuesta']]['estado_filtro'],decision_area_previa=oldmap[r['id_encuesta']]['decision_usuario'],decision_lista_previa=oldmap[r['id_encuesta']]['decision_lista_corta'])
  for side in ['origen','destino']:
   e=geos[r['id_encuesta'],side]
   for k in ['literal_original','localidad_original','lat_original','lon_original','localidad_homologada','comuna_homologada','cut_comuna','zona_modelacion_aristo','estado','metodo']:g[side+'_'+k]=e[k]
   g[side+'_alerta_limite']=e.get('alerta_limite_aprobada','')
  geo_records.append(g)
 writecsv(O/'csv/homologacion_cerrada_registros.csv',geo_records)
 strata,weights,stratumidx=expansion(obs,flows)
 oldfactors=None
 print('Originales leídos;29 homologaciones aplicadas; estratos:',len(strata),'compatibles:',sum(s['factor_expansion'] is not None for s in strata),flush=True)
 points=kmz(B);network=Red(B,obs,geos,points,scope);audit=[]
 for rid in scope:
  source=byid[rid];prior=oldmap[rid];r=network.audit(source);a,b=geos[rid,'origen'],geos[rid,'destino'];ea,eb=network.endpoint[rid,'origen'],network.endpoint[rid,'destino']
  r.update(archivo=source['archivo'],fila_excel=source['fila_excel'],campana=source['campana'],fecha=source['fecha'],segmento=source['segmento'],hash_registro_original=source['hash_registro'],origen_original=source['origen_original'],destino_original=source['destino_original'],origen_detallado=ea['localidad'] or a['localidad_original'],destino_detallado=eb['localidad'] or b['localidad_original'],comuna_origen=a['comuna_homologada'],comuna_destino=b['comuna_homologada'],CUT_origen=a['cut_comuna'],CUT_destino=b['cut_comuna'],zona_origen=a['zona_modelacion_aristo'],zona_destino=b['zona_modelacion_aristo'],origen_lat_original=a['lat_original'],origen_lon_original=a['lon_original'],destino_lat_original=b['lat_original'],destino_lon_original=b['lon_original'],tipo_punto_origen=ea['tipo_punto'],tipo_punto_destino=eb['tipo_punto'],fuente_punto_origen=ea['fuente'],fuente_punto_destino=eb['fuente'],OD_homologado=a['estado']=='Homologado comuna y zona' and b['estado']=='Homologado comuna y zona',referencia_internacional=any(word in norm(source['origen_original']+' '+source['destino_original']) for word in ['ARGENTINA','LOS ANTIGUOS','PERITO MORENO','JEINIMENI','HUEMULES']),clasificacion_anterior=prior['ruta_utilizada'],seleccionada_pantalla_anterior=prior['seleccionada_nueva']=='True',pantalla_anterior=prior['pantalla_nueva'],decision_area_previa=prior['decision_usuario'],decision_lista_previa=prior['decision_lista_corta'],ruta_observada_directamente=False,tramo_de_interceptacion_observado=True,ruta_inferida_topologicamente=True,control_KMZ_lat=points[source['control']]['lat'],control_KMZ_lon=points[source['control']]['lon'],id_estrato=weights[rid]['id_estrato'],factor_expansion=weights[rid]['factor'],motivo_sin_factor=weights[rid]['motivo'])
  primary,aux,why,whyaux=route_decision(r);r.update(aceptada_principal=primary,aceptada_corredor_auxiliar=aux,motivo_principal=why,motivo_auxiliar=whyaux,clasificacion_nueva='Chile completo inferido' if primary else 'Corredor chileno compatible; itinerario país desconocido' if aux else 'Excluida de estimación: evidencia insuficiente/contradictoria',Argentina_observada=False)
  audit.append(r)
 print('Rutas revalidadas:',len(audit),'principal:',sum(r['aceptada_principal'] for r in audit),'corredor auxiliar:',sum(r['aceptada_corredor_auxiliar'] for r in audit),flush=True)
 detail,zonal,screen,relations,daily,selected=build(audit,obs,strata,weights,network)
 for r in audit:r['seleccionada_corredor_auxiliar']=selected.get(r['id_encuesta'],False)
 # Con identificación insuficiente no se fabrica un flujo principal: las celdas quedan explícitamente no estimables.
 principal_detail=principal_sin_identificacion(detail)
 principal_zonal=principal_sin_identificacion(zonal)
 if any(r['aceptada_principal'] for r in audit):
  principal_input=[dict(a,aceptada_corredor_auxiliar=a['aceptada_principal']) for a in audit]
  pd,pz,ps,pr,pp,pselected=build(principal_input,obs,strata,weights,network)
  principal_detail=[];principal_zonal=[]
  for rows,out in [(pd,principal_detail),(pz,principal_zonal)]:
   for r in rows:
    v=dict(r);value=v.pop('flujo_corredor_auxiliar_12h');v.update(flujo_terrestre_Chile_confirmado_12h=value,estado='Expandido condicional de recorridos Chile identificados' if value is not None else 'Sin expansión compatible; no es cero',ruta='Chile completo inferido topológicamente')
    out.append(v)
 summary=[]
 for camp in ['Verano','Otoño']:
  for seg in ['VL','CAM']:
   aa=[a for a in audit if (a['campana'],a['segmento'])==(camp,seg)];vv=[a for a in aa if a['aceptada_corredor_auxiliar']];sel=[a for a in vv if a['seleccionada_corredor_auxiliar']];zz=[z for z in zonal if (z['campana'],z['segmento'])==(camp,seg)]
   summary.append(dict(campana=camp,segmento=seg,entrevistas_cohorte=len(aa),OD_homologado=sum(a['OD_homologado'] for a in aa),aceptadas_principal=sum(a['aceptada_principal'] for a in aa),aceptadas_corredor=len(vv),corredor_con_factor=sum(a['factor_expansion'] is not None for a in vv),seleccionadas_corredor=len(sel),seleccionadas_con_factor=sum(a['factor_expansion'] is not None for a in sel),OD_zonales_contraste=len(zz),OD_zonales_con_flujo_auxiliar=sum(z['flujo_corredor_auxiliar_12h'] is not None for z in zz),OD_zonales_sin_flujo=sum(z['flujo_corredor_auxiliar_12h'] is None for z in zz),flujo_principal_Chile=None,flujos_auxiliares='Ver celdas por pantalla y fechas: no se suman como total del sistema',unidad='Entrevistas, vehículos por clase/12h; sin personas,24h ni anualización'))
 data={'auditoria_112':audit,'matriz_principal_detallada':principal_detail,'matriz_principal_zonas_ARISTO':principal_zonal,'matriz_auxiliar_corredor_detallada':detail,'matriz_auxiliar_corredor_zonas':zonal,'seleccion_pantallas':screen,'relaciones_entre_controles':relations,'estimaciones_por_dia_y_pantalla':daily,'estratos_y_factores':strata,'testigos_itinerarios_Argentina_posibles':network.argument_paths,'resumen_cobertura':summary,'auxiliar_excluidos': [a for a in audit if not a['aceptada_principal']]}
 for name,rows in data.items():writecsv(O/'csv'/f'{name}.csv',rows)
 source_scope=[dict(id_encuesta=r['id_encuesta'],archivo=r['archivo'],fila_excel=r['fila_excel'],campana=r['campana'],fecha=r['fecha'],control=r['control'],sentido=r['sentido'],segmento=r['segmento'],minuto=r['minuto'],hash_registro=r['hash_registro'],id_estrato=weights[r['id_encuesta']]['id_estrato'],factor=weights[r['id_encuesta']]['factor'],motivo_sin_factor=weights[r['id_encuesta']]['motivo']) for r in obs]
 writecsv(O/'csv/universo_original_denominadores.csv',source_scope)
 checks={'8828_registros_originales_unicos':len(byid)==8828,'17656_extremos_conservados':len(geos)==17656,'29_aprobaciones_aplicadas_solo_en_sus_claves':len(changed)==29,'8_pendientes_sin_imputar':all(geos[r['id_encuesta'],r['extremo']]['estado']=='Pendiente' for r in pending),'112_auditadas_exactamente':len(audit)==112 and {a['id_encuesta'] for a in audit}==set(scope),'22_seleccionadas_antes_reproducidas':sum(a['seleccionada_pantalla_anterior'] for a in audit)==22,'denominadores_antes_de_todo_filtro':all(s['N_entrevistas_originales_ventana']==sum(r['minuto'] is not None and 480<=r['minuto']<1200 for r in obs if tuple(r[k] for k in ['campana','fecha','control','sentido','segmento'])==tuple(s[k] for k in ['campana','fecha','control','sentido','segmento'])) for s in strata),'balance_F_N_por_estrato':all(s['factor_expansion'] is None or math.isclose(s['N_entrevistas_originales_ventana']*s['factor_expansion'],s['aforo_vehiculos_12h'],rel_tol=1e-12) for s in strata),'ningun_peso_de_estrato_inconsistente':all(s['factor_expansion'] is None for s in strata if s['motivo']),'una_pantalla_por_OD_zonal_y_campana_clase':len([s for s in screen if s['seleccionada']])==len({(s['zona_origen'],s['zona_destino'],s['campana'],s['segmento']) for s in screen}),'no_suma_pantallas_alternativas':all(not r['se_suman'] for r in relations),'balance_detallado_zonal_auxiliar':all((z['flujo_corredor_auxiliar_12h'] is None and not any(r['flujo_corredor_auxiliar_12h'] is not None for r in detail if r['id_OD_zonal']==z['id_OD_zonal'])) or (z['flujo_corredor_auxiliar_12h'] is not None and math.isclose(z['flujo_corredor_auxiliar_12h'],sum(r['flujo_corredor_auxiliar_12h'] for r in detail if r['id_OD_zonal']==z['id_OD_zonal'] and r['flujo_corredor_auxiliar_12h'] is not None),rel_tol=1e-12)) for z in zonal),'principal_no_se_rellena_si_no_hay_identificacion':any(a['aceptada_principal'] for a in audit) or all(r['flujo_terrestre_Chile_confirmado_12h'] is None for r in principal_detail+principal_zonal),'testigos_Argentina_sin_repetir_nodos':all(r['prueba_no_repeticion'] for r in network.argument_paths),'sin_mezclar_personas_o_pesos_24h':all(r['segmento'] in ['VL','CAM'] for r in audit)}
 notes=[dict(tema='Resultado principal',detalle='No se identifica itinerario íntegro por Chile: matriz principal con flujos faltantes, nunca cero. Interceptación en Chile no descarta etapas argentinas.'),dict(tema='Auxiliar',detalle='Se estima únicamente fracción de entrevistas compatibles con corredor chileno observado; NO sustituye el bloque principal ni es demanda observada de Argentina. Ver fechas, clase y pantalla de cada celda.'),dict(tema='Expansión',detalle='F/N por misma fecha, sección EOD, sentido, clase y08–20, con todas las entrevistas originales del estrato. Duplicados no resueltos, horas imposibles o aforos incompatibles quedan sin peso. No usa aforos urbanos ni24h.'),dict(tema='Pantallas',detalle='Una por zonaOD/campaña/clase para que las localidades sumen únicamente dentro del mismo cordón/período. Se conserva cada alternativa y se identifican subOD sin muestra en pantalla elegida, sin rellenarlos con cero.'),dict(tema='Geografía',detalle='29 propuestas aprobadas aplicadas como capa adicional. Localidades/JSON/coordenadas originales y decisiones anteriores intactos;8 pendientes sin imputar. Puerto Río Tranquilo conserva alerta del límite.'),dict(tema='Argentina',detalle='Testigo de posibilidad geométrica, no ruta usada ni certificación de habilitación de pasos en cada fecha. No se descarta por ser largo ni se elige ruta por costo/distancia. La falta de itinerario impide confirmar el país integralmente.'),dict(tema='Exposición',detalle='Vehículos por ventana08–20 de las fechas válidas; promedio de jornadas de la misma pantalla. No24h, temporada completa ni anual. Cero muestral del estrato no es demanda inexistente.'),dict(tema='Privacidad',detalle='Excel/HTML/CSV detallados locales reservados. Publicación únicamente de código, método y resumen agregado revisado.')]
 workbook(O/'01_homologacion_cerrada.xlsx',[('Leer primero',notes,None),('29 cambios aplicados',changed,None),('Registros',geo_records,None),('Extremos',geolist,None),('8 pendientes conservados',pending,None)])
 workbook(O/'02_matrices_OD.xlsx',[('Leer primero',notes,None),('Principal detallada',principal_detail,None),('Principal zonas ARISTO',principal_zonal,None),('Aux corredor detallada',detail,None),('Aux corredor zonas',zonal,None),('Cobertura',summary,None)])
 workbook(O/'03_auditoria_112.xlsx',[('Leer primero',notes,None),('112 entrevistas',audit,None),('Excluidos principal',data['auxiliar_excluidos'],None),('Testigos Argentina',network.argument_paths,None)])
 workbook(O/'04_pantallas_y_expansion.xlsx',[('Leer primero',notes,None),('Selección pantallas',screen,None),('Relaciones serie ramas',relations,None),('Alternativas por jornada',daily,None),('Estratos y factores',strata,None),('Universo denominadores',source_scope,None),('Validaciones',[dict(comprobacion=k,correcto=v) for k,v in checks.items()],None),('Fuentes y hashes',hashes,None)])
 # Verificar lectura de resultados y correspondencia sin depender de una importación de modelos.
 w=load_workbook(O/'03_auditoria_112.xlsx',read_only=True,data_only=True);it=w['112 entrevistas'].values;hh=next(it);reread=[dict(zip(hh,v)) for v in it];w.close();checks['Excel112_unicos_y_decisiones_conservadas']=len(reread)==112 and len({r['id_encuesta'] for r in reread})==112 and all(r['decision_area_previa']==oldmap[r['id_encuesta']]['decision_usuario'] for r in reread if oldmap[r['id_encuesta']]['decision_usuario'])
 checks['fuentes_y_entregas_anteriores_intactas']=all(sha(r['archivo'])==r['sha256'] for r in hashes)
 assert all(checks.values()),checks
 result=dict(estado='correcto',verificaciones=checks,cohorte=112,homologaciones_aplicadas=29,pendientes_no_imputados=8,aceptadas_principal=sum(a['aceptada_principal'] for a in audit),aceptadas_corredor_auxiliar=sum(a['aceptada_corredor_auxiliar'] for a in audit),corredor_con_factor=sum(a['aceptada_corredor_auxiliar'] and a['factor_expansion'] is not None for a in audit),seleccionadas_corredor=sum(a['seleccionada_corredor_auxiliar'] for a in audit),seleccionadas_con_factor=sum(a['seleccionada_corredor_auxiliar'] and a['factor_expansion'] is not None for a in audit),celdas_auxiliares_con_flujo=sum(z['flujo_corredor_auxiliar_12h'] is not None for z in zonal),celdas_principales_con_flujo=sum(r['flujo_terrestre_Chile_confirmado_12h'] is not None for r in principal_zonal),estratos=len(strata),estratos_compatibles=sum(s['factor_expansion'] is not None for s in strata),archivos_protegidos=len(hashes),resultados_xlsx=[dict(archivo=x.name,sha256=sha(x)) for x in O.glob('*.xlsx')],fuentes_code=[dict(archivo=x.name,sha256=sha(x)) for x in Path(__file__).parent.glob('*.py')])
 (O/'validacion.json').write_text(json.dumps(result,ensure_ascii=False,indent=2));print(json.dumps(result,ensure_ascii=False,indent=2),flush=True)
 from informe import generar
 generar(O,notes,summary,detail,zonal,screen,audit,changed,result)

if __name__=='__main__':main()
