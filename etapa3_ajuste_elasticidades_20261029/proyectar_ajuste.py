"""Parametrización final de Etapa3: PIBcentral común y elasticidades del usuario.
Reutiliza escritor del script previo; no modifica etapas previas ni lee encuestas.
Las columnas anuales evitan repetir 25 veces los metadatos de cada OD.
"""
from pathlib import Path
import argparse,hashlib,importlib.util,json,time
import pandas as pd
from openpyxl import load_workbook

def sha(p):return hashlib.sha256(Path(p).read_bytes()).hexdigest()

def widen(base,primary,partial,cfg,ratios):
    result=[]
    for scenario in ['conservador','central','alto']:
        x=base.copy().rename(columns={'escenario':'escenario_base_Etapa2','anio':'anio_base_original'})
        x['id_fila_base']=range(len(base));x['escenario_crecimiento']=scenario
        x['familia_elasticidad']=base.segmento.map(cfg['familia_por_segmento'])
        unknown=x.familia_elasticidad.isna()
        assert base.loc[unknown,[primary,partial]].fillna(0).eq(0).all().all()
        x['elasticidad_PIB']=x.familia_elasticidad.map(lambda f:cfg['elasticidades'][f][scenario] if pd.notna(f) else 0.)
        x['regla_parametro']=x.familia_elasticidad.notna().map({True:'Supuesto transferido, no estimación local',False:'Sin parámetro: se conservan únicamente0oNULL de la base'})
        x['metrica_Q']=primary;x['anio_base']=2026
        for year in cfg['anios']:
            factor=ratios.loc[year]**x.elasticidad_PIB
            x[f'Q_{year}']=base[primary]*factor
            x[f'aporte_parcial_{year}']=base[partial]*factor
            assert x[f'Q_{year}'].isna().equals(base[primary].isna())
        result.append(x)
    return pd.concat(result,ignore_index=True)

def main():
    started=time.monotonic();ap=argparse.ArgumentParser();ap.add_argument('--anexos',type=Path,default=Path('.'));ap.add_argument('--config',type=Path,default=Path(__file__).with_name('configuracion.json'));ap.add_argument('--salida',type=Path,required=True);args=ap.parse_args()
    if args.salida.exists():raise FileExistsError('Conservar históricos: indicar carpeta nueva')
    root=args.anexos.resolve();cfg=json.loads(args.config.read_text());prior=root/'etapa3_proyecciones_20261029'
    spec=importlib.util.spec_from_file_location('proyeccion_previa',prior/'proyectar.py');old=importlib.util.module_from_spec(spec);spec.loader.exec_module(old)
    paths={n:root/'etapa2_demanda_anual_20261029/resultados/entrega_final/csv'/f'{n}.csv' for n in ['base_unificada_anual','aristo_anual_detalle','citra_anual_localidades','aristo_resumen_escenarios']}
    gdp_path=prior/'parametros/PIB_anual_2026_2050.csv';oldbook=prior/'resultados/entrega_final/proyecciones_2030_2040_2050.xlsx'
    protected=list(paths.values())+list((prior/'resultados/entrega_final').glob('*'))+[prior/'proyectar.py',gdp_path,args.config,Path(__file__)]
    hashes={str(p):sha(p) for p in protected if p.is_file()}
    data={k:pd.read_csv(p) for k,p in paths.items()}
    bases={k:d[d.escenario.str.lower().eq('central')].copy().reset_index(drop=True) for k,d in data.items()}
    gdp=pd.read_csv(gdp_path);ratios=gdp.set_index('anio').PIB_base/gdp.set_index('anio').PIB_base.loc[2026]
    assert ratios.loc[2026]==1. and set(cfg['anios'])<=set(ratios.index)
    b=bases['base_unificada_anual'];b['Q_referencia_2026']=b.demanda_anual_operacional.where(b.fuente.eq('ARISTO'),b.demanda_anual_ventana12h)
    z=widen(b,'Q_referencia_2026','aporte_parcial',cfg,ratios)
    ld=widen(bases['aristo_anual_detalle'],'demanda_anual_operacional_componente','aporte_anual_en_meses_con_perfil',cfg,ratios)
    cd=widen(bases['citra_anual_localidades'],'demanda_anual_ventana12h','aporte_parcial_meses_con_perfil_12h',cfg,ratios)
    factors=[]
    for family,ss in cfg['elasticidades'].items():
        for scenario,beta in ss.items():
            for year in cfg['anios']:
                factors.append(dict(familia=family,escenario=scenario,anio=year,PIB_relativo_2026=ratios.loc[year],elasticidad=beta,factor=ratios.loc[year]**beta,atribucion='Supuesto seleccionado por usuario, no parámetro estimado para el lago'))
    factors=pd.DataFrame(factors)
    totals=[];summarybase=bases['aristo_resumen_escenarios']
    for row in summarybase.to_dict('records'):
        family=cfg['familia_por_segmento'].get(row['segmento'])
        for scenario in ['conservador','central','alto']:
            beta=cfg['elasticidades'][family][scenario] if family else 0.
            for year in cfg['anios']:
                factor=ratios.loc[year]**beta
                totals.append({k:row[k] for k in ['version','segmento']}|dict(escenario_crecimiento=scenario,anio=year,elasticidad=beta,factor=factor,
                    **{c:row[c]*factor for c in ['flujo_fisico_anual','observado_OD_anualizado','imputado_OD_anualizado','residual_anualizado']}))
    totals=pd.DataFrame(totals)
    highlights=totals[totals.version.eq('B')&totals.segmento.isin(['PRIVADOS_COMPATIBLES','CAM','PEATON'])&totals.anio.isin(cfg['hitos'])].pivot(index=['anio','escenario_crecimiento'],columns='segmento',values='flujo_fisico_anual').reset_index()
    with pd.ExcelFile(oldbook) as book:
        previous=pd.read_excel(book,sheet_name='Resumen lacustre');previous_factors=pd.read_excel(book,sheet_name='Factores')
    previous['naturaleza']='Proyección anterior preservada; betaARISTOalta y trayectoriasPIB distintas entre escenarios; contraste independiente'
    newcentral=highlights[highlights.escenario_crecimiento.eq('central')]
    priorcentral=previous[previous.escenario_crecimiento.eq('central')]
    comparison=newcentral.merge(priorcentral,on='escenario_crecimiento',how='inner',suffixes=('_nuevo','_ARISTO'))
    comparison=comparison[comparison.anio.eq(comparison.anio_proyeccion)].drop(columns=['anio_proyeccion'])
    history=pd.DataFrame(cfg['contraste_historico'])
    checkbase=all(d.Q_2026.equals(d[basefield]) for d,basefield in [(z,'Q_referencia_2026'),(ld,'demanda_anual_operacional_componente'),(cd,'demanda_anual_ventana12h')])
    equation=all(abs(r.factor-r.PIB_relativo_2026**r.elasticidad)<1e-12 for r in factors.itertuples())
    residual_ids=set(b.index[b.componente.eq('residual')]);external_ids=set(b.index[b.subtipo.eq('residual_externo')])
    unitcheck=all(d.loc[d.escenario_crecimiento.eq('central'),'unidad'].reset_index(drop=True).equals(source.unidad) for d,source in [(z,b),(ld,bases['aristo_anual_detalle'])])
    # El soporte y las cantidades base son idénticos; los escenarios solo cambian beta.
    checks={
      'misma_base2026_todos_escenarios':checkbase,
      'sin_doble_crecimiento':factors[factors.anio.eq(2026)].factor.eq(1).all(),
      'crecimiento_consistente_con_beta_y_PIBcentral':equation and factors.groupby('anio').PIB_relativo_2026.nunique().eq(1).all(),
      'residuales_y_externos_conservados':set(z.loc[z.componente.eq('residual'),'id_fila_base'])==residual_ids and set(z.loc[z.subtipo.eq('residual_externo'),'id_fila_base'])==external_ids,
      'faltantes_conservados':all(z[f'Q_{y}'].isna().equals(z.Q_referencia_2026.isna()) for y in cfg['anios']),
      'sin_mercado_CITRA_nuevo':len(z[z.fuente.eq('CITRA')])==3*len(b[b.fuente.eq('CITRA')]),
      'unidades_compatibles_sin_VP_a_VL':unitcheck,
      'proyecciones_y_fuentes_anteriores_intactas':all(sha(p)==h for p,h in hashes.items()),
    };checks={k:bool(v) for k,v in checks.items()};assert all(checks.values()),checks
    params=pd.DataFrame([dict(familia=f,**betas,origen='Selección del usuario; supuesto transferido, no estimación local') for f,betas in cfg['elasticidades'].items()])
    notes=pd.DataFrame([
      dict(tema='Base',detalle='Etapa2central2026, versionesARISTOA/B intactas; no sumar versiones, vistas privadas o bloques.'),
      dict(tema='Escenarios',detalle='PIBcentralARISTOúnico. Solo varía beta. PIBoptimista y trayectoria anterior son contrastes, no inputs de los casos principales.'),
      dict(tema='Transferencia',detalle='VL,VP,motos,bicicletas y privados compatibles comparten beta de privados sin cambiar su unidad. Peatones usan su beta propia y no son pasajeros totales.'),
      dict(tema='Faltantes',detalle='CITRAparcial12h,24hNULL; clases/ODsinanual siguenNULL. Aporteparcial se proyecta aparte, no se transforma en demanda anual completa.'),
      dict(tema='Bibliografía',detalle='RAND2014,ITF/OECD2021 y estudios del lago/compraembarcación orientan método. No se afirma que publiquen exactamente las betas elegidas.'),
      dict(tema='Históricos',detalle=cfg['limite_contraste']),
    ])
    args.salida.mkdir(parents=True);xlsx=args.salida/'proyecciones_ajustadas_2026_2050.xlsx'
    sheets={'Leer primero':notes,'Elasticidades':params,'Factores anuales':factors,'PIB central y contraste':gdp,'Proyeccion zonas anual':z,'Lacustre detalle anual':ld,'CITRA localidades anual':cd,'Resumen lacustre anual':totals,'Hitos2030_2040_2050':highlights,'ARISTO anterior contraste':previous,'FactoresARISTO anteriores':previous_factors,'Central nuevo vs ARISTO':comparison,'Tasas historicas contraste':history,'Validaciones':pd.DataFrame([dict(control=k,correcto=v) for k,v in checks.items()]),'Trazabilidad':pd.DataFrame([dict(archivo=p,sha256=h) for p,h in hashes.items()])}
    old.write_excel(xlsx,sheets)
    w=load_workbook(xlsx,read_only=True,data_only=True);assert w.sheetnames==list(sheets);w.close()
    duration=round(time.monotonic()-started,2)
    css='body{font:16px/1.55 system-ui,sans-serif;max-width:1100px;margin:auto;padding:24px;color:#203e44}.scroll{overflow:auto}table{border-collapse:collapse;width:100%;font-size:13px;margin:16px 0}td,th{border:1px solid #cad9d5;padding:8px;text-align:left}th{background:#225d65;color:white}a{color:#096e8e}.box{padding:16px;background:#eaf3ef}pre{white-space:pre-wrap;background:#eef3f0;padding:12px}'
    text='<!doctype html><html lang="es"><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><title>Etapa3 — escenarios propios finales</title><style>'+css+'</style><body><h1>Etapa3: escenarios propios de crecimiento</h1>'
    text+='<p class="box">Proyecciones anuales2026–2050 desde la misma base central2026deEtapa2, usando el PIBcentralARISTOen todos los casos. Las diferencias principales se deben exclusivamente a elasticidades. <a href="'+xlsx.name+'">Descargar Excel completo</a>.</p>'
    text+='<h2>Parámetros elegidos</h2>'+old.table(params)
    text+='<p><code>Q_t=Q_2026×(PIBcentral_t/PIBcentral_2026)^β</code>. Son supuestos transferidos seleccionados por el usuario, no elasticidades estimadas para LagoGeneralCarrera. Las referencias metodológicas son RAND Europe(2014), <em>Road Traffic Demand Elasticities: Rapid Evidence Assessment</em>; ITF/OECD(2021), <em>ITF Transport Outlook2021</em>; el estudio de conectividad lacustre del lago y la metodología de compra de embarcaciones en zonas aisladas. <strong>No se afirma que esas publicaciones contengan exactamente los valores adoptados.</strong></p>'
    text+='<p>La base normalizada2026ya incluye la anualización aprobada: factor2026=1. No se reutiliza la demanda proyectada anterior como nueva base. Motos,bicicletas,VLyVPusan el ritmo de privados manteniendo sus categorías; VPno se convierte enVLestricto. Los otros pasajeros sin cantidad conocida siguenNULL; buses con base0no crean pasajeros.</p>'
    text+='<h2>Resultados lacustres: versiónB, ambos sentidos</h2><p>Unidades distintas por columna: privados compatibles(VL+moto+bicicleta/VP), vehículos de carga y personas peatones. No sumarlas como total de personas ni de vehículos. VersionesA/B y residuales están conservados en las matrices anuales.</p>'+old.table(highlights)
    text+='<h2>Comparación con las elasticidades altas de ARISTO</h2><p>Comparación central contra central: mismoPIBbase y misma demanda2026, aislando el cambio de beta. Los tres escenarios originales completos permanecen como contraste independiente enExcel; sus PIBalternativos no se combinan con los nuevos casos principales.</p>'+old.table(comparison[['anio','PRIVADOS_COMPATIBLES_nuevo','PRIVADOS_COMPATIBLES_ARISTO','CAM_nuevo','CAM_ARISTO','PEATON_nuevo','PEATON_ARISTO']])
    text+='<h2>Contraste histórico independiente</h2><p>Se recuperaron30tasas2026–2030de pasajeros sin auto del estudio <em>Mejoramiento de la Conectividad Lacustre en el LagoGeneralCarrera</em>, tablas53–54,p58. Son pronósticos históricos variables, no mediciones actuales. La numeración exacta del encabezadoAnexo22no se comprobó en la página consultada. No se extrapolan esas tasas a2050 ni se sustituyen las elasticidades elegidas.</p>'
    histshown=history[history.anio.isin([2026,2030])][['anio','periodo','escenario','tasa_porcentaje','segmento_original','tabla','pagina']]
    text+=old.table(histshown)
    text+='<p>No se recuperaron tasas numéricas verificadas de VL/CAMhistóricos en esta revisión acotada; el documento de compra de embarcación queda como antecedente metodológico, sin atribuirle cifras. La comparación histórica es parcial y no supone equivalencia exacta entre pasajeros sin auto y todos los pasajeros del ferry.</p>'
    text+='<h2>Cobertura conservada y reproducción</h2><p>CITRAcontinúa como mercado parcial de ventana12h; sus faltantes y24hsin estimación permanecenNULLtodos los años. El crecimiento no agrega OD ni mercado terrestre no observado. Los ODexternos y residualesARISTOsiguen identificados. Los aportes parciales aparecen separados deQanual. La composiciónODse conserva: no se calcula capacidad, captación, inducción o cambio de ruta.</p>'
    text+='<pre>python etapa3_ajuste_elasticidades_20261029/proyectar_ajuste.py --anexos . --config etapa3_ajuste_elasticidades_20261029/configuracion.json --salida /ruta/ENTREGA_NUEVA</pre>'
    text+=f'<p>{len(checks)}controles correctos; archivos anteriores intactos. Cálculo yExcelvalidados en{duration}s. Las columnas anuales evitan repetir los metadatosOD25veces. Solo se entregan este informe,unExcel yel código/configuración reproducibles.</p></body></html>'
    (args.salida/'informe_ajuste.html').write_text(text,encoding='utf-8')
    print(json.dumps({'estado':'correcto','segundos':duration,'validaciones':checks,'filas_zonas':len(z),'columnas_zonas':len(z.columns),'excel':str(xlsx),'html':str(args.salida/'informe_ajuste.html')},ensure_ascii=False,indent=2))

if __name__=='__main__':main()
