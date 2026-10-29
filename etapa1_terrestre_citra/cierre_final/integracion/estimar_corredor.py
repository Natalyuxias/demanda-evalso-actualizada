"""Cierre empírico de CITRA desde caché auditada, sin reconstruir red ni itinerarios.

Los aforos, N y geometrías provienen de la entrega previa validada. Se exige
equivalencia con sus Excel firmados por hash antes de reutilizar los CSV.
"""
from pathlib import Path
import argparse
import csv
import hashlib
import html
import json
import math
import time

import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Font, PatternFill


def sha(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def records(frame):
    return frame.astype(object).where(pd.notna(frame), None).to_dict('records')


def equivalent(a, b):
    if a is None or a == '':
        return b is None or b == ''
    if isinstance(b, bool):
        return str(a) == str(b)
    if isinstance(b, (int, float)):
        try:
            return math.isclose(float(a), b, rel_tol=1e-12, abs_tol=1e-10)
        except (ValueError, TypeError):
            return False
    return str(a) == str(b)


def validate_cache(cache):
    """Verificación acotada de artefactos reutilizados, sin recalcular sus métodos."""
    v = json.loads((cache / 'validacion.json').read_text())
    assert all(v['verificaciones'].values())
    for r in v['resultados_xlsx']:
        assert sha(cache / r['archivo']) == r['sha256'], r['archivo']
    mappings = {
        '03_auditoria_112.xlsx': [('112 entrevistas', 'auditoria_112.csv')],
        '04_pantallas_y_expansion.xlsx': [
            ('Selección pantallas', 'seleccion_pantallas.csv'),
            ('Relaciones serie ramas', 'relaciones_entre_controles.csv'),
            ('Alternativas por jornada', 'estimaciones_por_dia_y_pantalla.csv'),
            ('Estratos y factores', 'estratos_y_factores.csv')],
    }
    for filename, tables in mappings.items():
        w = load_workbook(cache / filename, read_only=True, data_only=True)
        for sheet, name in tables:
            it = w[sheet].iter_rows(values_only=True)
            headers = next(it)
            xx = [dict(zip(headers, row)) for row in it]
            with (cache / 'csv' / name).open(encoding='utf-8-sig', newline='') as f:
                cc = list(csv.DictReader(f))
            assert len(cc) == len(xx), name
            assert all(all(equivalent(row[k], x[k]) for k in row) for row, x in zip(cc, xx)), name
        w.close()
    # La caché geográfica y sus originales siguen siendo los de la entrega aprobada.
    protected = json.loads((cache / 'integridad_antes.json').read_text())
    assert all(sha(r['archivo']) == r['sha256'] for r in protected)
    return protected


def estimate_cell(rows, valid_strata):
    """Media de ventanas medidas de un SOLO cordón; n/N siempre del mismo día.

    Sin entrevista expandible en la celda, el flujo queda faltante. En presencia
    de muestra compatible se incluyen las jornadas con cero muestral y se
    descompone A+B usando exactamente el mismo denominador temporal.
    """
    n_days = len(valid_strata)
    known = rows[rows.factor_expansion.notna() & rows.id_estrato.isin(valid_strata.id_estrato)]
    result = {'jornadas_validas': n_days, 'entrevistas_pantalla': len(rows),
              'entrevistas_con_factor': len(known),
              'flujo_corredor_12h': None, 'flujo_A_12h': None, 'flujo_B_12h': None}
    for category in ('A', 'B'):
        result['entrevistas_' + category] = int(rows.categoria.eq(category).sum())
        result['entrevistas_' + category + '_con_factor'] = int(known.categoria.eq(category).sum())
    if n_days and len(known):
        result['flujo_corredor_12h'] = known.factor_expansion.sum() / n_days
        for category in ('A', 'B'):
            result['flujo_' + category + '_12h'] = known.loc[known.categoria.eq(category), 'factor_expansion'].sum() / n_days
    return result


def build_matrices(audit, screens, strata):
    valid = audit[audit.categoria.isin(['A', 'B'])].copy()
    detailed, zonal, daily = [], [], []
    chosen = screens[screens.seleccionada].set_index('id_OD_zonal', drop=False)
    for gid, group in valid.groupby('id_OD_zonal', sort=True):
        s = chosen.loc[gid]
        selected = group[group.control.eq(s.control) & group.sentido_observado.eq(s.sentido_control)]
        vs = strata[strata.campana.eq(s.campana) & strata.control.eq(s.control)
                    & strata.sentido.eq(s.sentido_control) & strata.segmento.eq(s.segmento)
                    & strata.factor_expansion.notna()]
        common = {k: s[k] for k in ['id_OD_zonal', 'zona_origen', 'zona_destino', 'campana', 'segmento', 'control', 'sentido_control']}
        common.update(fechas_representadas='|'.join(sorted(vs.fecha)), periodo='08:00–20:00 de fechas listadas; media de jornadas válidas',
                      unidad='vehículos de la clase / ventana de 12 horas',
                      cobertura='Parcial: fracción identificada en cohorte de 112; no total de comuna, corredor o proyecto',
                      ruta='Corredor observado; atribución a Chile inferida, no itinerario completo declarado',
                      incertidumbre='Tiempos, pasos operativos e itinerario individual no observados; representatividad muestral no certificada')
        common.update(comuna_origen=group.iloc[0].comuna_origen, comuna_destino=group.iloc[0].comuna_destino)
        for (origin, destination), fine in group.groupby(['origen_detallado', 'destino_detallado'], sort=True):
            sample = selected[selected.id_encuesta.isin(fine.id_encuesta)]
            r = common | {'origen': origin, 'destino': destination,
                          'CUT_origen': fine.iloc[0].CUT_origen, 'CUT_destino': fine.iloc[0].CUT_destino,
                          'sentido_OD': origin + ' → ' + destination,
                          'entrevistas_todos_controles': len(fine)} | estimate_cell(sample, vs)
            r['estado'] = ('Estimación parcial proporcional; cero de categoría es muestral, no ausencia poblacional'
                           if r['flujo_corredor_12h'] is not None else 'Sin muestra expandible en pantalla elegida; faltante, no cero')
            detailed.append(r)
        z = common | {'sentido_OD': f'{s.zona_origen} → {s.zona_destino}',
                      'entrevistas_todos_controles': len(group), 'OD_detallados': group.groupby(['origen_detallado', 'destino_detallado']).ngroups} | estimate_cell(selected, vs)
        z['OD_detallados_sin_flujo'] = sum(r['flujo_corredor_12h'] is None for r in detailed if r['id_OD_zonal'] == gid)
        z['estado'] = 'Suma únicamente de subOD identificados en la misma pantalla y fechas; otros subOD siguen faltantes'
        zonal.append(z)
    # Todas las pantallas alternativas y jornadas, con N y F originales intactos.
    for _, s in screens.iterrows():
        group = valid[valid.id_OD_zonal.eq(s.id_OD_zonal) & valid.control.eq(s.control)
                      & valid.sentido_observado.eq(s.sentido_control)]
        vs = strata[strata.campana.eq(s.campana) & strata.control.eq(s.control)
                    & strata.sentido.eq(s.sentido_control) & strata.segmento.eq(s.segmento)
                    & strata.factor_expansion.notna()]
        for _, st in vs.iterrows():
            sample = group[group.id_estrato.eq(st.id_estrato) & group.factor_expansion.notna()]
            r = {k: s[k] for k in ['id_OD_zonal', 'zona_origen', 'zona_destino', 'campana', 'segmento', 'control', 'sentido_control', 'seleccionada']}
            r.update(fecha=st.fecha, id_estrato=st.id_estrato, F_vehiculos_12h=st.aforo_vehiculos_12h,
                     N_original=st.N_entrevistas_originales_ventana, factor=st.factor_expansion,
                     n_corredor=len(sample), n_A=int(sample.categoria.eq('A').sum()), n_B=int(sample.categoria.eq('B').sum()))
            for category in ('corredor', 'A', 'B'):
                r['Q_' + category + '_12h'] = r['F_vehiculos_12h'] * r['n_' + category] / r['N_original']
            r.update(estado='Cero muestral del estimador proporcional; no demanda ausente' if not len(sample) else 'Expandido proporcional en ventana observada',
                     formula='Q=F*n/N; N original anterior a todo filtro', unidad='vehículos/12h del día indicado')
            daily.append(r)
    return pd.DataFrame(detailed), pd.DataFrame(zonal), pd.DataFrame(daily)


def export_excel(path, sheets):
    w = Workbook(write_only=True)
    for name, frame in sheets.items():
        s = w.create_sheet(name)
        s.freeze_panes = 'B2'
        s.append(list(frame.columns))
        for row in records(frame):
            cells = []
            for v in row.values():
                c = WriteOnlyCell(s, v)
                if isinstance(v, str):
                    assert len(v) <= 32767
                    c.data_type = 's'  # conservar textos originales, nunca fórmulas Excel
                cells.append(c)
            s.append(cells)
        s.auto_filter.ref = f'A1:{__import__("openpyxl").utils.get_column_letter(len(frame.columns))}{len(frame)+1}'
    w.save(path)


def main():
    started = time.monotonic()
    p = argparse.ArgumentParser()
    p.add_argument('--cache', type=Path, required=True)
    p.add_argument('--revision', type=Path, required=True)
    p.add_argument('--salida', type=Path, required=True)
    args = p.parse_args()
    cache, review, out = args.cache, args.revision, args.salida
    if out.exists():
        raise FileExistsError('Conservar entregas: elegir una carpeta nueva')
    protected = validate_cache(cache)
    read = lambda name: pd.read_csv(cache / 'csv' / (name + '.csv'))
    original = read('auditoria_112')
    audit = pd.read_csv(review / 'rutas/entrega_final/clasificacion_112.csv')
    screens = read('seleccion_pantallas')
    strata = read('estratos_y_factores')
    relations = read('relaciones_entre_controles')
    assert set(original.id_encuesta) == set(audit.id_encuesta) and len(audit) == audit.id_encuesta.nunique() == 112
    pd.testing.assert_frame_equal(original, audit[original.columns], check_dtype=False)
    eligible = audit.categoria.isin(['A', 'B'])
    # Misma población y geometrías: se puede reutilizar exactamente la selección cacheada.
    assert set(audit.loc[eligible, 'id_encuesta']) == set(original.loc[original.aceptada_corredor_auxiliar, 'id_encuesta'])
    for frame in (audit, screens):
        for col in ('zona_origen', 'zona_destino'):
            frame[col] = pd.to_numeric(frame[col], errors='coerce').astype('Int64')
    audit['id_OD_zonal'] = audit.apply(lambda r: '|'.join(map(str, [r.zona_origen, r.zona_destino, r.campana, r.segmento])) if pd.notna(r.zona_origen) and pd.notna(r.zona_destino) else None, axis=1)
    detailed, zonal, daily = build_matrices(audit, screens, strata)
    comparisons = pd.read_csv(review / 'pantallas/comparaciones_pantallas.csv')
    comparable = comparisons[comparisons.cobertura_completa_ambas & comparisons.fecha.notna()]
    contrast_ids = set(comparable.id_OD_zonal)
    for frame in (detailed, zonal):
        frame['contraste_misma_fecha_disponible'] = frame.id_OD_zonal.isin(contrast_ids)
        frame['advertencia_entre_controles'] = frame.id_OD_zonal.map(
            {r.id_OD_zonal: f'{r.fecha}: {r.control_1}={r.Q_1:.4f} frente a {r.control_2}={r.Q_2:.4f} vehículos/12h. Contraste descriptivo, sin intervalo de confianza; no determina cuál es mejor.' for r in comparable.itertuples()}
        ).fillna('Sin contraste equivalente por fecha y cobertura; consistencia no comprobada')
    selected = screens[screens.seleccionada]
    audit['seleccionada_final'] = audit.apply(lambda r: bool(eligible.loc[r.name] and ((selected.id_OD_zonal.eq(r.id_OD_zonal)) & selected.control.eq(r.control) & selected.sentido_control.eq(r.sentido_observado)).any()), axis=1)
    audit['utilizada_flujo_final'] = audit.seleccionada_final & audit.factor_expansion.notna()
    audit['decision_estimacion'] = ['C: excluida, conservada en auditoría' if c == 'C' else 'Sin expansión: ' + str(reason) if pd.isna(w) else 'Utilizada en pantalla elegida' if chosen else 'Pantalla alternativa, no sumada' for c, w, reason, chosen in zip(audit.categoria, audit.factor_expansion, audit.motivo_sin_factor, audit.seleccionada_final)]
    coverage = audit.groupby(['campana', 'segmento'], dropna=False).agg(entrevistas=('id_encuesta', 'size'), alta_confianza_A=('categoria', lambda s: s.eq('A').sum()), inciertas_B=('categoria', lambda s: s.eq('B').sum()), excluidas_C=('categoria', lambda s: s.eq('C').sum()), con_factor=('factor_expansion', 'count'), utilizadas_flujo=('utilizada_flujo_final', 'sum')).reset_index()
    old = read('matriz_auxiliar_corredor_zonas').set_index('id_OD_zonal')
    balance = all(math.isclose(r.flujo_corredor_12h, old.loc[r.id_OD_zonal, 'flujo_corredor_auxiliar_12h'], rel_tol=1e-12) for r in zonal.itertuples())
    checks = {
        'cache_cotejada_con_Excel_y_hashes_originales': True,
        '112_registros_identicos_mas_capa_clasificacion': True,
        'ningun_OD_ambiguo_reabierto': not audit.loc[eligible, 'OD_homologado'].eq(False).any(),
        'ningun_C_expandido_en_matriz': not audit.loc[audit.categoria.eq('C'), 'utilizada_flujo_final'].any(),
        'una_pantalla_por_OD_campana_segmento': not selected.id_OD_zonal.duplicated().any(),
        'sin_sumar_controles_superpuestos': not relations.se_suman.any(),
        'F_n_N_diario': bool((daily.Q_corredor_12h - daily.F_vehiculos_12h * daily.n_corredor / daily.N_original).abs().lt(1e-10).all()),
        'A_mas_B_igual_corredor_diario': bool((daily.Q_A_12h + daily.Q_B_12h - daily.Q_corredor_12h).abs().lt(1e-10).all()),
        'factores_sin_modificar': audit.factor_expansion.equals(original.factor_expansion),
        'balance_detalle_zona': all(math.isclose(r.flujo_corredor_12h, detailed.loc[detailed.id_OD_zonal.eq(r.id_OD_zonal), 'flujo_corredor_12h'].sum(), rel_tol=1e-12) for r in zonal.itertuples()),
        'mismos_flujos_corredor_previamente_validados': balance,
        'CITRA_y_ARISTO_no_sumados': True,
        'sin_24h_anualizacion_logit_o_captacion': True,
    }
    assert all(checks.values()), checks
    out.mkdir(parents=True)
    (out / 'csv').mkdir()
    tables = {'matriz_localidades': detailed, 'matriz_zonas_ARISTO': zonal, 'auditoria_112': audit,
              'flujos_fecha_pantalla': daily, 'pantallas': screens, 'relaciones_controles': relations,
              'comparaciones_controles': comparisons, 'factores': strata, 'cobertura': coverage}
    for name, frame in tables.items():
        frame.to_csv(out / 'csv' / (name + '.csv'), index=False, encoding='utf-8-sig')
    notes = pd.DataFrame([
        {'tema': 'Campos históricos de auditoría', 'detalle': 'aceptada_principal, motivo_principal, clasificacion_nueva y seleccionada_corredor_auxiliar conservan la ejecución estricta anterior. Las decisiones vigentes son categoria, motivo, seleccionada_final, utilizada_flujo_final y decision_estimacion. No confundir los estados anteriores con la clasificación actual.'},
        {'tema': 'Alcance principal', 'detalle': 'Demanda observada del corredor, expandida proporcionalmente. A es inferencia de alta confianza, no itinerario completo certificado ni demanda capturable.'},
        {'tema': 'Categorías', 'detalle': 'A alta confianza inferida; B corredor incierto; C excluida por OD insuficiente o evidencia incompatible. B=0 en esta cohorte no implica incertidumbre nula.'},
        {'tema': 'Exposición', 'detalle': 'Vehículos de clase por 12h de las fechas indicadas; media de jornadas válidas. No temporada completa, 24h o año. No sumar celdas como demanda total del proyecto.'},
        {'tema': 'Ceros y faltantes', 'detalle': 'Sin muestra expandible en pantalla elegida conserva faltante. Un cero diario o de categoría es muestral del estimador, no ausencia poblacional.'},
        {'tema': 'Consistencia', 'detalle': 'Ver contrastes: hay discrepancias entre controles comparables. La selección por cobertura/exposición se conserva sin favorecer el mayor flujo; no está demostrado cuál estima mejor.'},
        {'tema': 'Geografía', 'detalle': 'Se usa la homologación cerrada previa, preservada por hashes. 29 aprobaciones y 8 extremos pendientes sin imputar. No se reabre revisión manual.'},
        {'tema': 'Viabilidad', 'detalle': 'B: CITRA utilizable parcialmente como referencia terrestre por OD/corredor, no como matriz total ni referencia única de todas las celdas.'},
    ])
    export_excel(out / '01_matrices_corredor.xlsx', {'Leer primero': notes, 'Localidades': detailed, 'Zonas ARISTO': zonal, 'Por fecha y pantalla': daily, 'Cobertura': coverage})
    export_excel(out / '02_auditoria_112.xlsx', {'Leer primero': notes, '112 entrevistas': audit, 'A alta confianza': audit[audit.categoria.eq('A')], 'B inciertas': audit[audit.categoria.eq('B')], 'C excluidas': audit[audit.categoria.eq('C')]})
    export_excel(out / '03_factores_expansion.xlsx', {'Leer primero': notes, 'Estratos y factores': strata, 'Estratos cohorte': strata[strata.id_estrato.isin(audit.id_estrato)], 'Cobertura temporal': pd.read_csv(review / 'aforos/cobertura_temporal.csv')})
    export_excel(out / '04_pantallas_contrastes.xlsx', {'Leer primero': notes, 'Selección': screens, 'Serie y ramas': relations, 'Comparaciones zonales': comparisons, 'Comparaciones localidades': pd.read_csv(review / 'pantallas/comparaciones_detalladas_compatibles.csv'), 'Alternativas por fecha': daily})
    summary = {'viabilidad': 'B — Utilizable parcialmente', 'categorias': {c: int(audit.categoria.eq(c).sum()) for c in ['A', 'B', 'C']},
               'con_factor_A_B': int((eligible & audit.factor_expansion.notna()).sum()), 'utilizadas_pantalla_con_factor': int(audit.utilizada_flujo_final.sum()),
               'celdas_zonales': len(zonal), 'OD_zonales_distintos': int(zonal[['zona_origen', 'zona_destino']].drop_duplicates().shape[0]),
               'celdas_detalladas': len(detailed), 'celdas_detalladas_con_flujo': int(detailed.flujo_corredor_12h.notna().sum()),
               'pantallas_elegidas': len(selected), 'total_proyecto': None,
               'por_que_no_total': 'Cobertura parcial de cohorte, días y pantallas; no sumar distintas exposiciones ni transformar faltantes en cero',
               'rutas_recalculadas': 0, 'verificaciones': checks, 'segundos_integracion': round(time.monotonic() - started, 2)}
    (out / 'resumen_validacion.json').write_text(json.dumps(summary, ensure_ascii=False, indent=2) + '\n')
    from reporte_corredor import report
    report(out, review, summary, detailed, zonal, coverage, comparable, notes)
    sources = [cache / 'validacion.json', cache / 'integridad_antes.json'] + list((cache / 'csv').glob('*.csv'))
    sources += [review / 'rutas/entrega_final/clasificacion_112.csv', review / 'rutas/entrega_final/resumen_rutas.json']
    sources += list(review.glob('*/resumen*.json')) + list((review / 'pantallas').glob('*.csv')) + list((review / 'aforos').glob('*.csv'))
    sources += [review / 'aforos/auditar_aforos.py', review / 'rutas/clasificar_rutas.py', review / 'pantallas/analizar_pantallas.py', review / 'cobertura/auditar_cobertura.py']
    sources += list(Path(__file__).parent.glob('*.py'))
    manifest = {'cache': str(cache.resolve()), 'revision': str(review.resolve()), 'fuentes': [{'archivo': str(x.resolve()), 'sha256': sha(x)} for x in sources],
                'resultados': [{'archivo': str(x.relative_to(out)), 'sha256': sha(x)} for x in out.rglob('*') if x.is_file()],
                'originales_y_entregas_anteriores_sin_cambios': all(sha(r['archivo']) == r['sha256'] for r in protected)}
    assert manifest['originales_y_entregas_anteriores_sin_cambios']
    (out / 'manifest_integridad.json').write_text(json.dumps(manifest, ensure_ascii=False, indent=2) + '\n')
    print(json.dumps(summary, ensure_ascii=False, indent=2))


if __name__ == '__main__':
    main()
