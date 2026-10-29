"""Auditoría rápida del caché validado de expansión. No recalcula rutas ni modifica originales."""
from pathlib import Path
import argparse, csv, hashlib, json, collections, math
import pandas as pd
from openpyxl import load_workbook

def sha(p): return hashlib.sha256(Path(p).read_bytes()).hexdigest()
def main():
 p=argparse.ArgumentParser();p.add_argument('--cache',type=Path,default=Path('etapa1_terrestre_citra/resultados/entrega_02'));p.add_argument('--salida',type=Path,default=Path('etapa1_terrestre_citra/cierre_final/aforos'));args=p.parse_args();cache=args.cache;out=args.salida;out.mkdir(parents=True,exist_ok=True)
 validation=json.loads((cache/'validacion.json').read_text());assert all(validation['verificaciones'].values())
 before=json.loads((cache/'integridad_antes.json').read_text());protected=[]
 for r in before:
  path=Path(r['archivo'])
  if 'eod aysen' in str(path) and path.suffix.lower() in ['.xlsx','.kmz']:
   assert path.exists() and sha(path)==r['sha256'];protected.append(r)
 for r in validation['resultados_xlsx']:
  assert sha(cache/r['archivo'])==r['sha256']
 paths={n:cache/'csv'/n for n in ['estratos_y_factores.csv','universo_original_denominadores.csv','auditoria_112.csv']}
 ss=pd.read_csv(paths['estratos_y_factores.csv'],encoding='utf-8-sig',keep_default_na=False)
 universe=pd.read_csv(paths['universo_original_denominadores.csv'],encoding='utf-8-sig',keep_default_na=False)
 audit=pd.read_csv(paths['auditoria_112.csv'],encoding='utf-8-sig',keep_default_na=False)
 assert len(universe)==8828 and universe.id_encuesta.nunique()==8828
 # Se verifica N a partir de TODO el universo original cacheado, sin filtrar el proyecto.
 minute=pd.to_numeric(universe.minuto,errors='coerce');inside=universe[minute.ge(480)&minute.lt(1200)]
 counts=inside.groupby('id_estrato',sort=False).size();ss['N_recontado_universo_original']=ss.id_estrato.map(counts).fillna(0).astype(int)
 assert (ss.N_recontado_universo_original==ss.N_entrevistas_originales_ventana.astype(int)).all()
 factor=pd.to_numeric(ss.factor_expansion,errors='coerce');F=pd.to_numeric(ss.aforo_vehiculos_12h,errors='coerce');N=ss.N_recontado_universo_original;valid=factor.notna()
 assert ((factor[valid]*N[valid]-F[valid]).abs()<1e-8).all()
 assert (ss.loc[valid,'intervalos_unicos'].astype(int)==48).all()
 assert (ss.loc[valid,'intervalos_repetidos'].astype(int)==0).all()
 assert ss.loc[valid,'motivo'].eq('').all()
 assert ss.loc[valid,'fecha'].str.startswith('2026-').all()
 ss['validacion_N_universo']=True;ss['balance_F_N_validado']=valid;ss['factor_utilizable']=valid
 # Comparar los CSV reutilizados con el Excel cuyo SHA fue registrado en la entrega validada.
 wb=load_workbook(cache/'04_pantallas_y_expansion.xlsx',read_only=True,data_only=True)
 for sheet,csvframe in [('Estratos y factores',ss),('Universo denominadores',universe)]:
  data=wb[sheet].iter_rows(values_only=True);head=next(data);xx={str(r[0]):dict(zip(head,r)) for r in data};assert len(xx)==len(csvframe)
  k=csvframe.columns[0]
  for row in csvframe.to_dict('records'):
   rr=xx[str(row[k])]
   for col in ['factor_expansion','N_entrevistas_originales_ventana'] if sheet.startswith('Estratos') else ['hash_registro','id_estrato','factor']:
    a,b=row[col],rr[col]
    if a=='' and b is None:continue
    try:assert math.isclose(float(a),float(b),rel_tol=1e-12,abs_tol=1e-12)
    except (ValueError,TypeError):assert str(a)==str(b),(sheet,col,a,b)
 wb.close()
 ss.to_csv(out/'factores_validados.csv',index=False,encoding='utf-8-sig')
 relevant=ss[ss.id_estrato.isin(set(audit.id_estrato))];relevant.to_csv(out/'estratos_cohorte_112.csv',index=False,encoding='utf-8-sig')
 coverage=[]
 for key,group in ss.groupby(['campana','control','sentido','segmento'],sort=True):
  vv=pd.to_numeric(group.factor_expansion,errors='coerce').notna();gg=group[vv]
  coverage.append(dict(zip(['campana','control','sentido','segmento'],key))|{'estratos':len(group),'estratos_validos':int(vv.sum()),'fechas_validas':'|'.join(sorted(gg.fecha.unique())),'N_compatible_validado':int(gg.N_recontado_universo_original.sum()),'factor_minimo':float(pd.to_numeric(gg.factor_expansion).min()) if len(gg) else None,'factor_maximo':float(pd.to_numeric(gg.factor_expansion).max()) if len(gg) else None,'motivos_sin_expansion':';'.join(sorted({m for text in group.loc[~vv,'motivo'] for m in text.split(';') if m})),'periodo':'08:00–20:00, solo fechas listadas; sin extrapolacion diaria/anual'})
 pd.DataFrame(coverage).to_csv(out/'cobertura_temporal.csv',index=False,encoding='utf-8-sig')
 issues=collections.Counter(m for text in ss.loc[~valid,'motivo'] for m in text.split(';') if m)
 cohort_factor=pd.to_numeric(audit.factor_expansion,errors='coerce');corr=audit.aceptada_corredor_auxiliar.astype(str).eq('True');selected=audit.seleccionada_corredor_auxiliar.astype(str).eq('True')
 summary={'estado':'correcto','reutilizacion_cache':True,'sin_reconstruir_red_ni_rutas':True,'fuentes_originales_hash_verificado':protected,'hashes_entradas_cache':{name:sha(path) for name,path in paths.items()},'comparacion_csv_con_excel_validado':True,'estratos_totales':len(ss),'estratos_factor_valido':int(valid.sum()),'estratos_sin_expansion':int((~valid).sum()),'entrevistas_universo_original':len(universe),'entrevistas_en_ventana_original':len(inside),'entrevistas_en_estratos_validos':int(N[valid].sum()),'N_validado_antes_de_filtros':True,'campanas_fechas_validas':{camp:sorted(g.fecha.unique()) for camp,g in ss[valid].groupby('campana')},'problemas_estratos_no_sumables':dict(issues),'estratos_cohorte_112':len(relevant),'estratos_cohorte_validos':int(relevant.factor_utilizable.sum()),'fechas_cohorte_validas':sorted(relevant[relevant.factor_utilizable].fecha.unique()),'cohorte':len(audit),'cohorte_compatibles_corredor':int(corr.sum()),'cohorte_corredor_con_factor':int((corr&cohort_factor.notna()).sum()),'cohorte_seleccionada_con_factor':int((selected&cohort_factor.notna()).sum()),'recomendaciones':['Reutilizar los 186 factores validados, sin recalcular red o rutas.','Calcular Q=F*n/N por categoria usando N original; cambiar A/B solo modifica n.','No usar factores donde esten ausentes; mantener casos sin expansion y motivos.','No extrapolar 08–20 a 24h, temporada o año.','No sumar flujos de controles en serie ni de fechas dispares como total del proyecto.','La representatividad y no respuesta dentro del estrato no fueron certificadas por un protocolo de muestreo; expansion proporcional condicionada.','Los aforos urbanos no intervienen en esta auditoria ni sustituyen los aforos de seccion EOD.'],'fuente_metodologica':'etapa1_terrestre_citra/scripts/fuentes.py: expansion; entrega_02/04_pantallas_y_expansion.xlsx, Estratos y factores y Universo denominadores; EOD CITRA hojas EOD y Flujos'}
 (out/'resumen_aforos.json').write_text(json.dumps(summary,ensure_ascii=False,indent=2)+'\n')
 print(json.dumps({k:v for k,v in summary.items() if k not in ['fuentes_originales_hash_verificado','hashes_entradas_cache','campanas_fechas_validas']},ensure_ascii=False,indent=2))
if __name__=='__main__':main()
