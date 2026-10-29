"""Lectura de originales, estratos y capa geográfica aprobada. Sin modelos de elección."""
from pathlib import Path
import csv,json,hashlib,datetime,collections,math,re,zipfile,xml.etree.ElementTree as ET,unicodedata
from openpyxl import load_workbook

def readcsv(p):return list(csv.DictReader(Path(p).open(encoding='utf-8-sig',newline='')))
def sha(p):return hashlib.sha256(Path(p).read_bytes()).hexdigest()
def norm(v):return re.sub(r'\s+',' ',''.join(c for c in unicodedata.normalize('NFD',str(v or '').upper().replace('´',"'").replace('’',"'")) if unicodedata.category(c)!='Mn')).strip()
def canon(v):
 n=norm(v);return {'COIHAIQUE':'COYHAIQUE','COYHAYQUE':'COYHAIQUE','RIO TRANQUILO':'PUERTO RIO TRANQUILO','PUERTO TRANQUILO':'PUERTO RIO TRANQUILO','PUERTO INGENIERO IBANEZ':'PUERTO IBANEZ','PTO IBANEZ':'PUERTO IBANEZ','PTO. IBANEZ':'PUERTO IBANEZ','CERRO CASTILLO':'VILLA CERRO CASTILLO','GUADAL':'PUERTO GUADAL','BAHIA MURTA':'PUERTO MURTA','MURTA':'PUERTO MURTA','AISEN':'AYSEN'}.get(n,n)
def num(v):
 try:
  f=float(v);return f if math.isfinite(f) else None
 except (TypeError,ValueError):return None

def typed(v):
 if isinstance(v,(datetime.datetime,datetime.date,datetime.time)):return dict(tipo=type(v).__name__,valor=v.isoformat())
 return dict(tipo=type(v).__name__,valor=v)
def payload(vals):return json.dumps([typed(v) for v in vals],ensure_ascii=False,separators=(',',':'),allow_nan=False)
def date(v):
 if isinstance(v,(datetime.datetime,datetime.date)):return v.date().isoformat() if isinstance(v,datetime.datetime) else v.isoformat()
 if isinstance(v,(int,float)):return (datetime.datetime(1899,12,30)+datetime.timedelta(days=v)).date().isoformat()
 try:return datetime.date.fromisoformat(str(v).strip()[:10]).isoformat()
 except ValueError:return ''
def minutes(h,m=None,summer=False):
 if summer:
  a,b=num(h),num(m)
  return int(a)*60+int(b) if a is not None and b is not None and a==int(a) and b==int(b) and 0<=a<24 and 0<=b<60 else None
 if isinstance(h,datetime.datetime):h=h.time()
 if isinstance(h,datetime.time):return h.hour*60+h.minute
 if isinstance(h,(int,float)) and 0<=h<1:return int(round(h*1440)) if h*1440<1440 else None
 match=re.fullmatch(r'\s*(\d{1,2}):(\d{1,2})(?::\d{1,2}(?:\.\d+)?)?\s*',str(h))
 if not match:return None
 a,b=map(int,match.groups());return a*60+b if 0<=a<24 and 0<=b<60 else None

def read_originals(base):
 interviews=[];flows=[];headers=[]
 for camp,name in [('Verano','EOD Aysen Febrero 2026 vf.xlsx'),('Otoño','EOD Aysen Abril 2026.xlsx')]:
  path=Path(base)/'eod aysen'/name;summer=camp=='Verano'
  w=load_workbook(path,read_only=True,data_only=False,keep_links=False);it=w['EOD'].iter_rows(values_only=True);hh=next(it)
  headers.extend(dict(archivo=name,hoja='EOD',columna=i+1,encabezado_original=h) for i,h in enumerate(hh))
  for rn,vals in enumerate(it,2):
   if not any(v is not None for v in vals):continue
   typ=norm(vals[3]) if summer else str(vals[2]);seg={'LIVIANO':'VL','CAMION':'CAM','1':'VL','2':'CAM'}.get(typ,'DESCONOCIDO')
   clock=minutes(vals[5],vals[6],True) if summer else minutes(vals[4]);dt=date(vals[2] if summer else vals[0]);encoded=payload(vals)
   row=dict(id_encuesta=('V' if summer else 'O')+f'{rn:05d}',archivo=name,hoja='EOD',fila_excel=rn,campana=camp,fecha=dt,control=str(vals[0] if summer else vals[1]),sentido=str(vals[4] if summer else vals[3]),segmento=seg,modo_fuente=typ,minuto=clock,hora_original=json.dumps(typed(vals[5] if summer else vals[4]),ensure_ascii=False),hash_registro=hashlib.sha256(encoded.encode()).hexdigest(),registro_original_JSON=encoded,origen_original='' if vals[8 if summer else 7] is None else str(vals[8 if summer else 7]),destino_original='' if vals[12 if summer else 8] is None else str(vals[12 if summer else 8]),ejes_original=vals[23] if summer else vals[5])
   interviews.append(row)
  w.close();w=load_workbook(path,read_only=True,data_only=True,keep_links=False);it=w['Flujos'].iter_rows(values_only=True);next(it)
  for rn,v in enumerate(it,2):
   if not any(x is not None for x in v[:21]):continue
   components=[num(v[i]) for i in [6,7,16,12,13]];vl,cam=num(v[19]),num(v[20]);mm=minutes(v[3],v[4],True)
   problems=[]
   if any(x is None or x<0 for x in components) or vl is None or cam is None:problems.append('aforo_componentes_ausentes_o_invalidos')
   else:
    if not math.isclose(vl,sum(components[:3]),abs_tol=1e-9):problems.append('VL_no_concilia_con_componentes')
    if not math.isclose(cam,sum(components[3:]),abs_tol=1e-9):problems.append('CAM_no_concilia_con_componentes')
   flows.append(dict(archivo=name,hoja='Flujos',fila_excel=rn,campana=camp,fecha=date(v[2]),control=str(v[0]),sentido=str(v[5]),minuto=mm,VL=vl,CAM=cam,incidencias=';'.join(problems)))
  w.close()
 assert len(interviews)==8828 and len({r['id_encuesta'] for r in interviews})==8828
 assert len(flows)==5760
 return interviews,flows,headers

def kmz(base):
 p=Path(base)/'eod aysen/Mediciones Aysen (23-02-26).kmz';z=zipfile.ZipFile(p);root=ET.fromstring(z.read(next(n for n in z.namelist() if n.endswith('.kml'))));ns={'k':'http://www.opengis.net/kml/2.2'};out={}
 for pm in root.findall('.//k:Placemark',ns):
  name=pm.findtext('k:name',namespaces=ns);co=pm.findtext('.//k:Point/k:coordinates',namespaces=ns)
  if co:
   lon,lat,*_=map(float,co.strip().split(','));out[name]=dict(nombre=name,lat=lat,lon=lon,fuente=p.name+'; Placemark '+name)
 return out

def expansion(interviews,flows):
 by=collections.defaultdict(list);ff=collections.defaultdict(list)
 for r in interviews:by[(r['campana'],r['fecha'],r['control'],r['sentido'],r['segmento'])].append(r)
 for r in flows:
  for seg in ['VL','CAM']:ff[(r['campana'],r['fecha'],r['control'],r['sentido'],seg)].append(r)
 dup=collections.Counter((r['archivo'],r['hash_registro']) for r in interviews)
 strata=[];weights={};index={}
 for key in sorted(set(by)|set(ff)):
  rr,fs=by.get(key,[]),ff.get(key,[]);inside=[r for r in rr if r['minuto'] is not None and 480<=r['minuto']<1200];N=len(inside);problems=[];grid=collections.Counter(r['minuto'] for r in fs)
  if not fs:problems.append('sin_aforo_misma_fecha_punto_sentido_clase')
  elif set(grid)!=set(range(480,1200,15)):problems.append('grilla_12h_incompleta_o_invalida')
  if any(v>1 for v in grid.values()):problems.append('intervalos_duplicados_no_reconciliados')
  if not key[1].startswith('2026-'):problems.append('fecha_fuera_campana_no_corregida')
  if any(r['minuto'] is None for r in rr):problems.append('hora_entrevista_invalida_impide_N_completo')
  if key[4]=='DESCONOCIDO':problems.append('clase_no_identificada')
  if any(dup[r['archivo'],r['hash_registro']]>1 for r in rr):problems.append('filas_EOD_identicas_no_resueltas')
  if not N:problems.append('sin_entrevistas_compatibles_en_ventana')
  if any(r['incidencias'] for r in fs):problems.append('aforo_no_conciliado_con_clases_fisicas')
  F=sum(r[key[4]] for r in fs) if fs and key[4] in ['VL','CAM'] and all(r[key[4]] is not None and r[key[4]]>=0 for r in fs) else None
  if F is not None and N>F:problems.append('entrevistas_superan_aforo')
  factor=F/N if not problems and F is not None and N else None
  s=dict(id_estrato='|'.join(key),campana=key[0],fecha=key[1],control=key[2],sentido=key[3],segmento=key[4],aforo_vehiculos_12h=F,N_entrevistas_originales_ventana=N,entrevistas_estrato=len(rr),fuera_ventana=sum(r['minuto'] is not None and not(480<=r['minuto']<1200) for r in rr),horas_invalidas=sum(r['minuto'] is None for r in rr),filas_aforo=len(fs),intervalos_unicos=len(grid),intervalos_repetidos=sum(v-1 for v in grid.values() if v>1),factor_expansion=factor,estado='Compatible para expansión proporcional condicionada' if factor is not None else 'Sin expansión',motivo=';'.join(problems),fuente_aforo=';'.join(sorted({r['archivo']+'/Flujos' for r in fs})),filas_aforo_fuente='|'.join(str(r['fila_excel']) for r in fs),filas_entrevistas_fuente='|'.join(str(r['fila_excel']) for r in rr),balance_N_por_factor=N*factor if factor is not None else None,formula='w=F_vehiculos_clase_12h/N_todas_entrevistas_clase_original_08_20',unidad='vehículos de clase / ventana de12h del día indicado',diseno='Proporción dentro de estrato supuesta representativa; no ponderador oficial certificado')
  strata.append(s);index[key]=s
  for r in rr:
   reasons=list(problems)
   if r['minuto'] is None:reasons.append('hora_individual_invalida')
   elif not(480<=r['minuto']<1200):reasons.append('entrevista_individual_fuera_ventana')
   weights[r['id_encuesta']]=dict(id_estrato=s['id_estrato'],factor=factor if not reasons else None,motivo=';'.join(dict.fromkeys(reasons)))
 return strata,weights,index
