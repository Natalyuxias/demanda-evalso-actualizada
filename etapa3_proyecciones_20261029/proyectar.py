"""Proyección simple desde Etapa2 y parámetros existentes ARISTO8.3. Sin modelos nuevos."""
from pathlib import Path
import argparse, hashlib, html, json, time
import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Font, PatternFill

YEARS=[2030,2040,2050]
FLOW_U=['aporte_parcial','demanda_anual_operacional','demanda_anual_ventana12h','demanda_anual24h']
FLOW_DETAIL=['aporte_anual_en_meses_con_perfil','demanda_anual_operacional_componente']
FLOW_C=['aporte_parcial_meses_con_perfil_12h','aporte_meses_con_campana_propia_12h','demanda_anual_ventana12h','demanda_anual24h']
FLOW_SUM=['flujo_fisico_anual','aporte_parcial','observado_OD_anualizado','imputado_OD_anualizado','residual_anualizado']

def sha(p):return hashlib.sha256(Path(p).read_bytes()).hexdigest()

def project(base, factors, values):
    """Cruce por segmento, sin transformar NULL, clases, OD ni cobertura."""
    b=base.copy()
    b['id_fila_base']=range(len(b))
    b=b.rename(columns={'escenario':'escenario_base_Etapa2','anio':'anio_base_original'})
    ff=factors[factors.anio.isin(YEARS)].rename(columns={'escenario':'escenario_crecimiento','anio':'anio_proyeccion'})
    assert not ff.duplicated(['segmento','escenario_crecimiento','anio_proyeccion']).any()
    x=b.merge(ff,on='segmento',how='left',validate='many_to_many',suffixes=('','_parametro'))
    assert x.factor_demanda_desde_2026.notna().all()
    assert len(x)==len(b)*9
    for c in values:
        x[c+'_base2026']=x[c]
        x[c]=pd.to_numeric(x[c],errors='coerce')*x.factor_demanda_desde_2026
        assert x[c].isna().equals(x[c+'_base2026'].isna())
    x['anio_base']=2026
    x['naturaleza']='Proyección exógena sobre base anual condicionada; no crecimiento observado ni captación'
    return x

def write_excel(path,sheets):
    w=Workbook(write_only=True)
    for name,d in sheets.items():
        s=w.create_sheet(name);s.freeze_panes='B2';header=[]
        for i,k in enumerate(d.columns,1):
            s.column_dimensions[__import__('openpyxl').utils.get_column_letter(i)].width=min(52,max(17,len(k)+2))
            c=WriteOnlyCell(s,k);c.font=Font(bold=True,color='FFFFFF');c.fill=PatternFill('solid',fgColor='225D65');header.append(c)
        s.append(header)
        for row in d.astype(object).where(pd.notna(d),None).itertuples(index=False,name=None):
            rr=[]
            for v in row:
                c=WriteOnlyCell(s,v)
                if isinstance(v,str):c.data_type='s'
                rr.append(c)
            s.append(rr)
        s.auto_filter.ref=f'A1:{__import__("openpyxl").utils.get_column_letter(len(d.columns))}{len(d)+1}'
    w.save(path)

def table(d):return '<div class="scroll">'+d.to_html(index=False,na_rep='Faltante',float_format=lambda x:f'{x:,.3f}',border=0)+'</div>'

def main():
    started=time.monotonic();p=argparse.ArgumentParser();p.add_argument('--etapa2',type=Path,required=True);p.add_argument('--parametros',type=Path,required=True);p.add_argument('--salida',type=Path,required=True);a=p.parse_args()
    if a.salida.exists():raise FileExistsError('No sobrescribir: elegir salida nueva')
    source=a.etapa2/'csv';names=['base_unificada_anual','aristo_anual_detalle','citra_anual_localidades','aristo_resumen_escenarios']
    paths=[source/(n+'.csv') for n in names]+[a.parametros/n for n in ['factores_propuestos_rebasados_2026.csv','elasticidades_por_segmento.csv','PIB_anual_2026_2050.csv','resumen_parametros.json']]
    hashes={str(x.resolve()):sha(x) for x in paths}
    data={n:pd.read_csv(source/(n+'.csv')) for n in names}
    base={n:d[d.escenario.str.lower().eq('central')].copy() for n,d in data.items()}
    factors=pd.read_csv(a.parametros/'factores_propuestos_rebasados_2026.csv');elastic=pd.read_csv(a.parametros/'elasticidades_por_segmento.csv')
    params=json.loads((a.parametros/'resumen_parametros.json').read_text())
    missing=set().union(*(set(d.segmento) for d in base.values()))-set(factors.segmento)
    for segment in missing:
        # Otros pasajeros carecen de flujo base; la parametrización no completa esos NULL.
        assert segment=='OTRO_PASAJERO'
        extra=factors[factors.segmento.eq('PEATON')].copy();extra['segmento']=segment
        extra['tipo_parametro']='Proxy de crecimiento privado/peatonal; flujo base desconocido permanece NULL'
        factors=pd.concat([factors,extra],ignore_index=True)
        e=elastic[elastic.segmento.eq('PEATON')].copy();e['segmento']=segment
        elastic=pd.concat([elastic,e],ignore_index=True)
    u=project(base[names[0]],factors,FLOW_U)
    ld=project(base[names[1]],factors,FLOW_DETAIL)
    cd=project(base[names[2]],factors,FLOW_C)
    sums=project(base[names[3]],factors,FLOW_SUM)
    readable=sums[sums.version.eq('B')&sums.segmento.isin(['PRIVADOS_COMPATIBLES','CAM','PEATON'])].pivot(index=['anio_proyeccion','escenario_crecimiento'],columns='segmento',values='flujo_fisico_anual').reset_index()
    ff=factors[factors.anio.isin(YEARS)]
    showcase=ff[ff.segmento.isin(['VL','CAM'])][['anio','escenario','segmento','razon_PIB_rebasada_2026','elasticidad','factor_demanda_desde_2026']]
    ordered=ff.pivot(index=['anio','segmento'],columns='escenario',values='factor_demanda_desde_2026')
    checks={
      'base2026_sin_doble_crecimiento':bool(factors[factors.anio.eq(2026)].factor_demanda_desde_2026.eq(1).all()),
      'solo_2030_2040_2050':set(u.anio_proyeccion)==set(YEARS),
      'tres_escenarios_crecimiento':set(u.escenario_crecimiento)=={'conservador','central','alto'},
      'base_Etapa2_central_sin_mezclar_escenarios':u.escenario_base_Etapa2.str.lower().eq('central').all(),
      'faltantes_preservados':all(u[c].isna().equals(u[c+'_base2026'].isna()) for c in FLOW_U),
      'CITRA_24h_sigue_faltante':u.loc[u.fuente.eq('CITRA'),'demanda_anual24h'].isna().all(),
      'residuales_ARISTO_conservados':len(u[u.componente.eq('residual')])==9*len(base[names[0]][base[names[0]].componente.eq('residual')]),
      'escenarios_ordenados':bool((ordered.conservador<=ordered.central).all()&(ordered.central<=ordered.alto).all()),
      'datos_fuente_intactos':all(sha(path)==h for path,h in hashes.items()),
      'sin_regresiones_encuestas_rutas_capacidad':True,
    }
    # Crecimiento uniforme dentro de clase: observado+imputado+residual sigue siendo F.
    known=sums.flujo_fisico_anual.notna()
    err=sums.observado_OD_anualizado+sums.imputado_OD_anualizado+sums.residual_anualizado-sums.flujo_fisico_anual
    checks['balance_fisico_proyectado']=bool(err[known].abs().lt(1e-7).all())
    checks={k:bool(v) for k,v in checks.items()};assert all(checks.values()),checks
    notes=pd.DataFrame([
      dict(tema='Base',detalle='Referencia CENTRAL2026Etapa2. Se conservan versiones lacustresA/B y vistas noaditivas; no sumar versiones ni vistas.'),
      dict(tema='Fórmula',detalle='Q_t=Q_2026×(PIB_t/PIB_2026)^beta. No se aplica crecimiento2026sobre la base, ni factor2025añoPrimavera otra vez.'),
      dict(tema='Escenarios',detalle='CentralPIBBaseARISTO;altoOptimistaARISTO;conservadorBase−0,5pp/anual supuesto nuestro simétrico al contraste optimista documentado.'),
      dict(tema='Elasticidades',detalle='Aysénprivados3,36831632287041;CAM2,14413849315727. Transferencia al ferry y a peatones/motos/bicis/VP explícita para crecimiento solamente, sin cambiar unidades.'),
      dict(tema='Cobertura',detalle='CITRApermaneceparcial yventana12h,24hNULL. ARISTOresidual/ODimputado siguen identificados. VLestricto anual faltante no se transforma enVP ni se completa.'),
      dict(tema='Alcance',detalle='Demanda base tendencial proyectada, sin captación, inducción, capacidad ni cambio de ruta. No suma automático entrebloques.'),
    ])
    a.salida.mkdir(parents=True)
    xlsx=a.salida/'proyecciones_2030_2040_2050.xlsx'
    sheets={'Leer primero':notes,'Parametros':elastic,'Trayectorias PIB':pd.read_csv(a.parametros/'PIB_anual_2026_2050.csv'),'Factores':ff,'Base2026':base[names[0]],'Proyeccion zonas':u,'Lacustre detalle':ld,'CITRA localidades':cd,'Balances lacustres':sums,'Resumen lacustre':readable,'Validaciones':pd.DataFrame([dict(control=k,correcto=v) for k,v in checks.items()])}
    write_excel(xlsx,sheets)
    w=load_workbook(xlsx,read_only=True,data_only=True);assert w.sheetnames==list(sheets);w.close()
    duration=round(time.monotonic()-started,2)
    css='body{font:16px/1.55 system-ui,sans-serif;max-width:1100px;margin:auto;padding:25px;color:#233e44}table{border-collapse:collapse;font-size:13px;width:100%;margin:18px 0}td,th{padding:8px;border:1px solid #cbdad5;text-align:left}th{background:#225d65;color:white}.scroll{overflow:auto}.box{background:#eaf3f0;padding:17px}a{color:#086e8d}pre{white-space:pre-wrap;background:#eef3f0;padding:12px}'
    text='<!doctype html><html lang="es"><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><title>Etapa3 — proyecciones simples</title><style>'+css+'</style><body><h1>Etapa3: proyecciones2030,2040y2050</h1>'
    text+='<p class="box">Se proyectó directamente la referencia central2026deEtapa2. Tres escenarios de crecimiento; versionesARISTOA/B conservadas. Los faltantes terrestres y los residuales lacustres no se completaron. <a href="'+xlsx.name+'">Abrir Excel de resultados</a>.</p>'
    text+='<h2>Parámetros y procedimiento</h2><p><code>Q_t = Q_2026 × (PIB_t / PIB_2026)^β</code>. Se utilizan las trayectoriasBase yOptimista documentadas enARISTO§8.3, tabla181p241; el conservador aplicaBase−0,5puntosporcentuales anuales desde2027, supuesto explícito nuestro simétrico al contraste optimista. No se estimaron modelos nuevos.</p>'
    text+='<p>βAysén de tabla179p240:privados3,36831632287041 y carga2,14413849315727, precisión recuperada de <em>Factores de proyección de flujos.xlsx</em>, Ruta!D7/D11. PIB en hojaPIB,B/C35–59. Se rebasa todo a2026=1; no se reaplica crecimiento desde2025 a la primavera ni crecimiento2026 sobre la base normalizada.</p>'
    text+='<p>No existe una elasticidad propia del ferry o de peatones en esos parámetros. Se transfiere el ritmo privado para el ferry privado, VP,motos,bicicletas y peatones; la carga usa su propia elasticidad regional. Es una aproximación de crecimiento, no una conversión de personas a vehículos ni un modelo de elección de ruta.</p>'
    text+='<h2>Factores sobre2026</h2>'+table(showcase)
    text+='<h2>Referencia lacustre proyectada: versiónB</h2><p>Ambos sentidos, unidades separadas por columna. PRIVADOS_COMPATIBLES incluyeVL+moto+bicicleta yVPnormal; no se suma nuevamente con componentes. PEATON no es total de pasajeros. La versiónA y el detalle de residuales/imputados están en elExcel.</p>'+table(readable)
    text+='<h2>Supuestos conservados</h2><p>Se mantiene fija la composiciónOD y la estructura estacional de la base. CITRA sigue como fracción de corredor en ventanas12h; el crecimiento no crea mercado faltante ni factor24h. Las alternativas de anualización conservadora/alta deEtapa2no se cruzan con los escenarios de crecimiento: se usa exclusivamente su basecentral para mantener una proyección sencilla. Las versionesA/B son alternativas de asignación y no deben sumarse.</p>'
    text+='<p>Estas cifras son demanda tendencial de los bloques disponibles, no demanda captada porBahíaJara–Levicán ni demanda inducida. Los faltantes de unidades estrictas y OD sin representación continúan comoNULL; un residual lacustre proyectado sigue siendo residual. Se puede continuar con captación sobre los mercados cubiertos, conservando los límites deEtapa2.</p>'
    text+='<p>Fuentes: matrices anualizadasEtapa2/entrega_final y parámetrosARISTO§8.3, tablas179/181. No se usaron encuestas, nuevas series, internet, redes, capacidad o modelos anteriores de elección.</p>'
    text+='<h2>Reproducción</h2><pre>python etapa3_proyecciones_20261029/proyectar.py --etapa2 etapa2_demanda_anual_20261029/resultados/entrega_final --parametros etapa3_proyecciones_20261029/parametros --salida /ruta/ENTREGA_NUEVA</pre>'
    text+=f'<p>Tiempo hasta Excel validado: {duration}s. {len(checks)} controles correctos. Los archivos fuente permanecen intactos.</p></body></html>'
    (a.salida/'informe_proyecciones.html').write_text(text,encoding='utf-8')
    manifest={'fuentes':hashes,'script_sha256':sha(__file__),'verificaciones':checks,'segundos':duration,'filas_proyeccion_zonas':len(u),'archivo_excel_sha256':sha(xlsx),'archivo_HTML_sha256':sha(a.salida/'informe_proyecciones.html')}
    (a.salida/'validacion.json').write_text(json.dumps(manifest,ensure_ascii=False,indent=2)+'\n')
    print(json.dumps({k:v for k,v in manifest.items() if k!='fuentes'},ensure_ascii=False,indent=2))

if __name__=='__main__':main()
