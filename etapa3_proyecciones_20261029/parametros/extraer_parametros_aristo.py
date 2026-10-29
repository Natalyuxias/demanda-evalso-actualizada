"""Extracción ligera y escenarios de crecimiento existentes. Sin EOD, internet o calibración."""
from pathlib import Path
import csv,json,hashlib,math
from openpyxl import load_workbook
OUT=Path(__file__).parent
SOURCE=Path('Cap 8/Factores de proyección de flujos.xlsx')
def write(name,rows):
 with (OUT/name).open('w',encoding='utf-8-sig',newline='') as f:
  w=csv.DictWriter(f,fieldnames=list(rows[0]));w.writeheader();w.writerows(rows)
def main():
 before=hashlib.sha256(SOURCE.read_bytes()).hexdigest();w=load_workbook(SOURCE,read_only=True,data_only=True)
 pib={int(r[0]):dict(base=float(r[1]),optimista=float(r[2]),fila=i) for i,r in enumerate(w['PIB'].iter_rows(values_only=True),1) if isinstance(r[0],(int,float)) and r[0]>=2025}
 beta={'privados':float(w['Ruta'].cell(7,4).value),'CAM':float(w['Ruta'].cell(11,4).value)}
 gdp=[]
 for year in range(2026,2051):
  base=pib[year]['base']/pib[year-1]['base']-1;opt=pib[year]['optimista']/pib[year-1]['optimista']-1
  gdp.append(dict(anio=year,PIB_base=pib[year]['base'],PIB_optimista=pib[year]['optimista'],crecimiento_PIB_base=base,crecimiento_PIB_optimista=opt,fuente=f'Factores de proyección de flujos.xlsx; PIB fila{pib[year]["fila"]}; tabla181 Informe p241',aplicacion_base2026='2026 es nivelbase,NO aplicar crecimiento2026; productos comienzan en2027'))
 write('PIB_anual_2026_2050.csv',gdp)
 elastic=[]
 for seg in ['VL','VP','PRIVADOS_COMPATIBLES','PEATON','MOTO','BICI','CAM','BUS_TP']:
  cargo=seg=='CAM';b=beta['CAM' if cargo else 'privados'];exact=seg in ['VL','CAM']
  elastic.append(dict(segmento=seg,elasticidad_PIB=b,region_modelo='Aysén',grupo_ARISTO='Camiones' if cargo else 'Vehículos privados',tipo='Coeficiente regional transferido al segmento' if exact else 'PROXY de crecimiento privados; NO elasticidad propia del segmento',ferry='Proxy de crecimiento de carretera Aysén; no modelo lacustre calibrado',unidades='Conservar unidad fisica original. Proxy crecimiento NO convierte personas en vehiculos ni VP en VL.',fuente='Informe tabla179p240; ExcelRuta '+('D11' if cargo else 'D7')))
 write('elasticidades_por_segmento.csv',elastic)
 factors=[]
 for s in elastic:
  for scenario in ['conservador','central','alto']:
   ratio=1.
   for year in range(2026,2051):
    if year>2026:
     rate=(pib[year]['base']/pib[year-1]['base']-1)-.005 if scenario=='conservador' else pib[year]['base']/pib[year-1]['base']-1 if scenario=='central' else pib[year]['optimista']/pib[year-1]['optimista']-1
     ratio*=1+rate
    else:rate=0.
    F=ratio**s['elasticidad_PIB'];factors.append(dict(segmento=s['segmento'],escenario=scenario,anio=year,base=2026,tasa_PIB_aplicada_ese_anio=rate,razon_PIB_rebasada_2026=ratio,elasticidad=s['elasticidad_PIB'],factor_demanda_desde_2026=F,tipo_parametro=s['tipo'],supuesto_escenario='SUPUESTOnuestro PIBBaseARISTOmenos0.5puntosporcentuales2027–50; sensibilidadsimetrica NO escenarioestimadoARISTO' if scenario=='conservador' else 'Trayectoria PIBBase deARISTO' if scenario=='central' else 'Trayectoria PIBOptimista deARISTO +0.5pp',fuente=s['fuente']+'; PIB tabla181p241'))
 write('factores_propuestos_rebasados_2026.csv',factors)
 local=[]
 for rn in range(24,36):
  rr=[w['Ruta'].cell(rn,c).value for c in range(1,24)]
  for scenario,offset in [('Base',13),('Optimista',20)]:
   for j,year in enumerate([2030,2040,2050]):
    original=float(rr[offset+j]);rebased=(pib[year]['base' if scenario=='Base' else 'optimista']/pib[2026]['base' if scenario=='Base' else 'optimista'])**float(rr[4]);local.append(dict(punto=rr[0],region=rr[1],temporada=rr[2],clase=rr[3],elasticidad=rr[4],escenario_ARISTO=scenario,anio=year,anio_base_factor_original=rr[11],factor_original_Excel=original,factor_rebasado2026=rebased,fuente=f'Factores de proyección de flujos.xlsx; Ruta fila{rn}, tabla182Informe p242',regla='Si baseoriginal2025, dividir por factor2026/base2025 delMISMOescenario; no usarlo directo sobre base2026'))
 write('factores_locales_ARISTO_2030_2040_2050.csv',local)
 w.close();assert before==hashlib.sha256(SOURCE.read_bytes()).hexdigest()
 assert all(r['factor_demanda_desde_2026']==1 for r in factors if r['anio']==2026)
 for s in elastic:
  for y in range(2026,2051):
   rows={r['escenario']:r['factor_demanda_desde_2026'] for r in factors if r['segmento']==s['segmento'] and r['anio']==y};assert rows['conservador']<=rows['central']+1e-10<=rows['alto']+1e-10
 result={'estado':'correcto','fuente':str(SOURCE),'sha256':before,'beta_Aysen_privados':beta['privados'],'beta_Aysen_CAM':beta['CAM'],'sin_beta_peatones_o_lacustre_propio':True,'formula':'Q(t)=Q(2026)*(PIB(t)/PIB(2026))**beta','rebase':'No aplicar2026ni factor2025Primavera encima base2026; factor2026=1en todos escenarios.','escenarios':{'central':'PIBBase ARISTO,βregionalAysén','alto':'PIBOptimista ARISTO,mismaβ','conservador':'PIBBaseARISTOmenos0.5puntosporcentuales2027–2050; SUPUESTOnuestro simetrico, noescenarioestimadoARISTO'},'advertencias':['Tabla179NOcontienepeatones. Su crecimiento usa proxyprivados declarado, manteniendo unidades de personas.','No hay estimacionpropia de ferry: extrapolacion del ritmodecarretera, noidentifica cambiosdeoferta, turismo, captacion oinduccion.','Tabla182PDF tiene aparente desalineacion de contraste entreclases/temporadas respectoExcel. Se conservan factoresExcel yrecalculados conPIB/beta, sinusarTabla182directamente.','Primavera originalesbase2025; Verano/Inviernobase2026 segunInforme p243. Rebaseunico2026 elimina doblecrecimiento.'],'archivos':{'PIB_anual_2026_2050.csv':len(gdp),'elasticidades_por_segmento.csv':len(elastic),'factores_propuestos_rebasados_2026.csv':len(factors),'factores_locales_ARISTO_2030_2040_2050.csv':len(local)}}
 result['celdas_beta']={'privados':'Ruta!D7','CAM':'Ruta!D11'}
 result['PIB_periodos']=[dict(desde=a,hasta=b,tasa_base=r,tasa_alto=r+.005,tasa_conservadora_SUPUESTO=r-.005,fuente='InformeARISTO Tabla181p241; ExcelPIB columnasB/C filas35–59') for a,b,r in [(2026,2026,.02),(2027,2027,.025),(2028,2030,.021),(2031,2035,.018),(2036,2040,.017),(2041,2045,.016),(2046,2050,.015)]]
 (OUT/'resumen_parametros.json').write_text(json.dumps(result,ensure_ascii=False,indent=2)+'\n');print(json.dumps(result,ensure_ascii=False,indent=2))
if __name__=='__main__':main()
