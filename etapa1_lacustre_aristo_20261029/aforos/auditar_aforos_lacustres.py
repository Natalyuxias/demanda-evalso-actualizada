"""Reproduce conteos lacustres desde ARISTO sin estimacion temporal ni modelos previos."""
from pathlib import Path
import argparse,hashlib,json,unicodedata,math
from collections import defaultdict
import pandas as pd
from openpyxl import load_workbook

def norm(s):return ''.join(c for c in unicodedata.normalize('NFD',str(s or '').lower()) if unicodedata.category(c)!='Mn')
def direction(origin):return 'Chile Chico → Puerto Ibáñez' if norm(origin)=='chile chico' else 'Puerto Ibáñez → Chile Chico'
def sha(p):return hashlib.sha256(Path(p).read_bytes()).hexdigest()
def main():
 ap=argparse.ArgumentParser();ap.add_argument('--anexos',type=Path,default=Path('.'));ap.add_argument('--microdatos',type=Path,default=Path('etapa1_lacustre_aristo_20261029/microdatos/final/microdatos_ferry_264.csv'));ap.add_argument('--salida',type=Path,default=Path('etapa1_lacustre_aristo_20261029/aforos'));args=ap.parse_args();out=args.salida;out.mkdir(parents=True,exist_ok=True)
 cp=args.anexos/'Cap 6/Mediciones de flujo/Consolidado Flujos.xlsx';bp=args.anexos/'Cap 6/BBDD EOD/Base de datos agregada.xlsx';cf=load_workbook(cp,read_only=True,data_only=False,keep_links=False);cv=load_workbook(cp,read_only=True,data_only=True,keep_links=False)
 originalsha={str(cp):sha(cp),str(bp):sha(bp)}
 # 0-based field indices; G and Z are deliberately different.
 definitions={'VL':('Vehículos livianos de cuatro ruedas',[7,8,9,10,19],None,'vehículos'),'CAM':('Camión o furgón de carga',[11,15,16,17,18],24,'vehículos'),'PEATON':('Peatón sin sumar bicicletas',[6],6,'personas_peatones'),'MOTO':('Motocicleta',[20],20,'motocicletas'),'BICI':('Bicicleta',[21],21,'bicicletas'),'BUS_TP':('Bus o taxibus transporte público',[12,13,14],23,'buses'),'VP':('Privados agrupados: VL + motos + bicicletas',[7,8,9,10,19,20,21],22,'unidades_privadas_heterogeneas'),'PEATON_MAS_BICI':('Agregado Z original: peatones más bicicletas',[6,21],25,'unidades_heterogeneas_no_usar_como_personas')}
 rows=[];raw=[]
 headers=next(cf['Transbordadores'].iter_rows(values_only=True))
 for rn,(fr,vr) in enumerate(zip(cf['Transbordadores'].iter_rows(values_only=True),cv['Transbordadores'].iter_rows(values_only=True)),1):
  if rn==1 or not ('ibanez' in norm(fr[3]) and 'chile chico' in norm(fr[3])):continue
  camp_original=str(fr[4]).strip();camp='Otoño' if camp_original=='Normal' else camp_original;dt=fr[5].date().isoformat();sense=direction(fr[1]);raw.append({'fila_excel':rn,'campana':camp,'fecha':dt,'sentido_cruce':sense,'original_celdas_JSON':json.dumps({f'{i+1}:{headers[i]}':str(v) if hasattr(v,'isoformat') else v for i,v in enumerate(fr)},ensure_ascii=False)})
  for seg,(label,cols,agg,unit) in definitions.items():
   vals=[fr[i] for i in cols];known=[v for v in vals if isinstance(v,(float,int))];calculated=sum(known) if known else None;value=vr[agg] if agg is not None else calculated
   blanks=[f'{i+1}:{headers[i]}' for i in cols if fr[i] is None]
   assert value is None or isinstance(value,(float,int)) and value>=0
   if agg is not None and len(known)==len(cols):assert value==calculated,(rn,seg,value,calculated)
   rows.append(dict(fila_excel=rn,hoja='Transbordadores',campana=camp,campana_original=camp_original,fecha=dt,sentido_cruce=sense,segmento=seg,descripcion=label,unidad=unit,conteo_reportado=value,suma_componentes_disponibles=calculated,componentes_blancos=';'.join(blanks),componentes_blancos_no_imputados=len(blanks),origen_total='Valor directo de celda'+str(agg+1) if agg is not None and not isinstance(fr[agg],str) else 'Agregado/formula original ARISTO; SUM ignora blancos' if agg is not None else 'Suma de clases físicas H:K,T; mismo criterio VL07, blancos advertidos',formula_original=fr[agg] if agg is not None else 'SUM(H:K,T)',fila_identifica_zarpe='No hay hora ni identificador de zarpe en esta hoja',estatus='Conteo reportado; componentes faltantes conservados' if blanks else 'Conteo reportado con componentes completos',fuente=str(cp)+'; Transbordadores fila '+str(rn)))
 cf.close();cv.close();assert len(raw)==10 and len(rows)==80
 counts=pd.DataFrame(rows);counts.to_csv(out/'conteos_normalizados.csv',index=False,encoding='utf-8-sig');pd.DataFrame(raw).to_csv(out/'filas_originales_conteos.csv',index=False,encoding='utf-8-sig')
 micro=pd.read_csv(args.microdatos,keep_default_na=False,encoding='utf-8-sig');assert len(micro)==264
 micro['sentido_std']=micro.sentido_cruce.map(lambda s:'Chile Chico → Puerto Ibáñez' if 'De Chile Chico' in s else 'Puerto Ibáñez → Chile Chico')
 types={'VL':['Auto','Camioneta','Vehículo privado'],'CAM':['Camión o furgón de carga'],'PEATON':['Pasajero (peatón)'],'MOTO':['Moto'],'BICI':['Bicicleta'],'BUS_TP':[],'VP':['Auto','Camioneta','Vehículo privado','Moto','Bicicleta'],'PEATON_MAS_BICI':[]}
 strata=[]
 for (camp,sense,seg),g in counts.groupby(['campana','sentido_cruce','segmento'],sort=True):
  sample=micro[micro.campana.eq(camp)&micro.sentido_std.eq(sense)&micro.tipo_usuario_BBDD2.isin(types[seg])];n=len(sample);complete=g.conteo_reportado.notna().all();F=float(g.conteo_reportado.sum()) if complete else None;partial=float(g.conteo_reportado.sum()) if g.conteo_reportado.notna().any() else None;nd=len(g);dates='|'.join(sorted(g.fecha.unique()));survey_dates='|'.join(sorted(sample['b1__Día'].unique())) if n else ''
  reasons=[]
  if not complete:reasons.append('conteo_de_clase_faltante_en_al_menos_una_fila')
  if not n:reasons.append('sin_entrevistas_de_clase_compatible')
  if seg in ['BUS_TP','PEATON_MAS_BICI']:reasons.append('sin_unidad_muestral_compatible_no_expandir')
  if seg=='VL' and camp=='Otoño':reasons.append('Vehiculo_privado_No_desagrega_VL_MOTO_BICI; coeficiente_VL_solo_condicional07')
  if F is not None and n>F:reasons.append('N_supera_conteo_no_es_muestra_subconjunto_de_marco_contado')
  if int(g.componentes_blancos_no_imputados.sum()):reasons.append('agregado_ARISTO_con_blancos; no_implica_cero_observado_en_componente')
  usable=F is not None and n and seg not in ['BUS_TP','PEATON_MAS_BICI']
  mapping='Coincidencia de clase' if not(camp=='Otoño' and seg=='VL') else 'Supuesto: Vehículo privado equivale a VL; preferir VP para compatibilidad agrupada'
  strata.append(dict(campana=camp,sentido_cruce=sense,segmento=seg,descripcion=g.descripcion.iloc[0],unidad=g.unidad.iloc[0],filas_medidas=nd,fechas_medidas=dates,fechas_EOD=survey_dates,fechas_EOD_incluidas_conteo=all(dt in set(g.fecha) for dt in sample['b1__Día']),N_entrevistas_validas_clase=n,F_total_filas_medidas=F,F_parcial_disponible_no_rellenado=partial,media_por_fila=F/nd if F is not None else None,coef_Ftotal_sobre_N=F/n if usable else None,coef_media_fila_sobre_N=F/nd/n if usable else None,coef_expansion_individual_certificado=False,media_por_zarpe_acreditada=False,horarios_encuesta_planificados_tabla78=2 if camp=='Primavera' else None,uso_principal_sugerido=(seg in (['VP','CAM','PEATON','BUS_TP'] if camp=='Otoño' else ['VL','CAM','PEATON','MOTO','BICI','BUS_TP'])),compatibilidad_muestral=mapping,filas_fuente='|'.join(g.fila_excel.astype(str)),problemas=';'.join(reasons),tipo_coeficiente='Transferencia proporcional de composicion OD sobre total de filas medidas; no ponderador de inclusion individual',balance_N_por_coef=F if usable else None,exposicion='Total de filas medidas en fechas listadas; no 24h, semana, estacion o año'))
 st=pd.DataFrame(strata);st.to_csv(out/'estratos_y_coeficientes.csv',index=False,encoding='utf-8-sig')
 daily=[]
 for (camp,sense,dt,seg),g in counts.groupby(['campana','sentido_cruce','fecha','segmento'],sort=True):
  sample=micro[micro.campana.eq(camp)&micro.sentido_std.eq(sense)&micro.tipo_usuario_BBDD2.isin(types[seg])&micro['b1__Día'].eq(dt)];n=len(sample);total=float(g.conteo_reportado.sum()) if g.conteo_reportado.notna().all() else None
  daily.append(dict(campana=camp,sentido_cruce=sense,fecha=dt,segmento=seg,unidad=g.unidad.iloc[0],filas_medidas=len(g),F_total_filas_fecha=total,N_entrevistas_fecha=n,coef_F_fecha_sobre_N_fecha=total/n if total is not None and n and seg not in ['BUS_TP','PEATON_MAS_BICI'] else None,filas_fuente='|'.join(g.fila_excel.astype(str)),N_supera_F=(n>total) if total is not None else None,componentes_blancos=int(g.componentes_blancos_no_imputados.sum()),observacion='Suma exclusiva de filas de misma fecha/sentido; no total24h acreditado. Coeficiente transfiere composición, no inclusion certificada.'))
 pd.DataFrame(daily).to_csv(out/'conteos_y_coeficientes_por_fecha.csv',index=False,encoding='utf-8-sig')

 # Contrast of original TD: daily column uses SUM of rows, not mean. Keep it separate.
 bf=load_workbook(bp,read_only=True,data_only=False,keep_links=False);bv=load_workbook(bp,read_only=True,data_only=True,keep_links=False);td=[]
 for rn in range(20,26):
  fr=[bf['TD Transb'].cell(rn,c).value for c in range(1,9)];vr=[bv['TD Transb'].cell(rn,c).value for c in range(1,9)];camp_original=str(fr[1]).strip();camp='Otoño' if camp_original=='Normal' else camp_original;sense='Chile Chico → Puerto Ibáñez' if 'De Chile Chico' in fr[2] else 'Puerto Ibáñez → Chile Chico'
  for seg,col in [('VP',3),('CAM',4)]:
   r=st[st.campana.eq(camp)&st.sentido_cruce.eq(sense)&st.segmento.eq(seg)].iloc[0];td.append(dict(fila_TD=rn,campana=camp,sentido_cruce=sense,segmento=seg,total_TD=vr[col],total_filas_medidas=r.F_total_filas_medidas,total_concilia=vr[col]==r.F_total_filas_medidas,frecuencia_original_F=vr[5],formula_diaria=fr[6 if seg=='VP' else 7],valor_diario_cache=vr[6 if seg=='VP' else 7],filas_medidas=r.filas_medidas,media_fila=r.media_por_fila,correccion_si_se_quisiera_media=(vr[6 if seg=='VP' else 7]/r.filas_medidas),nota='TD multiplica suma de filas por frecuencia9/7, sin dividir numero de filas; no equivale a flujo diario observado ni promedio por zarpe'))
 bf.close();bv.close();pd.DataFrame(td).to_csv(out/'auditoria_TD_Transb.csv',index=False,encoding='utf-8-sig')
 # Scenarios only: user-known ordinary weekly service, independent from counted exposure.
 scenarios=[]
 for r in st.to_dict('records'):
  for label,freq in [('Escenario ordinario alta dic-abr',11),('Escenario ordinario baja may-nov',8)]:
   scenarios.append({k:r[k] for k in ['campana','sentido_cruce','segmento','media_por_fila','N_entrevistas_validas_clase','filas_fuente']}|dict(escenario=label,frecuencia_por_sentido_semana=freq,referencia_diaria_condicional=r['media_por_fila']*freq/7 if pd.notna(r['media_por_fila']) else None,tipo='NO observado; media/fila asimilada a zarpe representativo y multiplicada por frecuencia ordinaria de usuario; incluye incertidumbre sobre zarpe y representatividad',carga_peligrosa='No identificable por fila; no se puede separar embarque especial de ordinario',fuente_frecuencia='Confirmacion de usuario documentada en config_matrices_od.json; no conteo de zarpes de este archivo'))
 pd.DataFrame(scenarios).to_csv(out/'escenarios_frecuencia_NO_observados.csv',index=False,encoding='utf-8-sig')
 # Explicit printed report values: current source discrepancy never silently fixed.
 pdfrows=[('Primavera','Chile Chico → Puerto Ibáñez',54,11,6,17,188,107),('Primavera','Puerto Ibáñez → Chile Chico',69,22,3,25,188,107),('Verano','Chile Chico → Puerto Ibáñez',194,57,2,59,207,138),('Verano','Puerto Ibáñez → Chile Chico',120,46,9,55,207,138),('Otoño','Chile Chico → Puerto Ibáñez',48,16,4,20,223,169),('Otoño','Puerto Ibáñez → Chile Chico',62,16,3,19,223,169)]
 contrasts=[]
 for camp,sense,peat,priv,cam,total,page,tab in pdfrows:
  for seg,pdfval in [('PEATON',peat),('VP',priv),('CAM',cam)]:
   r=st[st.campana.eq(camp)&st.sentido_cruce.eq(sense)&st.segmento.eq(seg)].iloc[0];contrasts.append(dict(campana=camp,sentido_cruce=sense,segmento=seg,total_Excel=r.F_total_filas_medidas,total_Informe=pdfval,diferencia_Excel_menos_Informe=r.F_total_filas_medidas-pdfval,pagina_impresa=page,tabla=tab,nota='Normal/verano informe dice bicicletas sumadas a peatones; Excel G peatones, Z G+BICI. Normal VP y CAM discrepan; no se escoge automaticamente una cifra alternativa',fuente='Informe ARISTO.pdf tabla '+str(tab)+' p'+str(page)))
 pd.DataFrame(contrasts).to_csv(out/'contraste_Excel_Informe.csv',index=False,encoding='utf-8-sig')
 assert all(sha(Path(path))==value for path,value in originalsha.items())
 summary={'estado':'correcto','hashes_originales':originalsha,'hash_microdatos':sha(args.microdatos),'filas_cruce':10,'filas_fuente':[r['fila_excel'] for r in raw],'estratos':len(st),'N_microdatos':264,'campanas':['Primavera','Verano','Otoño'],'zarpes_identificados_individualmente_en_conteos':False,'Primavera_evidencia': 'Tabla78 p173 planifica encuesta a bordo en dos horarios por sentido08-11; no certifica identificadores individuales de zarpe. Tabla107 pp187–188 totaldia/sentido. No hay enlace hora-fila para cada una de las2filas.','Verano_Normal_evidencia':'Planificacion encuestas tablas109 p191 y140 p209 porfecha; conteos nohora/IDzarpe. Media porfila, no certificar porzarpe.','N_supera_F':st[st.N_entrevistas_validas_clase.gt(st.F_total_filas_medidas.fillna(float('inf')))][['campana','sentido_cruce','segmento','N_entrevistas_validas_clase','F_total_filas_medidas']].to_dict('records'),'clases_sin_N':st[st.N_entrevistas_validas_clase.eq(0)][['campana','sentido_cruce','segmento','F_total_filas_medidas']].fillna('').to_dict('records'),'errores_o_riesgos_07':['TD Transb usa9semanales fijo y sumasfilas, no media;07corrige mediafilas pero todavia es supuesto fila=zarpe fueraPrimavera.','07convierte automaticamente a referencia diaria por11/8; aqui queda solo escenario y el principal usa totalesmedidos.','Normal Vehiculo privado enEOD no desagregaVL/MOTO/BICI;07lo asignaVL4ruedas,supuesto no demostrado.','07 usa CAMExcel5/4 Normal sin explicitar discrepancia con Informe4/3.','N>Fprimavera no es expansion de muestra-subconjunto;07suponecomposicionrepresentativa, no admite validacion probabilistica.','07 sumaVL saltandoblancosNormal; mantener procedimientoARISTO y flags sin declarar ceros.'],'carga_peligrosa':'No hay identificador de carga peligrosa/especial ni indicador ordinario/especial en columnas de conteos. No puede estimarse ni excluirse separadamente.','recomendacion':'Principal: total de filas medidas y composicion OD por FECHA/sentido/usuario, sin transferir automaticamente muestra de otro dia. Coeficientes agrupados por campaña quedan como contraste de perfil transferido. No afirmar24h ni totalestacion. Verano01febPIaCC CAM5medidos quedan sinOD; enero27solo4expandidos desdeN2. Presentar media/fila y escenariosfrecuencia11/8 aparte, sin convertirlos en observado.'}
 (out/'resumen_aforos.json').write_text(json.dumps(summary,ensure_ascii=False,indent=2)+'\n')
 print(json.dumps({'estado':'correcto','filas':10,'estratos':len(st),'N_supera_F':summary['N_supera_F']},ensure_ascii=False,indent=2))
if __name__=='__main__':main()
