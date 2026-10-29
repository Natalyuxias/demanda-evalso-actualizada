"""Segunda revisión de 20 grupos, solo propuestas. No modifica homologación ni modelos.
Desde Anexos: python homologacion_geografica_etapa1/segunda_revision/scripts/revisar_prioritarios.py
"""
from pathlib import Path
import importlib.util,collections,json,hashlib,csv,re
import pandas as pd
import geopandas as gp
from shapely.geometry import Point
from openpyxl import load_workbook,Workbook
from openpyxl.styles import Font,PatternFill,Alignment
from openpyxl.worksheet.datavalidation import DataValidation

OUT=Path(__file__).resolve().parents[1];PRE=OUT.parent;BASE=PRE.parent
SEP=Path('/home/nataly/Escritorio/demanda_evalso')
spec=importlib.util.spec_from_file_location('homologacion_lectura',PRE/'scripts/homologar_geografia.py');m=importlib.util.module_from_spec(spec);spec.loader.exec_module(m)

def read(p):return list(csv.DictReader(Path(p).open(encoding='utf-8-sig',newline='')))
def sha(p):return hashlib.sha256(Path(p).read_bytes()).hexdigest()
def txt(v):return '' if v is None else str(v)
def boolean(v):return str(v) in ['True','true','1']
def write(name,rows):
 rows=list(rows);fields=list(dict.fromkeys(k for r in rows for k in r)) or ['Sin casos']
 with (OUT/(name+'.csv')).open('w',encoding='utf-8-sig',newline='') as f:
  w=csv.DictWriter(f,fieldnames=fields);w.writeheader();w.writerows(rows)

def main():
 protected=list(PRE.glob('*.*'))+[PRE/'scripts/homologar_geografia.py',PRE/'scripts/verificar_reglas.py']
 protected+= [Path(r['archivo']) for r in json.loads((PRE/'huellas_antes.json').read_text())]
 protected+=[SEP/'analisis_eod/tablas/catalogo_puntos.csv']
 protected+=list((BASE/'SIG/Coberturas/MINSAL').glob('Establecimientos_Salud.*'))
 protected=list(dict.fromkeys(p.resolve() for p in protected if p.exists()))
 hashes=[dict(archivo=str(p),sha256=sha(p)) for p in protected]
 (OUT/'huellas_antes.json').write_text(json.dumps(hashes,ensure_ascii=False,indent=2))
 w=load_workbook(PRE/'homologacion_CITRA_ARISTO.xlsx',read_only=True,data_only=True)
 it=w['Revisión prioritaria'].values;hh=next(it);groups=[{k:txt(v) for k,v in zip(hh,row)} for row in it if row[0]];w.close()
 assert len(groups)==20 and len({r['ID_caso'] for r in groups})==20
 bygroup={r['ID_caso']:r for r in groups}
 ends=read(PRE/'extremos_homologados.csv');previous={(r['id_encuesta'],r['extremo']):r for r in ends}
 selected=[r for r in ends if r['id_caso_geografico'] in bygroup]
 assert len(selected)==37 and len({r['id_encuesta'] for r in selected})==36
 assert all(boolean(r['relevante_revision']) and r['estado']=='Pendiente' for r in selected)
 raw,rawheaders=m.load_raw();rawmap={r['id_encuesta']:r for r in raw}
 records={r['id_encuesta']:r for r in read(PRE/'correspondencias_por_registro.csv')}
 legacy={r['id_encuesta']:r for r in read(SEP/'eod_filtrada/tablas/base_completa_filtrada.csv')}
 ext_raw=read(SEP/'analisis_eod/tablas/eod_recuperada_local.csv')
 controls={r['control']:r for r in read(SEP/'analisis_eod/tablas/catalogo_puntos.csv')}
 h=m.Homologador();health=gp.read_file(BASE/'SIG/Coberturas/MINSAL/Establecimientos_Salud.shp').to_crs(4326)
 refs=[];support=[]
 def ref(id_,name,lat,lon,source,kind):
  loc=h.intersect(float(lat),float(lon));assert loc['cut'] and loc['zona'] and not loc['conflicto_poligonos']
  r=dict(id_referencia=id_,nombre=name,lat_referencia=float(lat),lon_referencia=float(lon),CUT_referencia=loc['cut'],comuna_referencia=loc['comuna'],zona_referencia=loc['zona'],fuente=source,naturaleza=kind,advertencia='Punto de referencia; nunca sustituye la coordenada de una entrevista')
  refs.append(r);return r
 for name in ['PUERTO RIO TRANQUILO','PUERTO GUADAL','MALLIN GRANDE','COYHAIQUE','PUERTO CISNES','PUERTO AYSEN','EL BLANCO','LA JUNTA']:
  r=h.refs[name];ref(name,name,r['lat'],r['lon'],'matriz_od_proyecto/tablas/referencias_localidades.csv; '+r['fuente'],'KMZ' if r['fuerte'] else 'Referencia calculada de otras observaciones; no punto individual')
 for term in ['Posta de Salud Rural El Blanco','Centro de Salud Familiar La Junta','Hospital Dr. Jorge Ibar Bruce (Cisnes)','Hospital Dr. Leopoldo Ortega Rodríguez (Chile Chico)','Posta de Salud Rural Río Tranquilo']:
  q=health[health.NOMBRE.eq(term)];assert len(q)==1;r=q.iloc[0]
  rr=ref('MINSAL_'+str(int(r.COD_VIG)),r.NOMBRE,r.geometry.y,r.geometry.x,'SIG/Coberturas/MINSAL/Establecimientos_Salud.shp; COD_VIG='+str(int(r.COD_VIG))+'; comuna declarada MINSAL='+str(r.NOM_COM)+'; COD_COM='+str(r.COD_COM),'Establecimiento público de referencia; no identifica por sí solo el establecimiento visitado')
  assert rr['CUT_referencia']==str(r.COD_COM)
 refmap={r['id_referencia']:r for r in refs}
 # Validar pares que no deben confundirse; solo catálogo, sin cambios de otros registros.
 assert refmap['PUERTO RIO TRANQUILO']['CUT_referencia']=='11402'
 assert refmap['MALLIN GRANDE']['CUT_referencia']==refmap['PUERTO GUADAL']['CUT_referencia']=='11401'
 assert refmap['PUERTO CISNES']['CUT_referencia']=='11202' and refmap['PUERTO AYSEN']['CUT_referencia']=='11201'
 # Recuperar evidencia geocodificada de otros registros, siempre identificada como auxiliar.
 def peers(label,matcher):
  found=[]
  for r in ext_raw:
   for side in ['origen','destino']:
    if not matcher(r[side+'_original'],r[side+'_busqueda'],r[side+'_direccion']):continue
    lat,lon=m.num(r[side+'_lat']),m.num(r[side+'_lon'])
    if lat is None or lon is None:continue
    hit=h.intersect(lat,lon)
    rid=('V' if r['campana_archivo']=='febrero' else 'O')+f"{int(r['fila_excel']):05d}"
    rr=dict(conjunto_evidencia=label,id_registro_auxiliar=rid,extremo_auxiliar=side,archivo='EOD Aysen Abril 2026.xlsx' if r['campana_archivo']=='abril' else 'EOD Aysen Febrero 2026 vf.xlsx',fila_excel=r['fila_excel'],texto_auxiliar=r[side+'_original'],lat_auxiliar=lat,lon_auxiliar=lon,CUT_auxiliar=hit['cut'],zona_auxiliar=hit['zona'],uso='Corrobora nombre/comuna; NO transfiere coordenada a registros sin coordenada')
    found.append(rr);support.append(rr)
  return found
 losolmos=peers('LOS_OLMOS_PUERTO_GUADAL',lambda a,s,d:'LOS OLMOS' in m.norm(d) and 'PUERTO GUADAL' in m.norm(s))
 blanco=peers('EL_BLANCO',lambda a,s,d:m.canonical(s.split(',')[0])=='EL BLANCO')
 junta=peers('LA_JUNTA',lambda a,s,d:m.canonical(s.split(',')[0]) in ['LA JUNTA','LA JUNTA - CARRETERA AUSTRAL'])
 bilbao=peers('CALLES_BILBAO_QUINTAS_COYHAIQUE',lambda a,s,d:bool(re.search(r'BILBAO|LAS QUINTAS',m.norm(s+' '+d))) and bool(re.search(r'COYHAIQUE|COIHAIQUE',m.norm(s+' '+d))))
 peer_quality=[]
 for label,rows,cut in [('Los Olmos',losolmos,'11401'),('El Blanco',blanco,'11101'),('La Junta',junta,'11202'),('Bilbao/Quintas',bilbao,'11101')]:
  valid=[r for r in rows if r['CUT_auxiliar']==cut]
  invalid=[r for r in rows if not r['CUT_auxiliar']]
  contradictory=[r for r in rows if r['CUT_auxiliar'] and r['CUT_auxiliar']!=cut]
  assert valid and not contradictory,(label,collections.Counter(r['CUT_auxiliar'] for r in rows))
  for r in invalid:r['uso']='Sin comuna verificable por geometría; NO usado como apoyo positivo ni corregido'
  peer_quality.append(dict(conjunto=label,referencias_examinadas=len(rows),referencias_concordantes=len(valid),sin_comuna_verificable=len(invalid),en_otra_comuna=len(contradictory),CUT_esperado=cut))
 # Propuestas nominales específicas del lote; se validan literales, comunas y referencias en cada fila.
 nominal={
 '1ef5538938a42f':('11402','PUERTO RIO TRANQUILO','Topónimo explícito Río Tranquilo/Puerto Río Tranquilo; KMZ y dos polígonos lo ubican en Río Ibáñez. La comuna genérica Chile Chico contradice ese topónimo. Sin coordenada individual: propuesta nominal comunal, sin fijar punto de acceso.','KMZ_PUERTO_RIO_TRANQUILO'),
 'b53e153a78cac7':('11101','COYHAIQUE','El literal completo del extremo es COYHAIQUE; referencia KMZ y polígonos sitúan ese topónimo en comuna Coyhaique. Chile Chico es el campo comunal contradictorio. El control no decide cuál texto es correcto. Propuesta nominal, sin coordenada individual.','KMZ_COYHAIQUE'),
 '4f07c7789b87c9':('11101','MINERA EL TOQUI','La entidad minera homónima consta en la ficha oficial SMA/SNIFA 2509, comuna Coyhaique. Los registros carecen de coordenada individual. No se usa como referencia la calle Toqui de otra encuesta ni se localiza el acceso industrial.','PUB_TOQUI'),
 'ab563f0157136d':('11202','PUERTO CISNES','El literal es PUERTO CISNES; su referencia KMZ y el hospital MINSAL de Cisnes corroboran comuna Cisnes. Puerto Aysén se conserva como comuna declarada contradictoria, no como alias. Propuesta nominal sin coordenada individual.','KMZ_PUERTO_CISNES;MINSAL_125104'),
 '4350bf03bea926':('11401','PUERTO GUADAL','Contiene el topónimo íntegro PUERTO GUADAL y el sector LOS OLMOS; comuna declarada Chile Chico concordante. Otra encuesta identifica Los Olmos en Puerto Guadal con coordenada en CUT11401. Esa coordenada es auxiliar, no la del registro.','KMZ_PUERTO_GUADAL;LOS_OLMOS_PUERTO_GUADAL'),
 '320b48c8a2a850':('11101','','Comuna declarada Coyhaique y referencias geocodificadas de ambas calles en dicha comuna. Permite propuesta administrativa; no prueba la intersección ni fija el acceso del viaje.','CALLES_BILBAO_QUINTAS_COYHAIQUE'),
 'bed034143b896e':('11101','COYHAIQUE','COYHAIQUE CENTRO contiene un topónimo inequívoco y la comuna declarada coincide. Referencia KMZ y polígonos corroboran comuna. Centro no se convierte en coordenada.','KMZ_COYHAIQUE'),
 '136c7232a098eb':('11101','EL BLANCO','Localidad El Blanco corroborada por Posta de Salud Rural MINSAL COD125413, referencia calculada y otras observaciones geocodificadas. El dato normalizado previo Río Ibáñez es derivado y no se adopta frente a esta evidencia. No hay coordenada individual.','MINSAL_125413;EL_BLANCO'),
 '569e6ff1d5e211':('11201','PUERTO AYSEN','EL LEON PUERTO AYSEN contiene la localidad completa y coincide con comuna declarada Puerto Aysén. El predio El León de Mallín Grande es otro topónimo: no se usa como alias ni referencia de este registro. Se propone solo comuna/zona, sin ubicar predio.','KMZ_PUERTO_AYSEN'),
 'e72ccd1e735918':('11202','LA JUNTA','La localidad La Junta está corroborada por CESFAM MINSAL COD125442, referencia calculada y otras observaciones geocodificadas en comuna Cisnes. No se imputa una coordenada individual ni se valida una ruta.','MINSAL_125442;LA_JUNTA')}
 manual_info={
 'a2e1ed14613a40':('Alemania puede referir al extremo internacional o a residencia/nacionalidad informada; no existe coordenada ni topónimo chileno verificable. El campo normalizado Río Ibáñez no constituye evidencia independiente.','¿ALEMANIA es el origen completo del viaje, o una referencia de residencia/nacionalidad? Si el viaje empezó en Chile, indicar localidad.','',''),
 '8311abd0cf4b45':('La comuna declarada es Coyhaique y SERNATUR registra Cabaña Ruka Aike allí; CABAÑA RUKA es un nombre incompleto y no acredita identidad del establecimiento.','Confirmar si el extremo corresponde a una cabaña en comuna Coyhaique; no hace falta precisar domicilio para esta matriz.','11101','PUB_RUKA'),
 '7ff509d8d2a94f':('El literal CHILE CHICO (COYHAIQUE) contiene dos lugares distintos; se desconoce el significado del paréntesis. No hay coordenadas para dirimir.','¿El inicio del viaje es Chile Chico o Coyhaique? Aclarar qué indica el paréntesis.','11401 | 11101',''),
 '1e10245fa9054d':('CHILE CHICO (PUNTA ARENAS) puede mezclar inicio del trayecto y residencia/otro tramo; el registro no contiene coordenadas.','¿El origen completo es Chile Chico o Punta Arenas? Aclarar qué indica el paréntesis.','11401 | 12101',''),
 'a232e3cbd3c3ed':('La ficha oficial SMA registra Coca Cola Andina Coyhaique, compatible con comuna declarada, pero COCA COLA no identifica una sucursal/instalación inequívoca. Otras entrevistas con el mismo texto tampoco tienen coordenadas.','Confirmar si COCA COLA designa un lugar en comuna Coyhaique o una instalación de otra comuna.','11101','PUB_COCA'),
 '8eb299aff04ecd':('HOSPITAL es genérico. MINSAL contiene un hospital de Chile Chico, compatible con comuna declarada, y hospitales en otras comunas; falta identificar cuál fue el extremo.','Confirmar la localidad/comuna del hospital; no hace falta dirección exacta.','11401','MINSAL_125102'),
 'f4d811ed81d28a':('IBAÑEZ no distingue comuna Río Ibáñez, localidad Puerto Ibáñez ni otra referencia; no tiene coordenadas. El control de Puerto Río Tranquilo no es el origen del viaje.','¿El extremo corresponde a Puerto Ibáñez/comuna Río Ibáñez u otra referencia? Indicar localidad o comuna.','11402','')}
 proposals=[]
 for old in selected:
  rid,side=old['id_encuesta'],old['extremo'];src=rawmap[rid];gid=old['id_caso_geografico'];baserec=records[rid];leg=legacy[rid]
  assert src[side+'_original']==old['literal_original']
  assert src['hash_registro_original']==baserec['hash_registro_original']
  lat,lon=m.num(old['lat_original']),m.num(old['lon_original']);haspoint=lat is not None and lon is not None
  loc=h.intersect(lat,lon) if haspoint else None
  ctrl=controls[leg['control']]
  q=dict(ID_caso=gid,id_encuesta=rid,extremo=side,archivo=old['archivo'],hoja='EOD',fila_excel=old['fila_excel'],hash_registro_original=src['hash_registro_original'],literal_original=old['literal_original'],localidad_original=old['localidad_original'],busqueda_original=old['busqueda_original'],direccion_original=old['direccion_original'],lat_original=old['lat_original'],lon_original=old['lon_original'],coordenada_individual_disponible=haspoint,comuna_declarada_original=old['comuna_declarada_original'],comuna_corregida_fuente=next((r.get('comuna_'+side+'_corregida_fuente','') for r in ext_raw if str(r['fila_excel'])==old['fila_excel'] and r['campana_archivo']==old['campana_archivo']),''),comuna_normalizada_fuente=next((r.get('comuna_'+side+'_normalizada_fuente_no_validada','') for r in ext_raw if str(r['fila_excel'])==old['fila_excel'] and r['campana_archivo']==old['campana_archivo']),''),estado_homologacion_anterior=old['estado'],comuna_anterior=old['comuna_homologada'],CUT_anterior=old['cut_comuna'],zona_anterior=old['zona_modelacion_aristo'],motivo_anterior=old['motivo'],previo_estado_filtro=old['previo_estado_filtro'],previa_decision_area=old['previo_decision_area'],previa_decision_lista=old['previo_decision_lista'],control=leg['control'],localidad_control=ctrl['localidad'],lat_control=ctrl['latitud'],lon_control=ctrl['longitud'],sentido_control=leg['sentido_interceptacion'],OD_original_contraparte=src[('destino' if side=='origen' else 'origen')+'_original'],uso_control='Solo contexto regional del recorrido; no identifica extremo ni confirma ruta. No se elige comuna por cercanía al control.',CUT_poligono_original=loc['cut'] if loc else '',zona_poligono_original=loc['zona'] if loc else '',margen_limite_m=loc['dist_limite_m'] if loc else '',localidad_propuesta='',comuna_propuesta='',CUT_propuesto='',zona_propuesta='',estado_revision='',nivel_evidencia='',resuelto_propuesto=False,requiere_manual=False,fuentes_evidencia='',justificacion='',pregunta_minima='',candidato_manual='',aplicado=False,limitacion='Propuesta de comuna/CUT/zona; no confirma posición exacta, acceso, ruta, ni demanda física única.')
  if gid=='f8bf21e49bf88c':
   rr=refmap['PUERTO RIO TRANQUILO'];d=Point(*h.proj.transform(lon,lat)).distance(Point(*h.proj.transform(rr['lon_referencia'],rr['lat_referencia']))) if haspoint else None
   q['distancia_referencia_m']=round(d,1) if d is not None else ''
   official=refmap['MINSAL_125412']
   ok=haspoint and loc['cut']=='11402' and loc['zona']==40 and not loc['conflicto_poligonos'] and d<2500 and official['CUT_referencia']=='11402' and m.canonical(old['busqueda_original'].split(',')[0])=='PUERTO RIO TRANQUILO'
   q['alerta_precision_limite']='Proximidad al límite de '+str(loc['dist_limite_m'])+' m: se mantiene alerta de precisión; pertenencia del topónimo verificada en MINSAL, KMZ y ambos polígonos. No se mueve el punto.'
   if ok:
    q.update(CUT_propuesto='11402',localidad_propuesta='PUERTO RIO TRANQUILO',estado_revision='Resoluble con coordenada individual',nivel_evidencia='A: coordenada + topónimo + dos polígonos + referencia oficial MINSAL; alerta de límite',resuelto_propuesto=True,fuentes_evidencia='EOD original; KMZ Puerto Río Tranquilo; MINSAL_125412 (Río Ibáñez, CUT11402); IDE/COMUNAS; ARISTO/Zonificación',justificacion='Punto original dentro de Río Ibáñez/zona40 y próximo a referencia Puerto Río Tranquilo. El topónimo específico concuerda con MINSAL y KMZ; Chile Chico del texto administrativo es contradictorio. Se propone solo su pertenencia comunal corroborada, conservando alerta por cercanía al límite y coordenada aproximada sin desplazarla. No es una excepción aplicable a puntos anónimos cerca del límite.')
   else:q.update(requiere_manual=True,estado_revision='Requiere revisión manual',justificacion='No se cumplieron todas las comprobaciones individuales de geometría y referencia.')
  elif gid=='416608a103a768':
   city=h.refs['CHILE CHICO'];d=Point(*h.proj.transform(lon,lat)).distance(Point(*h.proj.transform(city['lon'],city['lat']))) if haspoint else None
   q['distancia_referencia_m']=round(d,1) if d is not None else ''
   ok=haspoint and loc['cut']=='11401' and loc['zona']==39 and not loc['conflicto_poligonos'] and loc['dist_limite_m']>250 and d<5000
   if ok:q.update(CUT_propuesto='11401',localidad_propuesta='CHILE CHICO',estado_revision='Resoluble con coordenada individual',nivel_evidencia='A: coordenada + búsqueda + referencia urbana + dos polígonos',resuelto_propuesto=True,fuentes_evidencia='EOD original; KMZ Chile Chico; IDE/COMUNAS; ARISTO/Zonificación',justificacion='Coordenada y búsqueda indican Chile Chico. Aysén aparece como componente postal/regional dentro de una dirección que también indica Chile Chico; no demuestra un extremo en Puerto Aysén. El primer parser produjo un conflicto nominal que estas evidencias permiten resolver.')
   else:q.update(requiere_manual=True,estado_revision='Requiere revisión manual',justificacion='Las comprobaciones de punto y ciudad no son suficientes.')
  elif gid=='7f34d075805c9c':
   rr=refmap['PUERTO GUADAL'];d=Point(*h.proj.transform(lon,lat)).distance(Point(*h.proj.transform(rr['lon_referencia'],rr['lat_referencia']))) if haspoint else None
   q['distancia_referencia_m']=round(d,1) if d is not None else ''
   ok=haspoint and loc['cut']=='11401' and loc['zona']==39 and not loc['conflicto_poligonos'] and loc['dist_limite_m']>250 and d<2500
   if ok:q.update(CUT_propuesto='11401',estado_revision='Resoluble con coordenada individual',nivel_evidencia='A administrativa: punto + dos polígonos + referencia próxima',resuelto_propuesto=True,fuentes_evidencia='EOD original; KMZ Puerto Guadal; IDE/COMUNAS; ARISTO/Zonificación',justificacion='JSON conserva coordenada válida en comuna Chile Chico/zona39, próxima a Puerto Guadal, aunque no tiene búsqueda/dirección. Se propone solo agregación administrativa según punto registrado, dejando localidad sin asignar. No se sustituye el punto ni se asegura que sea el lugar exacto del viaje.')
   else:q.update(requiere_manual=True,estado_revision='Requiere revisión manual',justificacion='Punto sin texto y sin geometría/referencia administrativa suficiente.')
  elif gid in nominal:
   cut,name,why,source=nominal[gid]
   # Nunca tratar una referencia de otro registro como coordenada de este.
   assert not haspoint,(rid,'Regla nominal revisable: ahora sí hay coordenada individual')
   n=m.norm(old['localidad_original'])
   if gid=='1ef5538938a42f':assert m.canonical(n)=='PUERTO RIO TRANQUILO'
   if gid=='b53e153a78cac7':assert m.canonical(n)=='COYHAIQUE'
   if gid=='ab563f0157136d':assert m.canonical(n)=='PUERTO CISNES'
   if gid=='4350bf03bea926':assert 'PUERTO GUADAL' in n and 'LOS OLMOS' in n and m.canonical(old['comuna_declarada_original'])=='CHILE CHICO'
   if gid=='4f07c7789b87c9':assert n=='MINERA TOQUI'
   q.update(CUT_propuesto=cut,localidad_propuesta=name,estado_revision='Resoluble nominalmente a escala comunal',nivel_evidencia='B: propuesta nominal sustentada; SIN coordenada individual verificable',resuelto_propuesto=True,fuentes_evidencia=source+'; IDE/COMUNAS; ARISTO/Zonificación; EOD original',justificacion=why)
  else:
   reason,question,candidate,source=manual_info[gid]
   q.update(estado_revision='Requiere revisión manual mínima',nivel_evidencia='C: candidato o contradicción no resuelta',requiere_manual=True,justificacion=reason,pregunta_minima=question,candidato_manual=candidate,fuentes_evidencia=source+'; EOD original')
  if q['CUT_propuesto']:
   cut=q['CUT_propuesto'];q['comuna_propuesta']=h.bycut[cut]['comuna'];q['zona_propuesta']=h.cut_zone[cut]
  proposals.append(q)
 pmap={(r['id_encuesta'],r['extremo']):r for r in proposals}
 assert len(pmap)==37
 impacts=[]
 for rid in sorted({r['id_encuesta'] for r in proposals}):
  rec=dict(id_encuesta=rid,estado_area_previo=legacy[rid]['estado_filtro'])
  complete={}
  for scenario in ['anterior','solo_A_coordenadas','A_mas_B_nominales']:
   resolved=[]
   for side in ['origen','destino']:
    old=previous[rid,side];p=pmap.get((rid,side));ok=old['estado']=='Homologado comuna y zona'
    if scenario!='anterior' and p and p['resuelto_propuesto']:
     ok=ok or scenario=='A_mas_B_nominales' or p['estado_revision']=='Resoluble con coordenada individual'
    resolved.append(ok)
   complete[scenario]=all(resolved);rec['OD_completo_'+scenario]=all(resolved)
  rec['adicional_completo_A']=not complete['anterior'] and complete['solo_A_coordenadas']
  rec['adicional_completo_A_B']=not complete['anterior'] and complete['A_mas_B_nominales']
  rec['pendiente_fuera_de_20']=' | '.join(side+': '+previous[rid,side]['localidad_original'] for side in ['origen','destino'] if previous[rid,side]['estado']!='Homologado comuna y zona' and (rid,side) not in pmap)
  rec['pendiente_manual_en_20']=' | '.join(side for side in ['origen','destino'] if (rid,side) in pmap and pmap[rid,side]['requiere_manual'])
  impacts.append(rec)
 summaries=[];manual=[]
 for g in groups:
  pp=[p for p in proposals if p['ID_caso']==g['ID_caso']];resolved=sum(p['resuelto_propuesto'] for p in pp);need=sum(p['requiere_manual'] for p in pp)
  summaries.append(dict(ID_caso=g['ID_caso'],toponimo=g['toponimo_o_busqueda'],extremos_en_grupo=len(pp),resueltos_A=sum(p['estado_revision']=='Resoluble con coordenada individual' for p in pp),resueltos_B=sum(p['estado_revision']=='Resoluble nominalmente a escala comunal' for p in pp),pendientes_manual=need,clasificacion='Propuesta resoluble' if not need else 'Requiere decisión mínima',detalle=' | '.join(sorted({p['justificacion'] for p in pp})),comunas_propuestas=' | '.join(sorted({p['comuna_propuesta'] for p in pp if p['comuna_propuesta']})),decision_geografica_anterior=g['Decisión geográfica'],CUT_usuario_anterior=g['CUT propuesto usuario'],zona_usuario_anterior=g['Zona propuesta usuario'],evidencia_usuario_anterior=g['Evidencia o comentario usuario']))
  if need:
   p=next(p for p in pp if p['requiere_manual']);manual.append(dict(ID_caso=g['ID_caso'],toponimo=g['toponimo_o_busqueda'],extremos=need,registros=' | '.join(sorted({p['id_encuesta'] for p in pp if p['requiere_manual']})),motivo=p['justificacion'],pregunta_minima=p['pregunta_minima'],candidato_no_asignado=p['candidato_manual'],fuentes_consultadas=p['fuentes_evidencia'],respuesta_nueva='',evidencia_nueva='',decision_anterior_intacta=g['Decisión geográfica']))
 relevant=[r for r in records.values() if boolean(r['relevante_revision'])]
 assert len(relevant)==316
 oldends=sum(r['estado']=='Homologado comuna y zona' for r in ends if boolean(r['relevante_revision']))
 oldfull=sum(boolean(r['ambos_extremos_completos']) for r in relevant)
 stats=[]
 for label,accept in [('Referencia actual',lambda p:False),('Propuesta A: coordenadas individuales',lambda p:p['estado_revision']=='Resoluble con coordenada individual'),('Propuesta A+B: incluye respaldo nominal sin coordenada',lambda p:p['resuelto_propuesto'])]:
  pe=[p for p in proposals if accept(p)];extra_full=sum(r['adicional_completo_A'] if label.startswith('Propuesta A:') else r['adicional_completo_A_B'] if label.startswith('Propuesta A+B') else False for r in impacts)
  stats.append(dict(escenario=label,registros_relevantes=316,extremos_relevantes=632,extremos_adicionales_propuestos=len(pe),extremos_completos_si_se_acepta=oldends+len(pe),cobertura_extremos_pct=round(100*(oldends+len(pe))/632,2),registros_OD_completos_adicionales=extra_full,registros_OD_completos_si_se_acepta=oldfull+extra_full,cobertura_registros_pct=round(100*(oldfull+extra_full)/316,2),estado='Actual sin cambios' if not pe else 'Condicional a adoptar propuesta; NO aplicado',unidad='Entrevistas/registros de interceptación; no viajes únicos ni vehículos expandidos'))
 assert all(not p['CUT_propuesto'] and not p['zona_propuesta'] for p in proposals if p['requiere_manual'])
 checks={'20_grupos_hoja_vigente':len(groups)==20,'37_extremos_sin_perder_duplicar':len(proposals)==37 and len(pmap)==37,'36_registros_distintos':len(impacts)==36,'solo_registros_relevantes':all(boolean(previous[p['id_encuesta'],p['extremo']]['relevante_revision']) for p in proposals),'literales_y_coordenadas_originales_intactos':all(all(p[k]==previous[p['id_encuesta'],p['extremo']][k] for k in ['literal_original','lat_original','lon_original','comuna_declarada_original']) for p in proposals),'hash_por_registro_contra_original':all(rawmap[p['id_encuesta']]['hash_registro_original']==records[p['id_encuesta']]['hash_registro_original'] for p in proposals),'decisiones_previas_copiadas':all(s['decision_geografica_anterior']==bygroup[s['ID_caso']]['Decisión geográfica'] for s in summaries),'manuales_sin_asignacion_propuesta':all(not p['CUT_propuesto'] for p in proposals if p['requiere_manual']),'CUT_zona_catalogo_correcto':all(h.cut_zone[p['CUT_propuesto']]==p['zona_propuesta'] for p in proposals if p['CUT_propuesto']),'A_geometria_y_limite_o_respaldo_oficial_verificados':all(p['CUT_poligono_original']==p['CUT_propuesto'] and p['zona_poligono_original']==p['zona_propuesta'] and (p['margen_limite_m']>250 or (p['ID_caso']=='f8bf21e49bf88c' and 'MINSAL_125412' in p['fuentes_evidencia'] and refmap['MINSAL_125412']['CUT_referencia']=='11402')) for p in proposals if p['estado_revision']=='Resoluble con coordenada individual'),'B_sin_coordenadas_imputadas':all(not p['coordenada_individual_disponible'] and p['lat_original']=='' and p['lon_original']=='' for p in proposals if p['estado_revision']=='Resoluble nominalmente a escala comunal'),'ninguna_propuesta_aplicada':all(not p['aplicado'] for p in proposals)}
 for name,rows in [('propuestas_por_extremo',proposals),('resumen_20_grupos',summaries),('revision_manual_minima',manual),('impacto_por_registro',impacts),('cobertura_condicional',stats),('referencias_verificadas',refs),('evidencia_geocodificada_auxiliar',support),('calidad_evidencia_auxiliar',peer_quality)]:write(name,rows)
 public=json.loads((OUT/'fuentes_publicas.json').read_text())
 instructions=[dict(tema='Alcance',detalle='Segunda revisión de 20 grupos (37 extremos,36 entrevistas), exclusivamente propuesta. No se modifica el Excel previo ni los originales ni decisiones. No se calculan flujos, expansión o logit.'),dict(tema='A con coordenada',detalle='Propuesta respaldada por punto original, dos polígonos concordantes y referencia/topónimo. En Puerto Río Tranquilo se conserva alerta de cercanía al límite; MINSAL y KMZ corroboran la pertenencia comunal del topónimo. No se reduce globalmente el umbral ni se aplica a puntos anónimos. No acredita acceso exacto. En el caso sin texto se propone solo comuna, no localidad.'),dict(tema='B nominal sin coordenada',detalle='Resuelve propuesta de agregación administrativa por topónimo específico, referencias y contexto. No fue posible verificar coordenada individual porque no existe. Debe leerse separada de A; su efecto en cobertura es condicional a aceptar este criterio nominal.'),dict(tema='Manual mínima',detalle='Solo los casos no resueltos de estos 20 grupos. Las preguntas solicitan el lugar/comuna o significado del texto; no revisar el universo regional restante. Los candidatos no son asignaciones.'),dict(tema='Controles',detalle='Se muestran control físico y sentido por registro; no se usa control como extremo ni el control más cercano para decidir comuna, ni se confirma ruta por carretera.'),dict(tema='Cobertura',detalle='Conteos de entrevistas/registros, nunca viajes únicos ni demanda física expandida. Se distinguen extremos resueltos y OD con ambos extremos completos. Un extremo fuera de estos20 grupos puede seguir pendiente.'),dict(tema='Fuentes',detalle='Referencias verificadas, evidencias de otras encuestas y fuentes públicas conservan procedencia. No se copian sus coordenadas a los registros sin coordenada. Las comunas corregidas/normalizadas de la fuente son auxiliares; no sustituyen prueba geográfica.'),dict(tema='Respuesta nueva',detalle='Puede completar Manual mínima. No se aplica ni sobrescribe ninguna decisión anterior; deberá aprobarse una actualización futura de la homologación.')]
 book=Workbook();book.remove(book.active)
 def sheet(title,rows):
  s=book.create_sheet(title);rows=list(rows);cols=list(dict.fromkeys(k for r in rows for k in r)) or ['Sin casos'];s.append(cols)
  for r in rows:
   s.append([r.get(c,'') for c in cols])
   for c in s[s.max_row]:
    if isinstance(c.value,str):c.data_type='s'
    c.alignment=Alignment(vertical='top',wrap_text=True)
  for c in s[1]:c.font=Font(bold=True,color='FFFFFF');c.fill=PatternFill('solid',fgColor='225D65')
  s.freeze_panes='B2';s.auto_filter.ref=s.dimensions
  for col in s.columns:s.column_dimensions[col[0].column_letter].width=min(60,max(20,len(str(col[0].value))+2))
  return s
 sheet('Leer primero',instructions);sheet('Cobertura condicional',stats);sheet('Resumen 20 grupos',summaries);sheet('Propuestas por extremo',proposals);sheet('Impacto por registro',impacts);sheet('Manual mínima',manual);sheet('Referencias verificadas',refs);sheet('Evidencia auxiliar',support);sheet('Calidad evidencia',peer_quality);sheet('Fuentes públicas',public);sheet('Decisiones previas',groups);sheet('Validaciones',[dict(comprobacion=k,correcto=v) for k,v in checks.items()]);sheet('Huellas de entrada',hashes)
 target=OUT/'propuesta_segunda_revision.xlsx';book.save(target)
 w=load_workbook(target,read_only=True,data_only=True);it=w['Propuestas por extremo'].values;hh=next(it);roundtrip=[dict(zip(hh,row)) for row in it];w.close()
 checks['Excel_37_extremos_unicos']=len(roundtrip)==37 and len({(r['id_encuesta'],r['extremo']) for r in roundtrip})==37
 checks['Excel_literales_preservados']=all(r['literal_original']==pmap[r['id_encuesta'],r['extremo']]['literal_original'] for r in roundtrip)
 checks['fuentes_y_entrega_anterior_sin_cambios']=all(sha(r['archivo'])==r['sha256'] for r in hashes)
 summary=dict(estado='correcto' if all(checks.values()) else 'ERROR',comprobaciones=checks,grupos=20,extremos=37,registros=36,resueltos_con_coordenada=sum(p['estado_revision']=='Resoluble con coordenada individual' for p in proposals),resueltos_nominales_sin_coordenada=sum(p['estado_revision']=='Resoluble nominalmente a escala comunal' for p in proposals),extremos_manual=sum(p['requiere_manual'] for p in proposals),grupos_manual=len(manual),cobertura_condicional=stats,archivos_protegidos=len(hashes),sha256_excel=sha(target),sha256_codigo=sha(Path(__file__)))
 (OUT/'validacion.json').write_text(json.dumps(summary,ensure_ascii=False,indent=2));assert all(checks.values()),summary
 print(json.dumps(summary,ensure_ascii=False,indent=2))

if __name__=='__main__':main()
