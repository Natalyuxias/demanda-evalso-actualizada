"""Homologación geográfica local de CITRA. No importa ni ejecuta modelos/expansiones.
Ejecutar: python homologacion_geografica_etapa1/scripts/homologar_geografia.py
Escribe solo dentro de homologacion_geografica_etapa1; lee fuentes originales.
"""
from pathlib import Path
import collections,csv,datetime,hashlib,json,math,re,unicodedata
import pandas as pd
import geopandas as gp
from shapely.geometry import Point
from pyproj import Transformer
from openpyxl import Workbook,load_workbook
from openpyxl.styles import Font,PatternFill,Alignment
from openpyxl.cell import WriteOnlyCell
from openpyxl.worksheet.datavalidation import DataValidation

OUT=Path(__file__).resolve().parents[1]; BASE=OUT.parent
SEPARADO=Path('/home/nataly/Escritorio/demanda_evalso')
CFG=json.loads((OUT/'configuracion.json').read_text())
ALIGNMENT=Alignment(vertical='top',wrap_text=True)

def norm(v):
 s=str(v or '').upper().replace('´',"'").replace('’',"'")
 s=''.join(c for c in unicodedata.normalize('NFD',s) if unicodedata.category(c)!='Mn')
 return re.sub(r'\s+',' ',s).strip().strip(' .')

def canonical(v):
 n=norm(v)
 aliases={'COIHAIQUE':'COYHAIQUE','COYHAYQUE':'COYHAIQUE','AISEN':'AYSEN','RIO TRANQUILO':'PUERTO RIO TRANQUILO','PUERTO TRANQUILO':'PUERTO RIO TRANQUILO','CERRO CASTILLO':'VILLA CERRO CASTILLO','PUERTO INGENIERO IBANEZ':'PUERTO IBANEZ','PTO. IBANEZ':'PUERTO IBANEZ','PTO IBANEZ':'PUERTO IBANEZ','PTO GUADAL':'PUERTO GUADAL','PTO. GUADAL':'PUERTO GUADAL','GUADAL':'PUERTO GUADAL','BAHIA MURTA':'PUERTO MURTA','MURTA':'PUERTO MURTA','CHACABUCO':'PUERTO CHACABUCO','PUERTO RAUL MARIN BALMACEDA':'RAUL MARIN BALMACEDA','OHIGGINS':"O'HIGGINS",'VILLA OHIGGINS':"VILLA O'HIGGINS",'VILLA O HIGGINS':"VILLA O'HIGGINS",'VILLA TORTEL':'CALETA TORTEL'}
 return aliases.get(n,n)

def sha(p):return hashlib.sha256(Path(p).read_bytes()).hexdigest()
def json_value(v):
 if isinstance(v,(datetime.datetime,datetime.date,datetime.time)):return {'tipo':type(v).__name__,'valor':v.isoformat()}
 return {'tipo':type(v).__name__,'valor':v}
def row_sha(values):return hashlib.sha256(json.dumps([json_value(v) for v in values],ensure_ascii=False,separators=(',',':'),allow_nan=False).encode()).hexdigest()
def readcsv(p):return list(csv.DictReader(Path(p).open(encoding='utf-8-sig',newline='')))
def writecsv(p,rows):
 rows=list(rows)
 fields=list(dict.fromkeys(k for r in rows for k in r)) or ['Sin registros']
 with Path(p).open('w',encoding='utf-8-sig',newline='') as f:
  w=csv.DictWriter(f,fieldnames=fields);w.writeheader();w.writerows(rows)
def num(v):
 try:
  f=float(v)
  return f if math.isfinite(f) else None
 except (ValueError,TypeError):return None

def sources():
 paths=[BASE/x for x in CFG['fuentes']]
 for base in ['Red_Vial_Final/Zonificación']:
  for ext in ['shp','dbf','shx','prj','cpg']:paths.append(BASE/'Cap 7/SIG Red_Vial_Final'/('Zonificación.'+ext))
 for ext in ['shp','dbf','shx','prj','cpg']:paths.append(BASE/'SIG/Coberturas/IDE'/('COMUNAS.'+ext))
 paths+=[BASE/'Cap 8/BBDD_para_Pijas_actualizada_EOD_Final.xlsx',BASE/'diagnostico_etapa1_2026-10-28/propuesta_correspondencias_geograficas.csv',SEPARADO/'matriz_od_proyecto/tablas/referencias_localidades.csv',SEPARADO/'eod_filtrada/tablas/referencias_recuperadas.csv',SEPARADO/'eod_filtrada/tablas/base_completa_filtrada.csv',SEPARADO/'analisis_eod/tablas/eod_recuperada_local.csv',SEPARADO/'eod_filtrada/revision_manual_prioritaria.xlsx',SEPARADO/'matriz_od_proyecto/revision_area_influencia.xlsx']
 return [p for p in paths if p.exists()]

class Homologador:
 def __init__(self):
  self.comunas=gp.read_file(BASE/'SIG/Coberturas/IDE/COMUNAS.shp').to_crs(4326)
  self.zonas=gp.read_file(BASE/'Cap 7/SIG Red_Vial_Final/Zonificación.shp').to_crs(4326)
  self.comutm=self.comunas.to_crs(32718);self.zonutm=self.zonas.to_crs(32718)
  self.proj=Transformer.from_crs(4326,32718,always_xy=True)
  self.bycut={str(r.CUT_COM):dict(comuna=r.COMUNA,cut=str(r.CUT_COM)) for _,r in self.comunas.iterrows()}
  self.byname={canonical(r.COMUNA):str(r.CUT_COM) for _,r in self.comunas.iterrows()}
  self.cut_zone={str(r.CUT_COM):int(r.Zona) for _,r in self.zonas.iterrows() if str(r.CUT_REG)!='100'}
  wb=load_workbook(BASE/'Cap 8/BBDD_para_Pijas_actualizada_EOD_Final.xlsx',read_only=True,data_only=True,keep_links=False)
  self.catalogo=[dict(zona_aristo=int(r[1]),nombre=r[0],region_fuente=r[2]) for r in list(wb['Zonas'].values)[1:] if r[0] is not None];wb.close()
  assert sorted(r['zona_aristo'] for r in self.catalogo)==list(range(1,56))
  self.zone_name={r['zona_aristo']:r['nombre'] for r in self.catalogo}
  for _,r in self.zonas.iterrows():
   if int(r.Zona)<=50:assert canonical(r.COMUNA)==canonical(self.zone_name[int(r.Zona)])
  self.refs={};self.cache={}
  for r in readcsv(SEPARADO/'matriz_od_proyecto/tablas/referencias_localidades.csv'):
   hit=self.intersect(num(r['lat']),num(r['lon']));self.refs[canonical(r['localidad'])]=dict(lat=float(r['lat']),lon=float(r['lon']),cut=hit['cut'],zona=hit['zona'],fuente=r['fuente'],fuerte='KMZ' in r['fuente'])
  self.learned={canonical(r['nombre']):r for r in readcsv(SEPARADO/'eod_filtrada/tablas/referencias_recuperadas.csv')}

 def intersect(self,lat,lon):
  out=dict(cut='',comuna='',zona='',nombre_zona='',pais='',dist_limite_m=None,conflicto_poligonos=False)
  if lat is None or lon is None or not(-90<=lat<=90 and -180<=lon<=180) or lat==0 or lon==0:return out
  p=Point(lon,lat);cm=list(self.comunas.sindex.query(p,predicate='intersects'));zm=list(self.zonas.sindex.query(p,predicate='intersects'))
  if len(cm)>1 or len(zm)>1:out['conflicto_poligonos']=True;return out
  pt=Point(*self.proj.transform(lon,lat));margins=[]
  if cm:
   q=self.comunas.iloc[cm[0]];out.update(cut=str(q.CUT_COM),comuna=q.COMUNA,pais='Chile');margins.append(self.comutm.iloc[cm[0]].geometry.boundary.distance(pt))
  if zm:
   q=self.zonas.iloc[zm[0]];out.update(zona=int(q.Zona),nombre_zona=self.zone_name[int(q.Zona)] if hasattr(self,'zone_name') else q.COMUNA)
   out['pais']='Argentina' if str(q.CUT_REG)=='100' else 'Chile'
   margins.append(self.zonutm.iloc[zm[0]].geometry.boundary.distance(pt))
   if cm and str(q.CUT_COM)!=out['cut']:out['conflicto_poligonos']=True
  if bool(cm)!=bool(zm) and (not zm or int(self.zonas.iloc[zm[0]].Zona)<=50):out['conflicto_poligonos']=True
  if margins:out['dist_limite_m']=round(min(margins),1)
  return out

 def names(self,texts):
  # Solo componentes completos separados por comas. Una calle Chile Chico no es comuna Chile Chico.
  tokens=[]
  for text in texts:
   components=str(text or '').split(',')
   for position,token in enumerate(components):
    t=canonical(token)
    # En direcciones geocodificadas «Aysén» después de una localidad es
    # un calificador regional. No convertirlo en comuna Aysén.
    if t=='AYSEN' and position>0:continue
    t=canonical(re.sub(r'^\d{7}\s+','',t))  # prefijo postal explícito
    if t:tokens.append(t)
  cuts=set();locals=[];evidence=[]
  for t in tokens:
   if t in self.byname:cuts.add(self.byname[t]);evidence.append('comuna:'+t)
   if t in self.refs and self.refs[t]['cut']:
    cuts.add(self.refs[t]['cut']);locals.append(t);evidence.append('localidad:'+t)
  return cuts,list(dict.fromkeys(locals)),list(dict.fromkeys(evidence)),tokens

 def endpoint(self,literal,context=''):
  cachekey=(literal,context)
  if cachekey in self.cache:return self.cache[cachekey].copy()
  out=dict(literal_original=literal,localidad_original='',formato_original='texto',busqueda_original='',direccion_original='',lat_original='',lon_original='',crs_original='',comuna_declarada_original=context,localidad_homologada='',escala_localidad='No resuelta; literal conservado',comuna_homologada='',cut_comuna='',zona_modelacion_aristo='',nombre_zona_aristo='',pais_homologado='',estado='Pendiente',metodo='',motivo='',comuna_candidata='',cut_candidato='',zona_candidata='',evidencia_textual='',evidencia_espacial='',distancia_limite_m='',calidad_coordenada='Sin coordenada original; no se sustituye',precision_acceso='No acredita ubicación exacta ni acceso puerta a puerta',fuente_asignacion='')
  j={};texts=[literal];invalid=False
  if literal.lstrip().startswith('{'):
   try:j=json.loads(literal)
   except (ValueError,TypeError):invalid=True;out['formato_original']='JSON inválido'
   if isinstance(j,dict) and j:
    out.update(formato_original='JSON',busqueda_original=j.get('search',''),direccion_original=j.get('address',''))
    coords=j.get('coords') or {};out['lat_original']=coords.get('lat','');out['lon_original']=coords.get('lon','');out['crs_original']='EPSG:4326 (según extracción entregada)'
    texts=[j.get('search',''),j.get('address','')]
  out['localidad_original']=str(texts[0]) if texts[0] else ''
  cuts,locals_,evid,tokens=self.names(texts)
  # En «Puerto Guadal, Chile Chico, Chile», el segundo componente es
  # administrativo. No contrastar el punto con el centro de Chile Chico.
  primary_cuts,primary_locals,_,_=self.names([str(texts[0]).split(',')[0]])
  ccuts,clocs,cevid,_=self.names([context])
  # Comuna original es contraste, no autorización para convertir una etiqueta vaga en localidad.
  allcuts=cuts|ccuts
  lat,lon=num(out['lat_original']),num(out['lon_original'])
  hascoords=out['lat_original']!='' or out['lon_original']!=''
  validcoords=lat is not None and lon is not None and -90<=lat<=90 and -180<=lon<=180 and lat!=0 and lon!=0
  spatial=self.intersect(lat,lon) if validcoords else dict(cut='',comuna='',zona='',nombre_zona='',pais='',dist_limite_m=None,conflicto_poligonos=False)
  out.update(distancia_limite_m=spatial['dist_limite_m'] if spatial['dist_limite_m'] is not None else '',evidencia_textual='; '.join(evid+cevid))
  if validcoords:
   out['calidad_coordenada']='Geocodificada original; precisión y carácter individual no certificados'
   out['evidencia_espacial']='Intersección punto original con IDE/COMUNAS y Zonificación ARISTO, CRS armonizado'
  elif hascoords:out['calidad_coordenada']='Coordenada original inválida o incompleta; no corregida'
  candidate_cuts=set(allcuts)
  if spatial['cut']:candidate_cuts.add(spatial['cut'])
  out['cut_candidato']=' | '.join(sorted(candidate_cuts));out['comuna_candidata']=' | '.join(self.bycut[k]['comuna'] for k in sorted(candidate_cuts) if k in self.bycut)
  zones={self.cut_zone[k] for k in candidate_cuts if k in self.cut_zone}
  if spatial['zona']:zones.add(spatial['zona'])
  out['zona_candidata']=' | '.join(str(k) for k in sorted(zones))
  foreign={'ARGENTINA','ALEMANIA','FRANCIA','ESPANA','ESTADOS UNIDOS','EEUU','BRASIL','SUIZA','ITALIA','INGLATERRA','BELGICA','ISRAEL','AUSTRALIA','CANADA','PERU','URUGUAY','LOS ANTIGUOS','PERITO MORENO','COMODORO RIVADAVIA','RIO GALLEGOS','BARILOCHE','ESQUEL','TRELEW','GOBERNADOR GREGORES'}
  foreign_primary=any(t in foreign for t in self.names(texts[:1])[3])
  refconflict=[]
  if validcoords:
   for name in primary_locals:
    rr=self.refs[name]
    dist=Point(*self.proj.transform(lon,lat)).distance(Point(*self.proj.transform(rr['lon'],rr['lat'])))
    if dist>5000:refconflict.append(name)
  method='';cut='';zone='';reason='Sin contraste geográfico suficiente'
  if invalid:reason='JSON original no interpretable'
  elif hascoords and not validcoords:reason='Coordenada original inválida; requiere resolver antes de asignar'
  elif spatial['conflicto_poligonos']:reason='Solape o discrepancia entre polígonos administrativos y zonificación'
  elif len(allcuts)>1:reason='Conflicto entre topónimos o comuna declarada original'
  elif validcoords and foreign_primary and spatial['pais']=='Chile':reason='Referencia internacional del extremo y coordenada chilena; naturaleza del extremo no resuelta'
  elif validcoords and spatial['cut'] and allcuts and spatial['cut'] not in allcuts:reason='Coordenada y texto/comuna original corresponden a comunas distintas'
  elif refconflict:reason='Topónimo de localidad y coordenada separados más de 5 km; posible referencia genérica'
  elif validcoords and spatial['dist_limite_m'] is not None and spatial['dist_limite_m']<=CFG['margen_preventivo_limite_m']:reason='Coordenada cercana al límite; margen preventivo de revisión de 250 m'
  elif validcoords and spatial['cut'] and len(cuts)==1 and spatial['cut'] in cuts:
   cut=spatial['cut'];zone=spatial['zona'];method='Coordenada y topónimo administrativo/localidad concordantes'
  elif validcoords and spatial['pais']=='Argentina':
   # Solo asignar provincia si está expresamente nombrada además del polígono.
   prov={'SANTA CRUZ':51,'PROVINCIA DE SANTA CRUZ':51,'CHUBUT':52,'PROVINCIA DEL CHUBUT':52,'RIO NEGRO':53,'PROVINCIA DE RIO NEGRO':53,'TIERRA DEL FUEGO':54,'PROVINCIA TIERRA DEL FUEGO':54}
   if any(prov.get(t)==spatial['zona'] for t in tokens):zone=spatial['zona'];method='Coordenada y provincia argentina declarada concordantes'
   else:reason='Provincia argentina sin contraste textual suficiente; conservar solo candidatos'
  elif not validcoords and len(cuts)==1:
   cut=next(iter(cuts));zone=self.cut_zone.get(cut,'')
   used=[self.refs[n] for n in locals_ if self.refs[n]['cut']==cut]
   if any(t in self.byname for t in tokens):method='Comuna explícita verificada en catálogo y polígonos; ubicación detallada no asignada'
   elif any(r['fuerte'] for r in used):method='Topónimo explícito y referencia KMZ dentro de comuna/zona concordantes'
   elif cut in ccuts and used:method='Topónimo y comuna declarada concordantes con referencia calculada; sin coordenada individual'
   else:cut='';zone='';reason='Referencia de localidad calculada sin contraste independiente de comuna original'
  elif not validcoords and canonical(literal) in {'SANTIAGO','METROPOLITANA','TEMUCO','VINA DEL MAR'}:
   zone=55;method='Macrozona externa por topónimo explícito del diagnóstico; comuna/CUT fuera de cobertura cartográfica';out['estado']='Solo zona; comuna no asignada'
  elif validcoords and spatial['cut']:reason='Intersección administrativa disponible, sin topónimo/comuna textual verificable'
  elif validcoords:reason='Coordenada fuera de cobertura cartográfica o no homologable con fuentes disponibles'
  if method:
   out.update(cut_comuna=cut,comuna_homologada=self.bycut[cut]['comuna'] if cut else '',zona_modelacion_aristo=zone,nombre_zona_aristo=self.zone_name.get(zone,''),pais_homologado='Chile' if cut or zone==55 else 'Argentina',metodo=method,motivo='Asignación administrativa; no acredita precisión del extremo ni itinerario',fuente_asignacion='IDE/COMUNAS + ARISTO Zonificación y catálogo Zonas; topónimo/referencia según método')
   if cut:out['estado']='Homologado comuna y zona'
   elif zone!=55:out['estado']='Solo zona extranjera; comuna no aplica'
   # No atribuir localidad desde cola administrativa de una dirección ni por cercanía.
   if len(primary_locals)==1:
    if primary_locals[0] in self.byname:
     out['escala_localidad']='Topónimo homónimo de comuna/localidad; no se atribuye escala urbana'
    else:
     out['localidad_homologada']=primary_locals[0];out['escala_localidad']='Topónimo explícito contrastado; ubicación aproximada'
   elif cut:out['escala_localidad']='Comuna resuelta; localidad original conservada sin asignación específica'
  else:out.update(motivo=reason,metodo='Sin asignación automática; candidatos separados')
  self.cache[cachekey]=out.copy();return out


def load_raw():
 raw=[];headers=[]
 for name in CFG['fuentes']:
  p=BASE/name;camp='febrero' if 'Febrero' in p.name else 'abril'
  w=load_workbook(p,read_only=True,data_only=False,keep_links=False);s=w['EOD'];it=s.iter_rows(values_only=True);hh=next(it)
  headers.extend(dict(archivo=p.name,columna=i+1,encabezado_original=h if h is not None else '') for i,h in enumerate(hh))
  n=0
  for rownum,vals in enumerate(it,2):
   if not any(v is not None for v in vals):continue
   n+=1;rid=('V' if camp=='febrero' else 'O')+f'{rownum:05d}'
   raw.append(dict(id_encuesta=rid,archivo=p.name,hoja='EOD',fila_excel=rownum,campana_archivo=camp,hash_registro_original=row_sha(vals),valores_originales=tuple(vals),origen_original='' if vals[8 if camp=='febrero' else 7] is None else str(vals[8 if camp=='febrero' else 7]),destino_original='' if vals[12 if camp=='febrero' else 8] is None else str(vals[12 if camp=='febrero' else 8]),comuna_origen_original='' if camp=='abril' or vals[9] is None else str(vals[9]),comuna_destino_original='' if camp=='abril' or vals[13] is None else str(vals[13])))
  assert n=={'febrero':5308,'abril':3520}[camp],(camp,n)
  w.close()
 return raw,headers


def sheet(w,title,rows,headers=None,manual=False):
 rows=list(rows);headers=headers or (list(rows[0]) if rows else ['Sin casos']);s=w.create_sheet(title);s.freeze_panes='A2';s.auto_filter.ref=f'A1:{__import__("openpyxl").utils.get_column_letter(len(headers))}{len(rows)+1}'
 for i,h in enumerate(headers,1):s.column_dimensions[__import__('openpyxl').utils.get_column_letter(i)].width=min(55,max(18,len(h)+2))
 hs=[]
 for h in headers:
  c=WriteOnlyCell(s,h);c.font=Font(bold=True,color='FFFFFF');c.fill=PatternFill('solid',fgColor='245F68');hs.append(c)
 s.append(hs)
 for r in rows:
  cells=[]
  for h in headers:
   v=r.get(h,'')
   if isinstance(v,str):
    if len(v)>32767:raise ValueError('Texto no preservable en Excel, no se truncará: '+h)
    c=WriteOnlyCell(s,v);c.data_type='s'
   else:c=WriteOnlyCell(s,v)
   if len(rows)<1000:c.alignment=ALIGNMENT
   cells.append(c)
  s.append(cells)
 if manual and rows:
  for field in ['CUT propuesto usuario','Zona propuesta usuario','Decisión geográfica']:
   col=headers.index(field)+1
   forcell=__import__('openpyxl').utils.get_column_letter(col)
   dv=DataValidation(type='list',formula1='"Pendiente,Confirmar candidato,Corregir,No resoluble"',allow_blank=True) if field=='Decisión geográfica' else None
   if dv:s.data_validations.append(dv);dv.add(f'{forcell}2:{forcell}{len(rows)+1}')
 return s


def main():
 previous=OUT/'validacion.json';existing=OUT/'homologacion_CITRA_ARISTO.xlsx'
 if previous.exists() and existing.exists():
  expected=json.loads(previous.read_text()).get('sha256_excel')
  if expected and sha(existing)!=expected:
   raise RuntimeError('El Excel de entrega fue editado después de generarse. Se conserva intacto; usar otra carpeta de salida antes de regenerar.')
 originalhash=[dict(archivo=str(p),sha256=sha(p),bytes=p.stat().st_size) for p in sources()]
 (OUT/'huellas_antes.json').write_text(json.dumps(originalhash,ensure_ascii=False,indent=2))
 raw,headers=load_raw();idx={r['id_encuesta']:r for r in raw};assert len(idx)==len(raw)
 legacy={r['id_encuesta']:r for r in readcsv(SEPARADO/'eod_filtrada/tablas/base_completa_filtrada.csv')};assert set(idx)==set(legacy)
 extracted=readcsv(SEPARADO/'analisis_eod/tablas/eod_recuperada_local.csv');assert {('V' if r['campana_archivo']=='febrero' else 'O')+f"{int(r['fila_excel']):05d}" for r in extracted}==set(idx)
 # Leer formularios vigentes sin importar generadores analíticos (tienen efectos de escritura).
 manual=[];manual_by_case={}
 for path,kind in [(SEPARADO/'eod_filtrada/revision_manual_prioritaria.xlsx','Lista corta'),(SEPARADO/'matriz_od_proyecto/revision_area_influencia.xlsx','Área anterior')]:
  w=load_workbook(path,read_only=True,data_only=True)
  for sn in w.sheetnames:
   if sn not in ['20 casos prioritarios','Pares OD','Revision OD','_Mapa pares']:continue
   it=w[sn].iter_rows(values_only=True);hh=next(it)
   for rn,vals in enumerate(it,2):
    d=dict(zip(hh,vals));key=d.get('ID caso') or d.get('ID par') or d.get('ID par sin sentido')
    if not key:continue
    rec=dict(formulario=str(path),hoja=sn,fila=rn,ID_caso_o_par=str(key),campana=d.get('Campaña') or '',decision=d.get('Decisión usuario') or '',comentario=d.get('Comentarios del usuario') or '')
    manual.append(rec)
    if sn=='20 casos prioritarios':manual_by_case[str(key)]=(str(d.get('Decisión usuario') or 'Pendiente'),str(d.get('Comentarios del usuario') or ''))
  w.close()
 h=Homologador();records=[];ends=[];orig=[];manual_mismatch=[];literal_mismatch=[]
 keep=['id_par','id_caso','decision_usuario','comentarios_usuario','decision_previa','comentario_previo','decision_lista_corta','estado_filtro','motivo_filtro','ambiguo_relevante','ruta_utilizada','ruta_probable','control']
 for r in raw:
  old=legacy[r['id_encuesta']];rec={k:r[k] for k in ['id_encuesta','archivo','hoja','fila_excel','campana_archivo','hash_registro_original']}
  for k in keep:rec['previo_'+k]=old.get(k,'')
  current=manual_by_case.get(old['id_caso'],('',''))
  rec['decision_formulario_vigente']=current[0];rec['comentario_formulario_vigente']=current[1]
  if current[0] and current[0]!=old.get('decision_lista_corta',''):manual_mismatch.append(r['id_encuesta'])
  relevant=old['estado_filtro']=='Incluido' or (old['estado_filtro']=='Excluido provisional' and old['ambiguo_relevante']=='True')
  rec['relevante_revision']=relevant
  for end in ['origen','destino']:
   if r[end+'_original']!=old.get(end+'_declarado',''):literal_mismatch.append((r['id_encuesta'],end))
   x=h.endpoint(r[end+'_original'],r['comuna_'+end+'_original'])
   # Valores exactos de fuente acompañan siempre cada resultado administrativo.
   rec[end+'_original']=r[end+'_original'];rec[end+'_comuna_original']=r['comuna_'+end+'_original']
   for k,v in x.items():rec[end+'_'+k]=v
   er=dict(id_encuesta=r['id_encuesta'],archivo=r['archivo'],hoja='EOD',fila_excel=r['fila_excel'],campana_archivo=r['campana_archivo'],extremo=end,relevante_revision=relevant,previo_estado_filtro=old['estado_filtro'],previo_decision_area=old['decision_usuario'],previo_decision_lista=old['decision_lista_corta'],decision_formulario_vigente=current[0],**x)
   ends.append(er)
  rec['ambos_extremos_completos']=all(rec[e+'_estado']=='Homologado comuna y zona' for e in ['origen','destino'])
  records.append(rec)
  payload=json.dumps([json_value(v) for v in r['valores_originales']],ensure_ascii=False,separators=(',',':'),allow_nan=False)
  orig.append(dict(id_encuesta=r['id_encuesta'],archivo=r['archivo'],hoja='EOD',fila_excel=r['fila_excel'],hash_registro_original=r['hash_registro_original'],registro_original_JSON=payload,**{f'original_col_{i+1:02d}':v for i,v in enumerate(r['valores_originales'])}))
 assert not manual_mismatch,'Formulario cambió frente al filtro: revisar sin recalcular área'
 assert not literal_mismatch,('Originales y auxiliar de filtrado discrepan',literal_mismatch[:5])
 # Revisar solo pendientes relevantes; conservar todo el resto en Extremos sin exigir revisión.
 groups=collections.defaultdict(list)
 for e in ends:
  e['id_caso_geografico']=''
  if e['relevante_revision'] and e['estado']=='Pendiente':
   label=canonical(e['busqueda_original'] or (e['localidad_original'] if e['formato_original']=='texto' else e['direccion_original']))
   key=(label,e['cut_candidato'],e['zona_candidata'],e['motivo'])
   groups[key].append(e)
 reviews=[]
 for key,ee in groups.items():
  cid=hashlib.sha256('|'.join(key).encode()).hexdigest()[:14]
  for e in ee:e['id_caso_geografico']=cid
  ids=sorted({e['id_encuesta'] for e in ee});incl=sum(e['previo_estado_filtro']=='Incluido' for e in ee)
  reviews.append(dict(ID_caso=cid,toponimo_o_busqueda=key[0],extremos=len(ee),registros_distintos=len(ids),extremos_de_incluidos=incl,prioridad='Alta: registros incluidos' if incl else 'Media: potencialmente relevantes',motivo=key[3],comunas_candidatas=' | '.join(sorted({e['comuna_candidata'] for e in ee if e['comuna_candidata']})),CUT_candidatos=key[1],zonas_candidatas=key[2],archivos=' | '.join(sorted({e['archivo'] for e in ee})),filas_ejemplo=' | '.join(e['id_encuesta']+':'+e['extremo'] for e in ee[:5]),todos_los_ID=' | '.join(ids),**{'Decisión geográfica':'Pendiente','CUT propuesto usuario':'','Zona propuesta usuario':'','Localidad propuesta usuario':'','Evidencia o comentario usuario':'','Precaución':'Decisión de geografía nueva; no modifica decisiones previas de área ni confirma ruta. Candidatos no son asignaciones.'}))
 reviews.sort(key=lambda r:(-r['extremos_de_incluidos'],-r['registros_distintos'],r['toponimo_o_busqueda']))
 for i,r in enumerate(reviews,1):r['orden_revision']=i
 top=reviews[:CFG['max_grupos_revision_prioritaria']]
 endsmap={(e['id_encuesta'],e['extremo']):e for e in ends}
 for r in records:
  for end in ['origen','destino']:r[end+'_id_caso_geografico']=endsmap[r['id_encuesta'],end]['id_caso_geografico']
 coverage=[]
 for camp in ['febrero','abril','Todas']:
  for universe in ['Base completa','Relevantes para revisión']:
   ee=[e for e in ends if (camp=='Todas' or e['campana_archivo']==camp) and (universe=='Base completa' or e['relevante_revision'])]
   rr=[r for r in records if (camp=='Todas' or r['campana_archivo']==camp) and (universe=='Base completa' or r['relevante_revision'])]
   for label,n,denom in [('Registros conservados',len(rr),len(rr)),('Extremos conservados',len(ee),len(ee)),('Extremos con comuna, CUT y zona',sum(e['estado']=='Homologado comuna y zona' for e in ee),len(ee)),('Extremos con alguna zona',sum(e['zona_modelacion_aristo']!='' for e in ee),len(ee)),('Registros con ambos extremos completos',sum(r['ambos_extremos_completos'] for r in rr),len(rr)),('Extremos pendientes o parciales',sum(e['estado']!='Homologado comuna y zona' for e in ee),len(ee))]:
    coverage.append(dict(campana=camp,universo=universe,indicador=label,cantidad=n,denominador=denom,porcentaje=round(100*n/denom,2) if denom else None,unidad='registros' if label.startswith('Registros') else 'extremos'))
 statuses=[dict(campana=k[0],estado=k[1],relevante_revision=k[2],extremos=n) for k,n in sorted(collections.Counter((e['campana_archivo'],e['estado'],e['relevante_revision']) for e in ends).items())]
 methods=[dict(metodo=k,extremos=n) for k,n in collections.Counter(e['metodo'] for e in ends).most_common()]
 # Conservation is checked independently against original row values and hashes after workbook save.
 checks={'8828_registros_originales':len(raw)==8828,'17656_extremos':len(ends)==2*len(raw),'IDs_unicos_y_conservados':len({r['id_encuesta'] for r in records})==len(raw) and {r['id_encuesta'] for r in records}==set(idx),'claves_extremos_unicas':len(endsmap)==2*len(raw),'literales_coinciden_con_auxiliar_anterior':not literal_mismatch,'decisiones_formulario_y_filtro_coinciden':not manual_mismatch,'campos_previos_sin_modificar':all(all(r['previo_'+k]==legacy[r['id_encuesta']].get(k,'') for k in keep) for r in records),'ambiguos_sin_asignacion_final':all(not(e['cut_comuna'] or e['zona_modelacion_aristo'] or e['comuna_homologada']) for e in ends if e['estado']=='Pendiente'),'zonas_en_catalogo_55':all(e['zona_modelacion_aristo']=='' or e['zona_modelacion_aristo'] in range(1,56) for e in ends),'CUT_y_zona_concordantes':all(h.cut_zone[e['cut_comuna']]==e['zona_modelacion_aristo'] for e in ends if e['cut_comuna']),'55_no_confundida_con_comuna':all(not e['cut_comuna'] and not e['comuna_homologada'] for e in ends if e['zona_modelacion_aristo']==55),'revision_solo_relevantes':all(e['relevante_revision'] for e in ends if e['id_caso_geografico'])}
 rules=json.loads((OUT/'verificacion_reglas.json').read_text())
 checks['pruebas_geograficas_adversas_correctas']=rules['estado']=='correcto' and all(rules['pruebas'].values())
 for name,rows in [('correspondencias_por_registro',records),('extremos_homologados',ends),('resumen_cobertura',coverage),('estados',statuses),('revision_prioritaria',top),('pendientes_relevantes',reviews),('decisiones_previas',manual)]:writecsv(OUT/(name+'.csv'),rows)
 instructions=[dict(tema='Alcance',detalle='Homologación geográfica exclusivamente. No recalcula demanda, pesos, rutas, área, logit ni proyecciones. Resultados locales reservados.'),dict(tema='Registros y Extremos',detalle='Una fila por entrevista en Registros; dos extremos por entrevista en Extremos. Origen/destino, localidad, JSON, coordenadas y campos previos se conservan. Originales guarda las columnas sin modificar y Campos originales identifica sus encabezados por libro.'),dict(tema='Asignación y candidatos',detalle='Usar solo comuna_homologada, cut_comuna, zona_modelacion_aristo finales. Campos candidatos NO están asignados. Pendiente mantiene finales vacíos. Zona55 es macrozona: CUT/comuna permanecen vacíos.'),dict(tema='Precisión',detalle='Las coordenadas son geocodificadas, no necesariamente puntos exactos del viaje. Nunca se sustituyen por referencias. No se asigna localidad por proximidad. Referencias de localidades contrastan solo nombres explícitos.'),dict(tema='Reglas preventivas',detalle='Conflicto de texto/comuna/coordenada, solape de polígonos, coordenada inválida o cercanía ≤250m al límite deja pendiente. Distancia >5km entre topónimo primario y referencia también deja pendiente. Son criterios de revisión, no precisiones oficiales.'),dict(tema='Fuentes',detalle='IDE/COMUNAS EPSG5360 reproyectado4326; ARISTO/Zonificación EPSG4326; distancias EPSG32718. Hoja Zonas del libro BBDD_para_Pijas:55 códigos. BBDD-2 no se usa como código de modelación.'),dict(tema='Revisión prioritaria',detalle=f'{len(top)} grupos prioritarios de {len(reviews)} relevantes pendientes/parciales. Orden por presencia en viajes incluidos y número de registros; no usa flujos expandidos. Pendientes relevantes conserva todos los demás grupos pertinentes. No se exige revisar el resto de Aysén.'),dict(tema='Decisiones geográficas',detalle='Puede completar la hoja Revisión prioritaria con evidencia y propuesta. El script no lee ni aplica automáticamente esas nuevas decisiones; requerirán validación explícita. Sí/No anteriores son decisiones de área y siguen intactas.'),dict(tema='No cambiar significado',detalle='Homologación administrativa no confirma ruta por Chile ni cruce ni cambia estado de inclusión. Solo zona extranjera y zona55 son resultados parciales con comuna vacía. No convertir blancos en cero.')]
 instructions.append(dict(tema='Preservación exacta',detalle='Originales incluye registro_original_JSON con tipo y valor de todas las celdas. Conserva texto vacío distinto de nulo, aunque Excel los muestre igual. SHA por fila comprueba ese payload contra el Excel fuente. validacion.json contiene también las comprobaciones de lectura posterior a guardar.'))
 instructions.append(dict(tema='Alcance de la revisión',detalle='Solo Pendiente de geografía relevante entra en la revisión. Las macrozonas externas ya identificadas sin CUT disponible y las zonas argentinas sin comuna aplicable conservan su estado parcial; no requieren revisar localidades externas sin efecto en la zonificación.'))
 w=Workbook(write_only=True)
 visible=['id_encuesta','archivo','fila_excel','campana_archivo','ambos_extremos_completos','relevante_revision']
 for end in ['origen','destino']:
  visible.extend(end+'_'+k for k in ['original','localidad_original','lat_original','lon_original','localidad_homologada','escala_localidad','comuna_homologada','cut_comuna','zona_modelacion_aristo','nombre_zona_aristo','estado','metodo','motivo','id_caso_geografico'])
 visible.extend(['previo_estado_filtro','previo_motivo_filtro','previo_decision_usuario','previo_comentarios_usuario','previo_decision_previa','previo_comentario_previo','previo_decision_lista_corta','decision_formulario_vigente','comentario_formulario_vigente','hash_registro_original'])
 sheet(w,'Leer primero',instructions);sheet(w,'Registros CITRA',records,headers=visible);sheet(w,'Extremos',ends);sheet(w,'Resumen cobertura',coverage);sheet(w,'Estados y métodos',statuses+[dict(metodo=r['metodo'],extremos=r['extremos']) for r in methods],headers=['campana','estado','relevante_revision','metodo','extremos']);sheet(w,'Revisión prioritaria',top,manual=True);sheet(w,'Pendientes relevantes',reviews);sheet(w,'Decisiones previas',manual);sheet(w,'Originales',orig);sheet(w,'Campos originales',headers);sheet(w,'Catálogo zonas ARISTO',h.catalogo);sheet(w,'Validaciones', [dict(comprobacion=k,correcto=v) for k,v in checks.items()]);sheet(w,'Fuentes y hashes',originalhash)
 target=OUT/'homologacion_CITRA_ARISTO.xlsx';w.save(target)
 # Lectura de vuelta: comprobar todas las celdas originales, no solo número de filas.
 w=load_workbook(target,read_only=True,data_only=False)
 it=w['Originales'].iter_rows(values_only=True);hh=next(it);verified=0
 for vals in it:
  rr=dict(zip(hh,vals));source=idx[rr['id_encuesta']];values=tuple(rr.get(f'original_col_{i+1:02d}') for i in range(len(source['valores_originales'])))
  # Excel representa tanto texto vacío como nulo con una celda visualmente
  # vacía. El payload tipado conserva exactamente la distinción original.
  assert all(a==b or (a=='' and b is None) for a,b in zip(source['valores_originales'],values)),('Cambio en copia de original',source['id_encuesta'])
  payload=rr['registro_original_JSON']
  assert json.loads(payload)==[json_value(v) for v in source['valores_originales']]
  assert hashlib.sha256(payload.encode()).hexdigest()==source['hash_registro_original'];verified+=1
 checks['celdas_originales_y_hashes_verificados_tras_Excel']=verified==len(raw)
 for sn,key in [('Registros CITRA','id_encuesta'),('Extremos','id_encuesta')]:
  it=w[sn].iter_rows(values_only=True);hh=next(it);cc=collections.Counter(vals[hh.index(key)] for vals in it)
  checks['lectura_Excel_'+sn]=set(cc)==set(idx) and all(n==(2 if sn=='Extremos' else 1) for n in cc.values())
 w.close()
 checks['fuentes_y_formularios_hash_sin_cambios']=all(sha(r['archivo'])==r['sha256'] for r in originalhash)
 validation=dict(estado='correcto' if all(checks.values()) else 'ERROR',comprobaciones=checks,registros=len(raw),extremos=len(ends),grupos_revision_prioritaria=len(top),grupos_pendientes_relevantes=len(reviews),filas_fuente_preservadas_con_sha=verified,campos_fuente='Todas las columnas originales de EOD por posición y encabezado; no deduplicadas; tipo y valor exactos en registro_original_JSON',sha256_excel=sha(target),codigo_sha256=sha(Path(__file__)),configuracion_sha256=sha(OUT/'configuracion.json'),fuentes_controladas=len(originalhash),pruebas_reglas=rules,alcance='Geografía únicamente; decisiones de área preservadas. No estimación, expansión, calibración ni proyección.')
 (OUT/'validacion.json').write_text(json.dumps(validation,ensure_ascii=False,indent=2))
 assert all(checks.values()),validation
 print(json.dumps(validation,ensure_ascii=False,indent=2));print('ESTADOS',collections.Counter(e['estado'] for e in ends));print('RELEVANTES',collections.Counter(e['estado'] for e in ends if e['relevante_revision']))

if __name__=='__main__':main()
