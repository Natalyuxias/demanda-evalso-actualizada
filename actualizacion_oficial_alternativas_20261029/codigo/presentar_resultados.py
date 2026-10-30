"""Presentación del Excel oficial ya calculado. No ejecuta ningún modelo.
Desde Anexos: python actualizacion_oficial_alternativas_20261029/codigo/presentar_resultados.py
"""
from pathlib import Path
import ast, hashlib, html, json, math, re
import pandas as pd
import numpy as np
from openpyxl import load_workbook, Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

PRINCIPALES=['LEER PRIMERO','Demanda actual 2026','Demanda anual','Demanda estacional','Captacion e induccion','Sensibilidades','Insumos evaluacion social']
UNIDADES={'personas_peatones':'personas a pie','vehículos':'vehículos','unidades_privadas_heterogeneas':'unidades VP (categoría original)'}
RUTAS=['Terrestre por Chile','Puerto Ibáñez–Chile Chico','Bahía Jara–Levicán']
BASICOS={'alternativa_proyecto':'Proyecto','anio':'Año','segmento':'Segmento','unidad':'Unidad','temporada':'Temporada','direccion_cruce':'Sentido','fuente':'Fuente','origen':'Origen','destino':'Destino','zona_origen':'Zona ARISTO de origen','zona_destino':'Zona ARISTO de destino','ruta_transporte':'Ruta de transporte'}
CONEXIONES={'OD_original':'Localización original disponible','circuito_Coyhaique':'Circuito coordinado por Coyhaique','parcial_o_sin_conexion':'Conexión parcial o sin conexión','Rio_Ibanez_proxy_PUERTO IBANEZ':'Aproximación: Puerto Ibáñez','Rio_Ibanez_proxy_PUERTO RIO TRANQUILO':'Aproximación: Puerto Río Tranquilo','Rio_Ibanez_proxy_VILLA CERRO CASTILLO':'Aproximación: Villa Cerro Castillo'}
NOTA='ND: información insuficiente. —: no aplica. 0: cero calculado. Las cantidades fuera del soporte se muestran aparte; no son captación cero.'

def central(x):
    return x[x.escenario_captacion.eq('central')&x.escenario_elasticidad.eq('central')].copy()
def legible(x,mapping=None):
    x=x.copy()
    if 'unidad' in x:x['unidad']=x.unidad.replace(UNIDADES)
    if 'segmento' in x:x['segmento']=x.segmento.replace({'PEATON':'Peatones'})
    return x.rename(columns={**BASICOS,**(mapping or {})})
def tabla_filas(ws,df,titulo=None,nota=None):
    encabezado=getattr(ws,'_presentation_next_row',1)+int(bool(titulo))+int(bool(nota))
    for i,h in enumerate(df.columns,1):ws.column_dimensions[get_column_letter(i)].width=min(32,max(16,len(h)*.8))
    if ws.title=='LEER PRIMERO':ws.column_dimensions['B'].width=110
    ws.row_dimensions[encabezado].height=44
    if titulo:
        c=WriteOnlyCell(ws,titulo);c.font=Font(bold=True,color='17364A',size=13);ws.append([c])
    if nota:ws.append([nota])
    hdr=[]
    for h in df.columns:
        c=WriteOnlyCell(ws,h);c.font=Font(bold=True,color='FFFFFF');c.fill=PatternFill('solid',fgColor='24566E');c.alignment=Alignment(wrap_text=True,vertical='top');hdr.append(c)
    ws.append(hdr)
    for row in df.itertuples(index=False,name=None):
        out=[]
        for h,v in zip(df.columns,row):
            if v is None or (not isinstance(v,(list,dict,str)) and pd.isna(v)):v='ND'
            if isinstance(v,np.generic):v=v.item()
            c=WriteOnlyCell(ws,v)
            if isinstance(v,(int,float)) and not isinstance(v,bool):
                c.number_format='0.0%' if ('Participación' in h or '% ' in h or h.startswith('Ahorro relativo') or h.startswith('Probabilidad')) else ('0' if h in ['Año','Mes'] else '#,##0.00')
            c.alignment=Alignment(vertical='top',wrap_text=ws.title=='LEER PRIMERO')
            out.append(c)
        ws.append(out)
    ws.append([])
    ws._presentation_next_row=encabezado+len(df)+2

def huella_filas(rows):
    h=hashlib.sha256()
    for row in rows:
        row=list(row)
        while row and row[-1] is None:row.pop()
        values=[format(v,'.12g') if isinstance(v,float) else v for v in row]
        h.update(json.dumps(values,ensure_ascii=False,default=str,separators=(',',':')).encode());h.update(b'\n')
    return h.hexdigest()

def presentar(salida=None):
    root=Path(__file__).resolve().parents[2]
    out=Path(salida) if salida else Path(__file__).resolve().parents[1]/'resultados'
    path=out/'demanda_oficial_tres_alternativas.xlsx'
    source=load_workbook(path,read_only=True,data_only=False)
    respaldos=[n for n in source.sheetnames if n not in PRINCIPALES]
    necesarios=['Referencia estacional','Referencia anual 2026','CP rutas anual','CP rutas estacional','Captacion procedencia','Resumen oficial','Demanda detallada','Costos por alternativa','Sensibilidades historicas','Crecimiento referencia','Curvas heredadas','Elasticidades heredadas','Extranjeros muestra','Validaciones','Frecuencias vigentes']
    data={}
    for name in necesarios:
        it=iter(source[name].values);headers=next(it);data[name]=pd.DataFrame(it,columns=headers)
    # Se agregan o seleccionan valores guardados; no se ejecutan funciones de demanda.
    ref=data['Referencia estacional'].copy()
    rk=['fuente','segmento','unidad','temporada']
    ref=ref.groupby(rk,dropna=False).agg(Q_referencia=('Q_referencia',lambda x:x.sum(min_count=1)),dias_representados=('dias_representados','first')).reset_index()
    ref['Promedio diario equivalente']=ref.Q_referencia/ref.dias_representados
    ref['Cobertura']=ref.fuente.map({'ARISTO':'Barcaza observada; anualización y OD según entrega aprobada','CITRA':'Solo corredores cubiertos; factor horario 1,20 y supuestos temporales aprobados'})
    ann=data['Referencia anual 2026'][['fuente','segmento','unidad','Q_referencia','alcance']].copy();ann['temporada']='Total anual';ann['dias_representados']=365;ann['Promedio diario equivalente']=ann.Q_referencia/365;ann=ann.rename(columns={'alcance':'Cobertura'})
    actual=legible(pd.concat([ref,ann],ignore_index=True)[['fuente','segmento','unidad','temporada','Q_referencia','dias_representados','Promedio diario equivalente','Cobertura']],{'Q_referencia':'Demanda de referencia 2026','dias_representados':'Días representados'})
    actual=actual.sort_values(['Fuente','Segmento','Temporada'])
    def rutas_presentadas(name):
        x=central(data[name])
        if name=='CP rutas anual':x=x[x.direccion_cruce.eq('Ambos sentidos')].copy()
        x['Fuera del soporte (ruta de origen)']=x.demanda_SBO_conocida-x.demanda_SBO_compatible
        columns=['alternativa_proyecto','anio','temporada','direccion_cruce','segmento','unidad','ruta_transporte','demanda_SBO_conocida','demanda_SBO_compatible','demanda_CP_existente','demanda_CP_inducida','demanda_CP_total_modelada','participacion_CP_compatible','Fuera del soporte (ruta de origen)']
        if name=='CP rutas estacional':columns+=['dias_equivalentes','CP_diario_equivalente','CP_por_zarpe_equivalente']
        x=x[columns]
        if 'CP_por_zarpe_equivalente' in x:
            x['CP_por_zarpe_equivalente']=x.CP_por_zarpe_equivalente.astype(object)
            x.loc[x.ruta_transporte.eq(RUTAS[0]),'CP_por_zarpe_equivalente']='—'
        return legible(x,{'demanda_SBO_conocida':'Demanda SBO conocida','demanda_SBO_compatible':'Demanda SBO en soporte','demanda_CP_existente':'Existente que permanece o llega (CP)','demanda_CP_inducida':'Inducida incremental estimada (CP)','demanda_CP_total_modelada':'Demanda total estimada (CP)','participacion_CP_compatible':'Participación CP en mercado compatible','dias_equivalentes':'Días de temporada','CP_diario_equivalente':'Promedio diario CP','CP_por_zarpe_equivalente':'Promedio por zarpe CP'})
    anual=rutas_presentadas('CP rutas anual');estacional=rutas_presentadas('CP rutas estacional')
    cap=central(data['Captacion procedencia'])
    cap=legible(cap[['alternativa_proyecto','anio','segmento','unidad','Q_compatible','D_desviada_tierra','D_desviada_actual','D_inducida_nuevo','D_nuevo_CP_total','D_residual_existente','porcentaje_captacion_existente_compatible']],{'Q_compatible':'Demanda existente compatible','D_desviada_tierra':'Desviada desde tierra','D_desviada_actual':'Desviada desde barcaza actual','D_inducida_nuevo':'Inducida incremental al nuevo cruce','D_nuevo_CP_total':'Total Bahía Jara–Levicán','D_residual_existente':'Demanda fuera del soporte','porcentaje_captacion_existente_compatible':'Participación captada de demanda existente'})
    sm=data['Resumen oficial'];sm=sm[sm.version_referencia.eq('B')&sm.segmento.isin(['VL','CAM','VP','PEATON'])].copy()
    sm['Escenario de captación']=sm.escenario_captacion.map({'conservador':'Conservador','central':'Central','alto':'Alto'})
    sm['Sensibilidad de inducción']=sm.escenario_elasticidad.map({'menor_induccion':'Menor','central':'Central','mayor_induccion':'Mayor'})
    vig=legible(sm[['alternativa_proyecto','anio','segmento','unidad','Escenario de captación','Sensibilidad de inducción','D_existente_nuevo','D_inducida_nuevo','D_nuevo_CP_total']],{'D_existente_nuevo':'Desviada al nuevo cruce','D_inducida_nuevo':'Inducida incremental al nuevo cruce','D_nuevo_CP_total':'Total Bahía Jara–Levicán'})
    hist=data['Sensibilidades historicas'].copy()
    casos={'actual_160min':'Navegación actual 160 min','base_con_buses':'Base histórica con buses','base_sin_buses':'Base histórica sin buses','nueva_21semanales':'Nuevo cruce 21 salidas semanales','nueva_40min':'Navegación nueva 40 min','tarifa_nueva_75pct':'Tarifa nueva al 75 %'}
    hist['Configuración histórica']=hist.caso_operacional.map(casos).fillna(hist.caso_operacional)
    hist['Escenario de captación']=hist.escenario_transferencia.replace({'conservador':'Conservador','central':'Central','alto':'Alto'})
    hist=legible(hist[['Configuración histórica','Escenario de captación','anio','segmento','unidad','captacion','frecuencia_de_su_base','navegacion_de_su_base']],{'captacion':'Captación histórica del nuevo cruce','frecuencia_de_su_base':'Frecuencias usadas entonces','navegacion_de_su_base':'Navegación usada entonces'})
    growth=legible(data['Crecimiento referencia'].drop(columns=['version','uso']),{'escenario_crecimiento':'Escenario de crecimiento','elasticidad':'Elasticidad respecto del PIB','factor':'Factor de crecimiento guardado','flujo_fisico_anual':'Demanda de referencia proyectada','observado_OD_anualizado':'OD distribuido sin imputación','imputado_OD_anualizado':'OD distribuido mediante supuestos','residual_anualizado':'Residual de referencia'})
    elast=legible(data['Elasticidades heredadas'],{'menor_induccion':'Menor inducción','central':'Inducción central','mayor_induccion':'Mayor inducción'})
    curve_rows=[]
    for r in data['Curvas heredadas'].itertuples(index=False):
        for ahorro,p in zip(ast.literal_eval(r.ahorros),ast.literal_eval(r.probabilidades)):
            curve_rows.append({'Segmento':{'privados':'VL y VP','carga':'CAM','peatones':'Peatones'}.get(r.familia,r.familia),'Escenario':{'conservador':'Conservador','central':'Central','alto':'Alto'}.get(r.escenario,r.escenario),'Ahorro relativo de costo':ahorro,'Probabilidad de transferencia de la curva':p})
    curvas=pd.DataFrame(curve_rows)
    # Cruce de valores guardados por cohorte: evita unir por OD ambiguos o duplicar conexiones.
    detail=central(data['Demanda detallada']);detail=detail[detail.version_referencia.eq('B')&detail.segmento.isin(['VL','CAM','VP','PEATON'])].copy()
    costs=data['Costos por alternativa'];merged=detail.merge(costs,on=['alternativa_proyecto','id_base'],suffixes=('','_coste'),validate='one_to_one')
    insumos=[]
    for route,pretty in [('tierra',RUTAS[0]),('actual',RUTAS[1]),('nuevo',RUTAS[2])]:
        q=merged[['alternativa_proyecto','anio','fuente','origen','destino','zona_origen','zona_destino','segmento','unidad','temporada','direccion_cruce']].copy()
        q['Ruta de transporte']=pretty
        q['Conexión de acceso']=merged.conexion.map(CONEXIONES).fillna('Localización original disponible')
        q['Demanda de referencia del bloque']=merged.Q_referencia
        q['Demanda SBO']=0. if route=='nuevo' else merged['D_'+route+'_SBO']
        q['Demanda existente CP']=merged.D_existente_nuevo if route=='nuevo' else merged['D_'+route+'_CP_existente']
        q['Inducida incremental asignada']=merged['D_inducida_'+route]
        q['Demanda total CP estimada']=merged['D_'+route+'_CP_total']
        orig=merged.fuente.eq('CITRA' if route=='tierra' else 'ARISTO') if route!='nuevo' else pd.Series(False,index=merged.index)
        q['Salidas hacia el nuevo cruce']=merged.D_existente_nuevo.where(orig,0.)
        q['Desviados recibidos en nuevo cruce']=merged.D_existente_nuevo if route=='nuevo' else 0.
        q['Demanda fuera del soporte de esta ruta']=merged.D_residual_existente.where(orig,0.)
        q['Estado de la demanda']=np.where(merged.p_nuevo.isna(),'Fuera del soporte: reparto CP sin calcular','Reparto existente calculado')
        q['Estado de la inducción']=np.where(merged.modelable_induccion,'Calculada con método oficial','ND: no hay costos completos comparables')
        q['Distancia vial (km)']=merged.nuevo_km_actualizado if route=='nuevo' else merged[route+'_km_base']
        q['Tiempo vial (h)']=merged.nuevo_h_vial_actualizado if route=='nuevo' else merged[route+'_h_base']
        q['Distancia lacustre (km)']=merged.distancia_lacustre_km if route=='nuevo' else ('—' if route=='tierra' else np.nan)
        q['Navegación (h)']=0. if route=='tierra' else (merged.tiempo_navegacion_nuevo_h if route=='nuevo' else 107/60)
        q['Presentación y embarque (h)']=0. if route=='tierra' else merged.tiempo_presentacion_h
        q['Desembarque (h)']=0. if route=='tierra' else merged.tiempo_desembarque_h
        q['Penalización por frecuencia (h)']=0. if route=='tierra' else (merged.tiempo_horario_nuevo_h if route=='nuevo' else merged.T_horario_actual)
        q['Transbordo adicional de buses (h)']=merged.nuevo_transbordo_adicional_h if route=='nuevo' else np.nan
        q['Transbordo adicional de buses (h)']=q['Transbordo adicional de buses (h)'].astype(object)
        q.loc[~merged.segmento.eq('PEATON'),'Transbordo adicional de buses (h)']='—'
        q['Tiempo total registrado (h)']=merged.nuevo_tiempo_total_h if route=='nuevo' else np.nan
        q['COV privado (CLP ene-2026)']=merged.nuevo_COV_actualizado if route=='nuevo' else merged[route+'_cov_base']
        q['COV privado (CLP ene-2026)']=q['COV privado (CLP ene-2026)'].astype(object)
        q.loc[merged.segmento.eq('PEATON'),'COV privado (CLP ene-2026)']='—'
        q['Tarifa lacustre (CLP ene-2026)']='—' if route=='tierra' else (merged.nuevo_fare_base if route=='nuevo' else merged.actual_fare_base)
        q['Tarifa de buses (CLP ene-2026)']=merged.nuevo_tarifa_buses_actualizada if route=='nuevo' else np.nan
        q['Tarifa de buses (CLP ene-2026)']=q['Tarifa de buses (CLP ene-2026)'].astype(object)
        q.loc[~merged.segmento.eq('PEATON'),'Tarifa de buses (CLP ene-2026)']='—'
        q['Costo monetario privado total (CLP)']=merged.nuevo_costo_monetario if route=='nuevo' else np.nan
        q['Costo temporal (CLP)']=merged.nuevo_costo_temporal if route=='nuevo' else np.nan
        q['CG SBO (CLP por unidad)']='—' if route=='nuevo' else merged['CG_'+route+'_SBO']
        q['CG con proyecto (CLP por unidad)']=merged['CG_'+route+'_CP']
        q['Alcance del costo']=np.where(merged.costo_solo_nucleo,'Solo núcleo de viaje: no costo completo del OD','Costo de viaje según soporte existente')
        q['Fuente de atributos']=merged['nuevo_fuente_componentes' if route=='nuevo' else route+'_fuente_componentes']
        q['COV social (CLP por unidad)']='ND';q['Costo de recursos sociales (CLP por unidad)']='ND'
        q['Tratamiento económico']='CG: monetario privado + VST social común. Tarifa separada; no es costo social puro.'
        q['VST común (CLP/persona-h)']=10047
        insumos.append(legible(q))
    insumos=pd.concat(insumos,ignore_index=True)
    # Los montos solo se copian de sus columnas originales: no se reconstruyen costos.
    assert len(insumos)==len(merged)*3
    assert np.allclose(insumos[insumos['Ruta de transporte'].eq(RUTAS[2])]['Demanda total CP estimada'],merged.D_nuevo_CP_total,equal_nan=True)
    lectura=pd.DataFrame([
      ['Objetivo','Demanda existente, captada e inducida para comparar cada proyecto por separado contra una SBO común.'],
      ['A1. Convencional','Acceso P-20a,31,14 km; tipo La Tehuelche; rampa fija; navegación43 min.'],
      ['A2. Rápida con linkspan','Acceso P-20d,24,62 km; tipo Capitán Wellmann; linkspan; navegación39 min.'],
      ['A3. Rápida con rampa fija','Acceso P-20b,22,60 km; tipo Capitán Wellmann; muelles/rampa fija; navegación39 min.'],
      ['SBO común','Terrestre por Chile y Puerto Ibáñez–Chile Chico con Capitán Wellmann, navegación107 min.'],
      ['Frecuencias oficiales','Ambos cruces:11 salidas/semana/sentido en diciembre–abril y8 en mayo–noviembre.'],
      ['Qué hojas leer','Las seis hojas siguientes contienen resultados centrales e insumos. Las hojas grises al final son respaldo técnico.'],
      ['Referencia2026','Anualización aprobada de las campañas: no es un conteo observado durante todo2026. ARISTO y CITRA se mantienen separados.'],
      ['Observaciones originales','Campanas ARISTO/CITRA y OD campana ARISTO conservan conteos y encuestas por exposición medida.'],
      ['Unidades','VL y CAM:vehículos; VP:categoría privada original heterogénea; peatones:personas a pie. Nunca sumar personas con vehículos ni VP con sus componentes.'],
      ['SBO conocida / en soporte','Conocida incluye el residual físico. En soporte es la parte con reparto de rutas calculable. Comparar CP contra SBO en soporte.'],
      ['Demanda con proyecto','Demanda existente que permanece o llega más la inducción incremental estimada. Los faltantes de inducción no son ceros calculados.'],
      ['Demanda fuera del soporte','Cantidad conocida sin reparto CP defendible. Permanece separada; no se distribuye ni se borra.'],
      ['Totales','A1,A2,A3 son proyectos excluyentes. Ambos sentidos ya suma sus componentes: no sumarlo nuevamente con los sentidos separados.'],
      ['Temporadas','Verano:dic–feb; otoño:mar–may; invierno:jun–ago; primavera:sep–nov. Invierno mantiene aproximaciones de la anualización aprobada.'],
      ['Equivalentes','Promedios diarios y por zarpe: no demanda punta ni capacidad requerida. Ambos sentidos utiliza dos exposiciones de servicio.'],
      ['Símbolos',NOTA],
      ['CITRA','Solo mercado y corredores cubiertos, con factor horario1,20 y supuestos posteriores aprobados; no total regional.'],
      ['Peatones','Escenario hipotético con buses coordinados. No representa todos los pasajeros transportados ni buses nuevos ya contratados.'],
      ['Escenarios','Tablas principales:crecimiento,captación e inducción centrales. Sensibilidades muestra otros casos ya calculados y separa los históricos no vigentes.'],
      ['Precios y VST','CLPenero2026 convencionales del modelo. VST10047 CLP/persona-h. CG combina componentes privados con VST social común.'],
      ['Evaluación social','Usar Insumos evaluacion social y conservar distinción privado/social. Los COV sociales no están en la entrega oficial y quedan ND; no tratar tarifas como consumo adicional de recursos.'],
      ['Usuarios extranjeros','Nacionalidad,patente,OD internacional y residencia son distintos y pueden solaparse. No hay desagregación anual defendible por ruta/proyecto. Los restantes viajes no se suponen nacionales.'],
      ['Reproducir presentación','python actualizacion_oficial_alternativas_20261029/codigo/presentar_resultados.py. Lee solo este Excel; no ejecuta modelos.']
    ],columns=['Tema','Explicación'])
    fx=data['Extranjeros muestra'].copy()
    fx['Indicador']=fx.indicador.map({'OD_internacional':'Origen o destino internacional','nacionalidad_declarada_peaton':'Nacionalidad extranjera declarada (peatones)','patente_vehicular':'Patente extranjera (vehículos)','residencia_extranjera':'Residencia extranjera'})
    fx=fx[['Indicador','universo_aplicable_muestra','entrevistas_indicador_positivo']].rename(columns={'universo_aplicable_muestra':'Entrevistas con indicador aplicable','entrevistas_indicador_positivo':'Indicador positivo en muestra'})
    fx['Magnitud anual por alternativa']='ND'
    new=Workbook(write_only=True)
    dfs={'Demanda actual 2026':actual,'Demanda anual':anual,'Demanda estacional':estacional,'Captacion e induccion':cap,'Insumos evaluacion social':insumos}
    for name in PRINCIPALES:
        ws=new.create_sheet(name);ws.sheet_properties.tabColor='217D71' if name=='LEER PRIMERO' else '24566E';ws.freeze_panes='G4' if name=='Insumos evaluacion social' else 'D4'
        if name=='LEER PRIMERO':
            ws.column_dimensions['B'].width=110
            ws.freeze_panes='B4';tabla_filas(ws,lectura,'GUÍA PARA EL EQUIPO DE EVALUACIÓN SOCIAL');tabla_filas(ws,fx,'Usuarios extranjeros: indicadores disponibles','Son recuentos muestrales solapados; no sumar. ND no significa residencia nacional.');ws.column_dimensions['B'].width=110
        elif name=='Sensibilidades':
            tabla_filas(ws,vig,'VIGENTES — captación e inducción con frecuencia11/8 común','Cinco casos calculados por proyecto: tres captaciones con inducción central; menor/mayor inducción con captación central. No sumar casos.')
            tabla_filas(ws,elast,'VIGENTES — elasticidades de inducción','Elasticidades respecto del costo generalizado; no son elasticidades-PIB.')
            tabla_filas(ws,curvas,'VIGENTES — puntos de las curvas de transferencia','Supuestos del proyecto guardados; no calibración estadística. Las factibilidades y reglas adicionales permanecen en el respaldo técnico.')
            tabla_filas(ws,growth,'CONTRASTE — crecimiento de referencia ya calculado','No son nuevas captaciones. La presentación principal utiliza crecimiento central.')
            tabla_filas(ws,hist,'HISTÓRICOS — sensibilidades operacionales NO VIGENTES','Frecuencias y navegación diferentes de las oficiales actuales. No comparar como si fueran sensibilidades de A1/A2/A3 con frecuencia11/8.')
        else:
            note=NOTA+' Escenario central; no sumar rutas con su residual ni sumar alternativas de proyecto.'
            if name=='Insumos evaluacion social':note+=' Cada bloque de procedencia aparece en tres rutas; sumar demanda de cada ruta, no la referencia repetida. Costos sociales pendientes. En núcleos de viajes externos, distancias y CG no representan el viaje completo.'
            tabla_filas(ws,dfs[name],name.upper(),note)
            ws.auto_filter.ref=f'A3:{get_column_letter(len(dfs[name].columns))}{len(dfs[name])+3}'
    # Mantener todos los valores de respaldo; solo cambia su ubicación y estilo.
    originales={};filas={};columnas={}
    for name in respaldos:
        original=source[name];ncols=original.max_column or len(next(original.values));ws=new.create_sheet(name);ws.sheet_properties.tabColor='A6ADB4';ws.freeze_panes='D2';h=hashlib.sha256();n=0
        for i in range(1,ncols+1):ws.column_dimensions[get_column_letter(i)].width=22
        for row in original.values:
            row=list(row)
            while row and row[-1] is None:row.pop()
            norm=[format(v,'.12g') if isinstance(v,float) else v for v in row]
            h.update(json.dumps(norm,ensure_ascii=False,default=str,separators=(',',':')).encode());h.update(b'\n')
            outrow=[]
            for i,v in enumerate(row):
                c=WriteOnlyCell(ws,v)
                if n==0:c.font=Font(bold=True,color='FFFFFF');c.fill=PatternFill('solid',fgColor='667580')
                elif isinstance(v,(int,float)) and not isinstance(v,bool):c.number_format='#,##0.0000'
                outrow.append(c)
            ws.append(outrow);n+=1
        originales[name]=h.hexdigest();filas[name]=n;columnas[name]=ncols
        ws.auto_filter.ref=f'A1:{get_column_letter(ncols)}{n}'
        for i in range(1,ncols+1):ws.column_dimensions[get_column_letter(i)].width=22
    source.close()
    temp=out/'.presentacion_temporal.xlsx';new.save(temp)
    check=load_workbook(temp,read_only=True,data_only=False)
    assert check.sheetnames[:7]==PRINCIPALES
    for name in respaldos:
        assert huella_filas(check[name].iter_rows(max_col=columnas[name],values_only=True))==originales[name],f'Cambio inesperado en respaldo {name}'
    check.close()
    report=crear_informe(actual,anual,estacional,cap,vig,hist,fx)
    temp.replace(path)
    (out/'informe_alternativas_oficiales.html').write_text(report,encoding='utf-8')
    print(json.dumps({'hojas_principales':PRINCIPALES,'hojas_tecnicas_intactas':len(respaldos),'filas_insumos':len(insumos),'demanda_y_costos':'Sin cambios; valores técnicos verificados por huella normalizada a12 cifras significativas'},ensure_ascii=False),flush=True)


def crear_informe(actual,anual,estacional,cap,vig,hist,fx):
    def fmt(v):
        if v is None or (not isinstance(v,str) and pd.isna(v)):return 'ND'
        if isinstance(v,(int,float,np.number)):
            return f'{v:,.2f}'.replace(',','X').replace('.',',').replace('X','.')
        return html.escape(str(v))
    def tabla(df,pct=()):
        headers=''.join('<th>'+html.escape(str(c))+'</th>' for c in df.columns)
        rows=[]
        for row in df.itertuples(index=False,name=None):
            vals=[]
            for c,v in zip(df.columns,row):
                text=(fmt(v*100)+' %') if c in pct and pd.notna(v) else (str(int(v)) if c=='Año' else fmt(v))
                vals.append('<td>'+text+'</td>')
            rows.append('<tr>'+''.join(vals)+'</tr>')
        return '<div class="tabla"><table><thead><tr>'+headers+'</tr></thead><tbody>'+''.join(rows)+'</tbody></table></div>'
    reference=actual[actual.Temporada.eq('Total anual')][['Fuente','Segmento','Unidad','Demanda de referencia 2026','Cobertura']]
    a=anual.copy();keys=['Proyecto','Año','Segmento','Unidad'];parts=[]
    for ruta,label in zip(RUTAS,['Tierra','Barcaza actual','Nuevo cruce']):
        r=a[a['Ruta de transporte'].eq(ruta)].set_index(keys)[['Demanda SBO en soporte','Demanda total estimada (CP)']].rename(columns={'Demanda SBO en soporte':label+' SBO','Demanda total estimada (CP)':label+' CP'})
        parts.append(r)
    main=pd.concat(parts,axis=1).reset_index().drop(columns=['Nuevo cruce SBO'])
    # Una fila por proyecto/año/segmento, residual sumado solo en sus rutas de origen.
    residual=a.groupby(keys,dropna=False)['Fuera del soporte (ruta de origen)'].sum(min_count=1)
    main=main.merge(residual.rename('Fuera del soporte').reset_index(),on=keys,validate='one_to_one')
    def records(df):
        return json.loads(df.astype(object).where(pd.notna(df),None).to_json(orient='records',force_ascii=False))
    seascols=['Proyecto','Año','Temporada','Sentido','Segmento','Unidad','Ruta de transporte','Demanda SBO en soporte','Existente que permanece o llega (CP)','Inducida incremental estimada (CP)','Demanda total estimada (CP)','Promedio diario CP','Promedio por zarpe CP','Fuera del soporte (ruta de origen)']
    payload=json.dumps(records(estacional[seascols]),ensure_ascii=False).replace('</','<\\/')
    senscols=['Proyecto','Año','Segmento','Unidad','Escenario de captación','Sensibilidad de inducción','Desviada al nuevo cruce','Inducida incremental al nuevo cruce','Total Bahía Jara–Levicán']
    senspayload=json.dumps(records(vig[senscols]),ensure_ascii=False).replace('</','<\\/')
    return f'''<!doctype html><html lang="es"><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><title>Demanda oficial · Bahía Jara–Levicán</title>
<style>body{{font:16px system-ui;color:#193747;max-width:1450px;margin:32px auto;padding:0 20px}}h1,h2{{color:#145c68}}p,li{{line-height:1.55}}.tabla{{overflow:auto;max-height:650px;border:1px solid #d7e2e7;margin:16px 0}}table{{border-collapse:collapse;width:100%;font-size:13px}}th,td{{padding:9px;border-bottom:1px solid #d7e2e7;text-align:right;white-space:nowrap}}th{{position:sticky;top:0;background:#24566e;color:white;text-align:left}}td:first-child{{text-align:left}}.nota{{background:#eef5f5;padding:14px;border-left:4px solid #398776}}select{{padding:8px;margin:6px 18px 6px 4px}}a{{color:#125f8b}}details{{margin:18px 0}}summary{{cursor:pointer;font-weight:bold}}@media print{{.tabla{{max-height:none;overflow:visible}}th{{position:static}}select{{display:none}}}}</style>
<h1>Demanda oficial de Bahía Jara–Levicán</h1><p>Resultados centrales para el equipo de Evaluación Social de Proyectos. Esta entrega reorganiza valores ya calculados: no cambia demanda, costos, parámetros ni metodología.</p>
<p><a href="demanda_oficial_tres_alternativas.xlsx">Descargar Excel de resultados e insumos</a></p>
<h2>1. Demanda de referencia 2026</h2><p>ARISTO corresponde a usuarios de la barcaza; CITRA cubre parcialmente corredores terrestres. La referencia es una anualización con los supuestos aprobados, no un conteo completo de2026. VL y VP conservan las clases y coberturas de las campañas; sus aportes no constituyen por sí solos anuales completos de una clase estricta.</p>{tabla(reference)}
<p class="nota">No sumar personas con vehículos, VP con sus componentes, alternativas de proyecto entre sí ni total de ambos sentidos con sus sentidos separados. CITRA no representa toda la demanda terrestre regional. {NOTA}</p>
<h2>2. Demanda anual central de las tres alternativas</h2><p><b>A1:</b> P-20a31,14km, navegación43min. <b>A2:</b> P-20d24,62km,39min y linkspan. <b>A3:</b> P-20b22,60km,39min y rampa fija. Cada proyecto se compara por separado con la misma SBO: terrestre y Puerto Ibáñez–Chile Chico, navegación107min. Ambos cruces mantienen11 salidas/semana/sentido en diciembre–abril y8 en mayo–noviembre.</p>
<p>La tabla compara el mismo mercado con soporte. CP incluye la inducción incremental estimada. El residual conocido se conserva aparte y no se reparte artificialmente. Los peatones incluyen buses coordinados hipotéticos.</p>{tabla(main)}
<h2>3. Demanda por temporada y sentido</h2><p>Los promedios diarios y por zarpe son equivalentes normalizados, no demanda punta. Para ambos sentidos, el denominador por zarpe incluye ambos servicios.</p>
<label>Proyecto <select id="proyecto"><option>A1</option><option>A2</option><option>A3</option></select></label><label>Año <select id="anio"><option>2026</option><option>2030</option><option>2040</option><option>2050</option></select></label><label>Sentido <select id="sentido"><option>Ambos sentidos</option><option>Sur → Norte</option><option>Norte → Sur</option><option>Sin dirección determinable</option></select></label><div id="estaciones"></div><noscript>Consulte la hoja Demanda estacional del Excel.</noscript>
<h2>4. Procedencia de la demanda de Bahía Jara–Levicán</h2><p>Total del nuevo cruce = desvío desde tierra + desvío desde barcaza actual + inducción incremental asignada al nuevo. Los dos primeros ya existían; el tercero corresponde a viajes adicionales atribuibles al proyecto.</p>{tabla(cap,['Participación captada de demanda existente'])}
<h2>5. Sensibilidades</h2><p><b>Vigentes:</b> curvas conservadora, central y alta con inducción central; y elasticidades menor/mayor con captación central. Frecuencia11/8 común y crecimiento central. No se suman escenarios. Los puntos de las curvas y elasticidades se encuentran en Sensibilidades del Excel.</p>
<label>Proyecto <select id="sensproyecto"><option>A1</option><option>A2</option><option>A3</option></select></label><label>Año <select id="sensanio"><option>2026</option><option>2030</option><option>2040</option><option>2050</option></select></label><div id="sensibilidad"></div>
<details><summary>Contrastes históricos: no representan la operación oficial vigente</summary><p>Las variantes de navegación120/160min, nuevo80/40min, frecuencia14/21, tarifas y buses mantienen las condiciones originales con que se calcularon. Se conservan en Sensibilidades, junto con los contrastes de crecimiento; no son nuevas sensibilidades de A1–A3 con11/8.</p>{tabla(hist[hist['Año'].eq(2026)])}</details>
<h2>6. Usuarios extranjeros y limitaciones</h2>{tabla(fx)}<p>Son indicadores muestrales diferentes y solapados: no se suman. No existe una desagregación anual defendible por alternativa, y los demás viajes no se suponen nacionales. Nacionalidad, patente, OD internacional y residencia no son equivalentes.</p><p>Las limitaciones principales son la cobertura terrestre parcial, OD residuales, accesos aproximados de algunas zonas, hipótesis de buses y costos limitados al núcleo de ciertos viajes externos. La inducción requiere costos completos comparables; sus faltantes permanecen ND. Un0 de captación dentro del soporte es un resultado calculado.</p>
<h2>7. Insumos para continuar la evaluación social</h2><p>La hoja <b>Insumos evaluacion social</b> reúne por OD, proyecto, ruta, temporada, segmento y año: demandaSBO/CP, desvíos, inducción, distancia vial, componentes temporales, COVprivado, tarifas y costos generalizados disponibles. Cada bloque de procedencia aparece en tres rutas; no sumar la demanda de referencia repetida.</p><p>Los importes están en CLPenero2026 convencionales. El CG combina costos monetarios privados con VSTsocial común de10047CLP/persona-h: <b>no es un costo social puro</b>. COVsocial y costo de recursos sociales quedanND porque no están calculados en esta entrega. Las tarifas se mantienen separadas y no deben tratarse como consumo adicional de recursos. Los costos de operación de buses coordinados deben incorporarse posteriormente.</p><p>Las hojas técnicas grises al final conservan los valores originales para consulta. <a href="demanda_oficial_tres_alternativas.xlsx">Abrir Excel oficial</a>.</p>
<script>
const estaciones={payload}, sensibilidades={senspayload};
function render(id,rows){{const e=document.getElementById(id);if(!rows.length){{e.textContent='Sin filas para esta selección; consultar cobertura en el Excel.';return;}}const cols=Object.keys(rows[0]);const table=document.createElement('table');const head=table.createTHead().insertRow();cols.forEach(c=>{{const t=document.createElement('th');t.textContent=c;head.append(t)}});const body=table.createTBody();rows.forEach(r=>{{const row=body.insertRow();cols.forEach(c=>{{const t=row.insertCell();const v=r[c];t.textContent=v===null?'ND':typeof v==='number'?(c==='Año'?v:new Intl.NumberFormat('es-CL',{{maximumFractionDigits:2}}).format(v)):v;}})}});e.replaceChildren(table);e.className='tabla';}}
function update(){{render('estaciones',estaciones.filter(r=>r.Proyecto===document.getElementById('proyecto').value&&String(r['Año'])===document.getElementById('anio').value&&r.Sentido===document.getElementById('sentido').value));render('sensibilidad',sensibilidades.filter(r=>r.Proyecto===document.getElementById('sensproyecto').value&&String(r['Año'])===document.getElementById('sensanio').value));}}document.querySelectorAll('select').forEach(s=>s.addEventListener('change',update));update();
</script></html>'''

if __name__=='__main__':presentar()
